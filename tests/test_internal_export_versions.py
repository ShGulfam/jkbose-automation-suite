import json,unittest,copy
from unittest.mock import Mock,patch
from pathlib import Path
from openpyxl import load_workbook
import test_internal_batch as fixture
import test_internal_saved_report as saved_fixture
from jkbose_tool.features.internal_assessment import Audit,match_batch
from jkbose_tool.features import internal_assessment_browser as browser
from jkbose_tool.features import internal_assessment_batch as batch
from jkbose_tool.features import internal_assessment_saved_report as saved
from jkbose_tool.features.internal_assessment_source_versions import record_source_version
from jkbose_tool.features.internal_assessment_tracking import refresh_tracking_copy,tracking_payload

class RejectedCorrectionTests(unittest.TestCase):
    setUp=fixture.ClassRunTests.setUp
    write_book=fixture.ClassRunTests.write_book
    plan=fixture.ClassRunTests.plan
    state=fixture.ClassRunTests.state
    report=saved_fixture.SavedReportTests.report
    sign=saved_fixture.SavedReportTests.sign
    def rejected(self,attendance='A',marks='0'):
        p=self.plan('Mathematics');r=self.report(p);r['records'][0].update(status='Rejected',attendance=attendance,marks=marks,rejection_reason='review');self.sign(r)
        saved.recognise_saved_awards(self.history,[p],'school-test',r,self.root/'official.json');return p,r
    def test_rejected_absence_can_change_to_present_but_normal_absence_cannot(self):
        p,r=self.rejected();d=self.state(p);d['rows'][0].update(attendance='A',disabled=True,value='0',present_input_id='present',present_enabled=True)
        self.assertFalse(match_batch(p,d,allow_absent=True)[1]);self.assertTrue(match_batch(p,d)[1]);del p.rejected_receipts
        self.assertTrue(match_batch(p,d,allow_absent=True)[1])
    def test_changed_rejected_marks_can_only_replace_exact_old_value(self):
        p,r=self.rejected(attendance='P',marks='15');d=self.state(p);d['rows'][0]['value']='15'
        self.assertFalse(match_batch(p,d,allow_absent=True)[1]);d['rows'][0]['value']='13';self.assertTrue(match_batch(p,d,allow_absent=True)[1])
    def test_present_radio_missing_or_disabled_blocks_correction(self):
        p,r=self.rejected();d=self.state(p);d['rows'][0].update(attendance='A',disabled=True,value='0',present_input_id='present',present_enabled=False)
        self.assertTrue(match_batch(p,d,allow_absent=True)[1])
    def test_rejected_correction_full_fill_save_receipt_then_history_skip(self):
        p,r=self.rejected();d=self.state(p);row=d['rows'][0];row.update(attendance='A',disabled=True,value='0',present_input_id='present',present_enabled=True);driver=Mock();clicks=[]
        def element(by,identity):
            e=Mock()
            if identity=='present':
                e.is_enabled.return_value=True;e.get_attribute.return_value='P';e.click.side_effect=lambda:(clicks.append(identity),row.update(attendance='P',disabled=False))
            elif identity=='bodypH_btnSaveMarks':e.click.side_effect=lambda:clicks.append(identity)
            elif identity==row['input_id']:
                e.get_attribute.side_effect=lambda key:row['value'] if key=='value' else ''
                e.clear.side_effect=lambda:row.update(value='');e.send_keys.side_effect=lambda value:row.update(value=value)
            else:raise AssertionError(identity)
            return e
        driver.find_element.side_effect=element
        def receipt(plan,awards,batch_id):return dict(official_statuses=saved.verify_saved_batch(plan,awards,'school-test',self.report(plan)))
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=lambda *a,**k:copy.deepcopy(d)),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_begin_read'),patch.object(browser,'_wait_page'):
            result=browser.run(driver,p,self.history,'school-test',Mock(return_value=True),submit=True,confirm_session=True,receipt_callback=receipt)
        self.assertEqual(clicks,['present','bodypH_btnSaveMarks']);self.assertEqual(row['value'],'12');self.assertEqual(result['status'],'saved_subject_verified')
        self.assertEqual(batch.verified_history(self.history,p,'school-test')['verified'],{'REG1'})
    def test_new_export_differing_from_accepted_receipt_is_report_only(self):
        old=self.plan('Mathematics');r=self.report(old);self.write_book(maths=19);p=self.plan('Mathematics')
        result=saved.recognise_saved_awards(self.history,[p],'school-test',r,self.root/'official.json')
        self.assertEqual(result['Mathematics'],0);prior=batch.verified_history(self.history,p,'school-test');self.assertEqual(prior['verified'],{'REG1'});self.assertEqual(prior['proofs'][0]['kind'],'recorded_different')
        audit=Audit(self.history,p);self.assertFalse(audit.verified());self.assertFalse(audit.unresolved());self.assertIn('official_recorded_award',[e['kind'] for e in audit.events])
    def test_unknown_identity_never_gets_report_only_skip(self):
        p=self.plan('Mathematics');r=self.report(p);r['records'][0]['exam_roll']='301000009';self.sign(r)
        with self.assertRaisesRegex(ValueError,'registration/exam pair conflicts'):saved.recognise_saved_awards(self.history,[p],'school-test',r,self.root/'official.json')

class ExportVersionTests(unittest.TestCase):
    setUp=fixture.ClassRunTests.setUp
    write_book=fixture.ClassRunTests.write_book
    plan=fixture.ClassRunTests.plan
    report=saved_fixture.SavedReportTests.report
    sign=saved_fixture.SavedReportTests.sign
    def version(self):return record_source_version(self.file,self.history,'11','2025-26','school-test')
    def test_same_export_is_idempotent_and_next_export_lists_changes(self):
        a=self.version();self.assertFalse(a['changes']);self.assertEqual(self.version()['time'],a['time']);self.write_book(maths=19);b=self.version();changes=[x for x in b['changes'] if x['subject']=='Mathematics'];self.assertEqual([(x['field'],str(x['previous']),str(x['current'])) for x in changes],[('Marks','12','19')]);self.assertEqual(self.version()['changes'],b['changes'])
    def test_changed_filename_still_uses_same_class_session_history(self):
        self.version();self.write_book(maths=19);new=self.file.with_name('new-consolidated-export.xlsx');new.write_bytes(self.file.read_bytes());b=record_source_version(new,self.history,'11','2025-26','school-test');self.assertTrue(b['changes']);self.assertEqual(b['previous_source'],self.file.name)
    def test_previous_tracking_copy_is_first_run_baseline(self):
        self.version();output=refresh_tracking_copy(self.file,self.history,'11','2025-26','school-test');other=self.root/'other-history';self.write_book(maths=19);v=record_source_version(self.file,other,'11','2025-26','school-test');self.assertTrue(any(c['field']=='Marks' and str(c['previous'])=='12' for c in v['changes']))
    def test_tracking_changes_distinguish_recorded_from_pending_and_keep_source_immutable(self):
        old=self.plan('Mathematics');r=self.report(old);folder=self.history/'_class_runs'/'synthetic';folder.mkdir(parents=True);(folder/'official_saved_assessment_report.json').write_text(json.dumps(r));self.version();self.write_book(maths=19,physics=9);before=self.file.read_bytes();output=refresh_tracking_copy(self.file,self.history,'11','2025-26','school-test');self.assertEqual(self.file.read_bytes(),before)
        book=load_workbook(output);self.assertIn('Source_Changes',book.sheetnames);rows=list(book['Source_Changes'].values)[5:];maths=next(x for x in rows if x[2]=='Mathematics');physics=next(x for x in rows if x[2]=='Physics');self.assertEqual(maths[10],'Submitted');self.assertIn('never retry',maths[11]);self.assertEqual(physics[10],'Pending');self.assertIn('current export',physics[11]);book.close()
        p=tracking_payload(self.file,self.history,'11','2025-26','school-test');self.assertEqual(p['flags'][1][0],'Submitted');self.assertIn('no retry',p['flags'][1][5]);json.dumps(p)
    def test_tampered_snapshot_is_rejected(self):
        self.version();path=next((self.history/'_source_versions').rglob('*.json'));path.write_text('{}')
        with self.assertRaisesRegex(ValueError,'snapshot was changed'):self.version()
    def test_scope_change_never_imports_another_class(self):
        with self.assertRaisesRegex(ValueError,'cohort differs'):record_source_version(self.file,self.history,'12','2025-26','school-test')
    def test_removed_award_is_reported_without_inference_of_portal_deletion(self):
        self.version();book=load_workbook(self.file);book.active.delete_rows(3);book.save(self.file);book.close();v=self.version();self.assertTrue(any(x['change']=='Removed' and x['subject']=='Mathematics' for x in v['changes']))

if __name__=='__main__':unittest.main()
