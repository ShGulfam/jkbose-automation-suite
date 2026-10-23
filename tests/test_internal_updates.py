"""Regression coverage for Excel limits, optional Biology and safe live navigation."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
from openpyxl import Workbook
from jkbose_tool.features import internal_assessment as model
from jkbose_tool.features import internal_assessment_browser as browser

class InternalUpdateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name); self.file = self.root / 'synthetic.xlsx'

    def workbook(self, entries):
        book = Workbook(); ws = book.active; ws.title = 'JKBOSE_Internal_Assessment'
        ws.append(['Class','Session','Evaluation Type','Subject Code','Subject Name','Board Reg. No.','Exam Roll No.','Class Roll No.','Student Name','Marks','Max Marks','Attendance','Source Issues','Export Schema','Source Award ID'])
        for code, label, marks, maximum, reg in entries:
            ws.append(['12th','2025-26','Internal Assessment',code,label,reg,'301000001','1','Synthetic Student',marks,maximum,'Present','','hss-jkbose-ia-v1','synthetic'])
        book.save(self.file); book.close()

    def plan(self, **kwargs):
        return model.load_plan(self.file, student_class='12', session='2025-26', **kwargs)

    def test_auto_maximum_and_threshold_come_from_excel(self):
        self.workbook([('MA','Mathematics',8,20,'REG1')])
        plan = self.plan(subject='Mathematics')
        self.assertEqual(plan.maximum, '20'); self.assertEqual(plan.pass_marks, '8')
        self.assertFalse(plan.awards[0].issues)
        self.assertTrue(self.plan(subject='Mathematics', maximum='10').awards[0].issues)

    def test_unequal_biology_parts_sum_both_maxima_and_validate_live(self):
        self.workbook([('BO','Botany',2.5,4,'REG1'),('ZO','Zoology',3.5,6,'REG1')])
        plan = self.plan()
        self.assertEqual(plan.maximum,'10'); self.assertEqual(plan.awards[0].marks,'6.0')
        self.assertEqual(plan.component_maxima, {'Botany':'4','Zoology':'6'})
        self.assertFalse(plan.awards[0].issues)
        row = dict(regno='REG1',name='Synthetic Student',exam_roll='301000001',subject='Biology',attendance='P',max='10',min='0',step='any',value='')
        self.assertFalse(model.match_batch(plan, {'rows':[row]})[1])
        plan.awards[0].components['Botany']='5'
        self.assertTrue(model.match_batch(plan, {'rows':[row]})[1])

    def test_checkbox_can_import_completed_biology_or_separate_parts(self):
        self.workbook([('BI','Biology',6,10,'REG1')])
        plan = self.plan(combine_biology=False)
        self.assertEqual(plan.awards[0].marks,'6'); self.assertFalse(plan.awards[0].issues)
        with self.assertRaises(ValueError): self.plan()
        self.workbook([('BO','Botany',3,5,'REG1')])
        self.assertFalse(self.plan(subject='Botany',combine_biology=False).awards[0].issues)
        with self.assertRaises(ValueError): self.plan(subject='Botany')

    def test_conflicting_missing_zero_and_formula_maxima_stop_auto(self):
        for limit in ['',0,'=5+5']:
            self.workbook([('PH','Physics',5,limit,'REG1')])
            with self.subTest(limit=limit), self.assertRaises(ValueError): self.plan(subject='Physics')
        self.workbook([('PH','Physics',5,10,'REG1'),('PH','Physics',5,20,'REG2')])
        with self.assertRaises(ValueError): self.plan(subject='Physics')

    def test_matrix_auto_reads_maximum_row_and_direct_total_when_unchecked(self):
        book=Workbook(); ws=book.active; ws.title='Awards_Matrix'
        ws.append(['INTERNAL ASSESSMENT — Class 12th']); ws.append(['Session: 2025-26'])
        ws.append(['Maximum Marks','','','','',4,6,10])
        ws.append(['S.No.','Class Roll No.','Exam Roll No.','Board Reg. No.','Student Name','Botany (BO)','Zoology (ZO)','Biology (BI)'])
        ws.append([1,'1','301000001','REG1','Synthetic Student',2.5,3.5,6])
        book.save(self.file); book.close()
        self.assertEqual(self.plan().maximum,'10')
        self.assertFalse(self.plan().awards[0].issues)
        self.assertEqual(self.plan(combine_biology=False).awards[0].marks,'6')

    def test_history_migrates_intact_and_preserves_subsequent_new_events(self):
        legacy=self.root/'.jkbose_tool_private'/'internal_assessment'
        legacy.mkdir(parents=True); (legacy/'evidence.txt').write_text('original')
        target=self.root/'tool'/'history'/'internal_assessment'
        with patch.object(model,'DEFAULT_HISTORY_ROOT',target), patch.object(Path,'home',return_value=self.root):
            self.assertEqual(model.prepare_history_root(),target.resolve()/'References')
            self.assertEqual((target/'References'/'evidence.txt').read_text(),'original')
            (target/'References'/'evidence.txt').write_text('new events')
            model.prepare_history_root()
        self.assertEqual((target/'References'/'evidence.txt').read_text(),'new events')
        self.assertEqual((legacy/'evidence.txt').read_text(),'original')

    def test_history_collision_and_legacy_lock_stop_migration(self):
        legacy=self.root/'.jkbose_tool_private'/'internal_assessment'; legacy.mkdir(parents=True)
        target=self.root/'history'; target.mkdir()
        (legacy/'evidence').write_text('original'); (target/'evidence').write_text('different')
        with patch.object(model,'DEFAULT_HISTORY_ROOT',target), patch.object(Path,'home',return_value=self.root):
            with self.assertRaises(ValueError): model.prepare_history_root()
            (legacy/'internal-assessment.lock').write_text('123')
            with self.assertRaises(ValueError): model.prepare_history_root()

    def test_old_history_fingerprints_still_load(self):
        self.workbook([('MA','Mathematics',8,20,'REG1')]); plan=self.plan(subject='Mathematics')
        plan.component_maxima=None
        audit=model.Audit(self.root/'history',plan)
        audit.append('offline_preview',awards=plan.rows())
        recovered, old=model.Audit.existing(audit.folder)
        self.assertEqual(old.digest(),plan.digest())
        self.assertEqual(recovered.folder,audit.folder)

    def test_open_subject_uses_verified_sidebar_instead_of_dashboard_redirect(self):
        self.workbook([('MA','Mathematics',8,20,'REG1')]); plan=self.plan(subject='Mathematics')
        driver=Mock(); driver.current_url='https://jkbose.jk.gov.in/utilities/auth_pages/subadminpanelforclerk.aspx'
        link=Mock(); link.get_attribute.return_value='practicalMarks.aspx'
        link.click.side_effect=lambda: setattr(driver,'current_url',browser.URL)
        driver.find_elements.return_value=[link]
        selected=Mock(); selected.first_selected_option.get_attribute.side_effect=['12','Mathematics']
        option=Mock(); option.get_attribute.return_value='Mathematics'; selected.options=[option]
        with patch.object(browser,'Select',return_value=selected), patch.object(browser,'WebDriverWait'), patch.object(browser,'show_students') as show:
            browser.open_subject(driver,plan,None)
        link.click.assert_called_once(); driver.get.assert_not_called(); show.assert_called_once()

    def test_dashboard_redirect_never_passes_navigation_validation(self):
        self.workbook([('MA','Mathematics',8,20,'REG1')]); plan=self.plan(subject='Mathematics')
        driver=Mock(); driver.current_url='https://jkbose.jk.gov.in/utilities/auth_pages/subadminpanelforclerk.aspx'
        driver.find_elements.return_value=[]
        with patch.object(browser,'WebDriverWait'), self.assertRaises(ValueError): browser.open_subject(driver,plan,None)
        driver.find_element.assert_not_called()

    def test_explicit_sessions_do_not_collapse_academic_years(self):
        for raw, expected in [('2026-27','2026-27'),('Annual Regular 2027','2026-27'),('2023–2024','2023-24')]:
            self.assertEqual(model.session_key(raw),expected)
        with self.assertRaises(ValueError): model.session_key('2025-27')

    def test_all_absent_fail_and_zero_sources_block_before_filling_or_save(self):
        for mark, attendance, source_issue in [('AB','Absent',''),(0,'Present',''),(3.99,'Present',''),(8,'Absent',''),(8,'Present','Failed: manual submission only')]:
            self.workbook([('MA','Mathematics',mark,20,'REG1')])
            from openpyxl import load_workbook
            book=load_workbook(self.file); ws=book.active
            ws.cell(2,12,attendance); ws.cell(2,13,source_issue); book.save(self.file); book.close()
            plan=self.plan(subject='Mathematics')
            row=dict(regno='REG1',name='Synthetic Student',exam_roll='301000001',subject='Mathematics',attendance='P',max='20',min='0',step='any',value='')
            state={'rows':[row],'total':1}; driver=Mock()
            with self.subTest(mark=mark,attendance=attendance,issue=source_issue), patch.object(browser,'open_subject'), patch.object(browser,'snapshot',return_value=state), patch.object(browser,'_evidence',return_value='synthetic'), patch.object(browser,'_fill') as fill:
                browser.run(driver,plan,self.root/'history','school-test',Mock(return_value=True),submit=True,confirm_session=True)
                fill.assert_not_called(); driver.find_element.assert_not_called()

    def test_official_maximum_is_authority_for_fail_and_over_max_checks(self):
        self.workbook([('MA','Mathematics',9,20,'REG1')]); plan=self.plan(subject='Mathematics')
        row=dict(regno='REG1',name='Synthetic Student',exam_roll='301000001',subject='Mathematics',attendance='P',max='25',min='0',step='any',value='')
        problems=model.match_batch(plan, {'rows':[row],'maximum':'25','hidden_maximum':'25'})[1]
        self.assertTrue(any('below 40% of the official' in message for message in problems))
        row['max']='5'
        problems=model.match_batch(plan, {'rows':[row],'maximum':'5','hidden_maximum':'5'})[1]
        self.assertTrue(any('exceeds the official' in message for message in problems))
        with self.assertRaises(ValueError): model.live_maximum({'rows':[row],'maximum':'20','hidden_maximum':'5'})

    def test_history_reports_both_limits_and_manual_only_official_failures(self):
        self.workbook([('MA','Mathematics',9,20,'REG1')]); plan=self.plan(subject='Mathematics')
        audit=model.Audit(self.root/'history',plan)
        audit.append('official_maximum_checked',maximum='25',pass_marks='10',excel_maximum='20',mismatch=True)
        audit.report(plan)
        report=model.json.loads((audit.folder/'report.json').read_text())
        self.assertEqual(report['official_jkbose_maximum'],'25')
        self.assertEqual(report['records'][0]['status'],'manual_required')
        self.assertTrue(any('Below official 40%' in reason for reason in report['records'][0]['issues']))

if __name__=='__main__': unittest.main()
