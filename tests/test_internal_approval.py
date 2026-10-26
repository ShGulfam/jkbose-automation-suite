"""Synthetic Principal workflow; never connects to JKBOSE or clicks a live decision."""
from pathlib import Path
from unittest.mock import Mock,patch
from copy import deepcopy
from html import escape
import json,unittest,threading
from openpyxl import load_workbook
import test_internal_batch as fixtures
import test_internal_students_report as roster_fixtures
from jkbose_tool.features import internal_assessment_approval as approval
from jkbose_tool.features.internal_assessment import prepare_history_root

class FakePortal:
    def __init__(self,rows):self.rows=deepcopy(rows);self.opens=[];self.confirms=[];self.reads=0;self.fail_after_save=False;self.wrong_target=False
    def read_all(self):
        self.reads+=1
        return deepcopy(self.rows),dict(file='synthetic-export.xls',sha256='synthetic',total=len(self.rows))
    def open_approval(self,r):
        self.opens.append(r['id'])
        if self.wrong_target:raise ValueError('Popup target differs')
    def confirm(self,r):
        self.confirms.append(r['id']);self.rows=[x for x in self.rows if x['id']!=r['id']]
        if self.fail_after_save:raise RuntimeError('Disconnected after decision')

class PrincipalTests(unittest.TestCase):
    setUp=fixtures.ClassRunTests.setUp
    write_book=fixtures.ClassRunTests.write_book
    def jobs(self,subjects=None):
        from jkbose_tool.features.internal_assessment_batch import load_jobs
        return load_jobs(self.file,student_class='11',session='2025-26',subjects=subjects or ['MA','PH'])
    def rows(self,jobs=None):
        return [dict(id=str(i),class_name='11',regno=a.regno,exam_roll=a.exam_roll,name=a.name,subject=j['plan'].subject,
                     marks=a.marks,maximum=j['plan'].maximum,attendance='P',requested_on='23 Oct 2026')
                for i,(j,a) in enumerate([(j,a) for j in jobs or self.jobs() for a in j['plan'].awards],1)]
    def roster(self):
        data=roster_fixtures.StudentsReportTests().snapshot();data['account']='prin000001'
        if any(r['regno']=='REG2' for r in self.rows()):
            second=deepcopy(data['rows'][0]);second.update(sequence='2',exam_roll='301000002',regno='REG2',name='Synthetic Two');data['rows'].append(second);data['row_widths'].append(21)
        return data
    def run_approval(self,portal,execute=False,confirm=None,jobs=None,**kwargs):
        driver=Mock()
        options=dict(execute=execute,confirm_session=True,driver_factory=Mock(return_value=driver),login_callback=Mock(),report_reader=Mock(return_value=self.roster()),
                     portal_factory=Mock(return_value=portal),confirm_callback=confirm or Mock(return_value=True))
        options.update(kwargs)
        result=approval.run_approval(jobs or self.jobs(),self.file,self.history,'sch000001','prin000001',**options)
        return result,driver,options
    def journal(self):return approval.ApprovalJournal(prepare_history_root(self.history),dict(school='sch000001',principal='prin000001',**{'class':'11'},session='2025-26'))
    def test_readonly_never_opens_or_confirms_decisions(self):
        portal=FakePortal(self.rows());confirm=Mock(side_effect=AssertionError('No decision confirmation in read-only'))
        result,driver,options=self.run_approval(portal,confirm=confirm)
        self.assertEqual(result['status'],'checked_read_only');self.assertEqual(portal.opens,[]);self.assertEqual(portal.confirms,[]);confirm.assert_not_called();driver.quit.assert_called_once()
        self.assertIn('.chrome_profile_principal',options['driver_factory'].call_args.kwargs['profile_dir']);self.assertEqual(options['login_callback'].call_args.kwargs['role'],'principal')
    def test_multiple_subjects_complete_with_one_preview_and_exact_history(self):
        portal=FakePortal(self.rows());confirm=Mock(return_value=True)
        result,_,_=self.run_approval(portal,True,confirm)
        self.assertEqual(result['status'],'finished');self.assertEqual(portal.confirms,['1','2']);confirm.assert_called_once();self.assertEqual(self.journal().unresolved(),{})
        self.assertEqual(self.journal().approved(),{'1','2'})
        self.assertTrue((self.journal().folder/'OPEN APPROVAL HISTORY.html').is_file())
    def test_decline_preview_stops_before_popup(self):
        portal=FakePortal(self.rows());result,_,_=self.run_approval(portal,True,Mock(return_value=False))
        self.assertEqual(result['status'],'cancelled');self.assertEqual(portal.opens,[])
    def test_wrong_accounts_or_missing_session_stops_before_browser(self):
        for principal in ['sch000001','prin999999']:
            factory=Mock()
            with self.assertRaises(ValueError):approval.run_approval(self.jobs(),self.file,self.history,'sch000001',principal,confirm_session=True,driver_factory=factory)
            factory.assert_not_called()
        with self.assertRaises(ValueError):self.run_approval(FakePortal(self.rows()),True,confirm_session=False)
    def test_queue_change_after_preview_stops_without_decision(self):
        portal=FakePortal(self.rows())
        def change(data):portal.rows[0]['marks']='1';return True
        with self.assertRaisesRegex(ValueError,'changed after preview'):self.run_approval(portal,True,change)
        self.assertEqual(portal.opens,[])
    def test_mixed_valid_and_absent_students_only_approves_valid_student(self):
        self.write_book(second=True);book=load_workbook(self.file);book.active.cell(2,10,'AB');book.active.cell(2,12,'Absent');book.save(self.file);book.close()
        jobs=self.jobs(['MA']);portal=FakePortal(self.rows(jobs));result,_,_=self.run_approval(portal,True,jobs=jobs)
        self.assertEqual(result['status'],'finished_with_holds');self.assertEqual(portal.confirms,['2'])
    def test_wrong_name_marks_maximum_or_attendance_never_approved(self):
        for field,value in [('name','Other Person'),('marks','11'),('maximum','30'),('attendance','A'),('exam_roll','301000099')]:
            with self.subTest(field=field):
                jobs=self.jobs(['MA']);rows=self.rows(jobs);rows[0][field]=value;portal=FakePortal(rows)
                result,_,_=self.run_approval(portal,True,jobs=jobs)
                self.assertEqual(result['status'],'nothing_eligible');self.assertEqual(portal.confirms,[])
    def test_uncertain_after_save_never_retries_and_requires_reconciliation(self):
        portal=FakePortal(self.rows());portal.fail_after_save=True
        with self.assertRaisesRegex(RuntimeError,'Disconnected'):self.run_approval(portal,True)
        self.assertEqual(portal.confirms,['1']);self.assertEqual(set(self.journal().unresolved()),{'1'})
        second=FakePortal(self.rows())
        result,_,_=self.run_approval(second,True)
        self.assertEqual(result['status'],'finished_with_holds');self.assertEqual(second.confirms,['2']);self.assertEqual(set(self.journal().unresolved()),{'1'})
        journal=self.journal();journal.reconcile('1','approved','Synthetic Admin','Exact request verified Approved in the official report.')
        self.assertEqual(self.journal().unresolved(),{})
    def test_wrong_popup_target_does_not_confirm_and_keeps_uncertainty(self):
        portal=FakePortal(self.rows());portal.wrong_target=True
        with self.assertRaisesRegex(ValueError,'Popup target'):self.run_approval(portal,True)
        self.assertEqual(portal.confirms,[]);self.assertEqual(set(self.journal().unresolved()),{'1'})
    def test_no_export_removal_of_different_or_additional_awards_is_accepted(self):
        rows=self.rows()
        for after in [rows,[],[rows[0]],[dict(rows[1],marks='1')]]:
            with self.assertRaises(ValueError):approval.verify_removal(rows,after,rows[0])
        approval.verify_removal(rows,[rows[1]],rows[0])
    def test_integrity_and_manual_resolution_require_specific_evidence(self):
        journal=self.journal();journal.append('approval_open_attempt',request=self.rows()[0])
        with self.assertRaises(ValueError):journal.reconcile('1','approved','Admin','okay')
        text=journal.path.read_text();journal.path.write_text(text.replace('Synthetic One','Changed Person'))
        with self.assertRaisesRegex(ValueError,'integrity'):self.journal()
    def test_changed_source_after_preview_prevents_any_decision(self):
        portal=FakePortal(self.rows())
        def change(data):self.write_book(maths=13);return True
        with self.assertRaisesRegex(ValueError,'Source workbook changed'):self.run_approval(portal,True,change)
        self.assertEqual(portal.opens,[])
    def test_stop_interrupts_before_opening_browser(self):
        event=threading.Event();event.set();factory=Mock()
        with self.assertRaises(Exception):self.run_approval(FakePortal(self.rows()),True,stop_event=event,driver_factory=factory)
        factory.assert_not_called()
    def test_parser_reads_all_pages_export_and_rejects_duplicate_or_bad_headers(self):
        rows=self.rows();path=self.root/'pending.xls'
        def write(records,headers=approval.EXPORT_HEADERS):
            records=[[r['id'],'11',r['exam_roll'],r['regno'],r['name'],r['subject'],'Present',r['marks'],r['maximum'],r['requested_on']] for r in records]
            path.write_text('<table>'+''.join('<tr>'+''.join('<td>'+escape(v)+'</td>' for v in row)+'</tr>' for row in [headers]+records)+'</table>')
        write(rows);self.assertEqual(approval.parse_export(path,'11',2),rows)
        write(rows+rows)
        with self.assertRaises(ValueError):approval.parse_export(path,'11',4)
        write(rows,headers=approval.EXPORT_HEADERS[:-1])
        with self.assertRaises(ValueError):approval.parse_export(path,'11',2)
    def test_parser_empty_queue_and_incomplete_coverage(self):
        path=self.root/'pending.xls';path.write_text('<table><tr>'+''.join('<th>'+h+'</th>' for h in approval.EXPORT_HEADERS)+'</tr></table>')
        self.assertEqual(approval.parse_export(path,'11',0),[])
        with self.assertRaises(ValueError):approval.parse_export(path,'11',1)
    def test_popup_hidden_id_binding_blocks_confirmation(self):
        portal=approval.PrincipalPortal(Mock(),self.root/'downloads','11','prin000001')
        portal.snapshot=Mock(return_value=dict(approve_visible=True,reject_visible=False,approve_id='WRONG'))
        with self.assertRaises(ValueError):portal.confirm(self.rows()[0])
        portal.driver.find_element.assert_not_called()
    def test_seek_finds_target_after_pagination_without_decision_clicks(self):
        portal=approval.PrincipalPortal(Mock(),self.root/'downloads','11','prin000001');target=self.rows()[0]
        raw=dict(cells=['11',target['exam_roll'],target['regno'],target['name'],target['subject'],target['marks'],target['maximum'],target['requested_on']],approve='bodypH_gridRR_btnApprove_0')
        portal.page=Mock(side_effect=[dict(total='11',rows=[]),dict(total='11',rows=[raw])])
        self.assertEqual(portal.seek(target)[0],raw);self.assertEqual([c.args[0] for c in portal.page.call_args_list],[1,2]);portal.driver.find_element.assert_not_called()
    def test_verified_approval_updates_tracking_without_marking_unapproved_subjects(self):
        portal=FakePortal(self.rows());result,_,_=self.run_approval(portal,True)
        from jkbose_tool.features.internal_assessment_tracking import tracking_payload
        p=tracking_payload(self.file,self.history,'11','2025-26','sch000001')
        maths=next(r for r in p['summary'] if r[0]=='Mathematics');self.assertEqual(maths[4],1)
        self.assertEqual(p['flags'][1][:3],['Submitted','Approved','Approved (tool verified)'])
        bio=next(r for r in p['summary'] if r[0]=='Biology');self.assertEqual(bio[4],0)
    def test_cli_principal_defaults_are_read_only(self):
        import main
        args=main.build_parser().parse_args(['internal-approve','--excel-file','synthetic.xlsx','--class','11','--session','2025-26','--subjects','MA','--school-account','sch000001','--principal-account','prin000001'])
        self.assertFalse(args.execute);self.assertFalse(args.confirm_session)
    def test_final_empty_queue_with_no_total_label_is_explicitly_recognised(self):
        data=dict(account='prin000001',**{'class':'11'},headers=['There are no pending requests to display.'],rows=[],total='',empty_text='There are no pending requests to display.',page=[],pager=[])
        self.assertEqual(approval.validate_grid(data,'11','prin000001'),0)
        data['total']='';data['empty_text']=''
        with self.assertRaises(ValueError):approval.validate_grid(data,'11','prin000001')
    def test_single_page_queue_without_a_pager_requires_every_row(self):
        request=self.rows()[0];raw=dict(cells=['11',request['exam_roll'],request['regno'],request['name'],request['subject'],request['marks'],request['maximum'],request['requested_on']],approve='bodypH_gridRR_btnApprove_0')
        data=dict(account='prin000001',**{'class':'11'},headers=approval.HEADERS,rows=[raw],total='1',page=[],pager=[])
        self.assertEqual(approval.validate_grid(data,'11','prin000001'),1);self.assertEqual(data['page'],['1'])
        data['total']='2'
        with self.assertRaises(ValueError):approval.validate_grid(data,'11','prin000001')
    def test_empty_portal_does_not_try_missing_download_button(self):
        portal=approval.PrincipalPortal(Mock(),self.root/'downloads','11','prin000001');portal.open=Mock(return_value=dict(total='0',approve_visible=False,reject_visible=False))
        self.assertEqual(portal.read_all()[0],[]);portal.driver.find_element.assert_not_called()
    def test_resume_empty_queue_keeps_approved_receipts_without_repeating_actions(self):
        first=FakePortal(self.rows());self.run_approval(first,True)
        second=FakePortal([]);confirm=Mock(side_effect=AssertionError('No repeat confirmation'))
        result,_,_=self.run_approval(second,True,confirm)
        self.assertEqual(result['status'],'nothing_eligible');self.assertEqual(second.opens,[]);confirm.assert_not_called();self.assertEqual(self.journal().approved(),{'1','2'})
    def test_principal_popup_launches_readonly_mode_without_live_browser(self):
        import tkinter as tk
        import gui_main
        try:root=tk.Tk()
        except tk.TclError:self.skipTest('No desktop available')
        root.withdraw()
        try:
            app=gui_main.App(root);app._show_section('internal');app.ia['account'].set('sch000001')
            with patch.object(app,'_start_task') as start,patch.object(app,'_internal_jobs',return_value=self.jobs()):
                app._open_principal_approval();root.update_idletasks()
                def children(widget):
                    for child in widget.winfo_children():yield child;yield from children(child)
                button=next(w for w in children(root) if w.winfo_class()=='TButton' and w.cget('text')=='Check requests (read only)')
                button.invoke();self.assertFalse(start.call_args.args[2]['execute']);self.assertEqual(start.call_args.args[2]['principal_account'],'prin000001')
        finally:root.destroy()
    def test_unqueued_subject_roster_does_not_block_queued_selected_subject(self):
        jobs=self.jobs();portal=FakePortal(self.rows(self.jobs(['MA'])));roster=self.roster();roster['rows'][0]['subjects']=['Mathematics','','','','']
        result,_,_=self.run_approval(portal,True,jobs=jobs,report_reader=Mock(return_value=roster))
        self.assertEqual(result['status'],'finished');self.assertEqual(portal.confirms,['1'])
    def test_one_roster_identity_conflict_holds_only_that_request(self):
        self.write_book(second=True);jobs=self.jobs(['MA']);portal=FakePortal(self.rows(jobs));roster=self.roster();roster['rows'][0]['name']='Conflicting Student'
        result,_,_=self.run_approval(portal,True,jobs=jobs,report_reader=Mock(return_value=roster))
        self.assertEqual(result['status'],'finished_with_holds');self.assertEqual(portal.confirms,['2'])
    def test_popup_opener_validates_hidden_request_before_confirmation(self):
        request=self.rows()[0];portal=approval.PrincipalPortal(Mock(),self.root/'downloads','11','prin000001')
        data=dict(total='1',approve_visible=False,reject_visible=False)
        row=dict(approve='bodypH_gridRR_btnApprove_0');portal.seek=Mock(return_value=(row,data))
        portal.snapshot=Mock(return_value=dict(total='1',approve_visible=True,reject_visible=False,approve_id='OTHER',rows=[]))
        with patch.object(approval,'begin_read'),patch.object(approval,'wait_read'),self.assertRaisesRegex(ValueError,'not bound'):
            portal.open_approval(request)
        self.assertEqual(portal.driver.find_element.call_args.args,(approval.By.ID,'bodypH_gridRR_btnApprove_0'))
    def test_principal_report_wrong_school_stops_without_decision(self):
        portal=FakePortal(self.rows());roster=self.roster();roster['account']='prin999999'
        with self.assertRaisesRegex(ValueError,'account/session differs'):self.run_approval(portal,True,report_reader=Mock(return_value=roster))
        self.assertEqual(portal.opens,[])
    def test_reconciliation_scope_cannot_escape_private_history(self):
        scope=dict(school='sch000001',principal='prin000001',**{'class':'11'},session='2025-26/../../outside')
        with self.assertRaises(ValueError):approval.ApprovalJournal(self.history,scope)
    def test_manual_approval_provenance_is_not_claimed_as_a_tool_action(self):
        journal=self.journal();journal.append('approval_open_attempt',request=self.rows()[0]);journal.reconcile('1','approved','Synthetic Admin','Exact approval confirmed from the official saved report.')
        receipts,_=approval.tracking_receipts(prepare_history_root(self.history),'sch000001','11','2025-26')
        self.assertEqual(next(iter(receipts.values()))['verified_method'],'manual')
