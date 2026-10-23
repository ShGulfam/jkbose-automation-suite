import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from openpyxl import Workbook, load_workbook
from selenium.common.exceptions import NoAlertPresentException
from jkbose_tool.features.exam_input import read_exam_dataframe, load_check_inputs, EXAM_SCHEMA
from jkbose_tool.features.exam_guard import require_exam_form, ExamPortalUnavailable
from jkbose_tool.features import inspect, submit_form
import main

class ExamSafetyTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.file = Path(self.tmp.name) / 'exam.xlsx'

    def workbook(self, cls='12', **extra):
        row = dict(RegNo='2301000000000037', Class=cls, Session='2025-26', ClassRollNo='37', Name='Synthetic Student', ExportSchema=EXAM_SCHEMA, ApprovedForAutomation='Yes', Status='', **extra)
        book=Workbook();ws=book.active;ws.title='Sheet2';ws.append(list(row));ws.append(list(row.values()));book.save(self.file);book.close()

    def test_export_is_accepted_and_exact_text_is_preserved(self):
        self.workbook()
        before=self.file.read_bytes()
        rows=load_check_inputs(self.file)
        self.assertEqual(rows[0]['RegNo'],'2301000000000037')
        self.assertEqual(before,self.file.read_bytes())

    def test_wrong_class_and_formula_and_unassigned_are_blocked_before_browser(self):
        self.workbook()
        with self.assertRaisesRegex(ValueError,'another class'):read_exam_dataframe(self.file,class_override='11')
        for column,value in [('RegNo','=1+1'),('ClassRollNo','')]:
            self.workbook();book=load_workbook(self.file);ws=book['Sheet2'];headers=[c.value for c in ws[1]];ws.cell(2,headers.index(column)+1,value);book.save(self.file);book.close()
            with patch.object(main,'get_driver') as browser:
                with self.assertRaises(ValueError):main.cmd_submit(self.file)
                browser.assert_not_called()

    def test_dropped_and_completed_only_batches_do_not_open_browser(self):
        for update in [{'ExamDropped':'Yes'}, {'Status':'Submitted - verified'}]:
            self.workbook();book=load_workbook(self.file);ws=book['Sheet2'];headers=[c.value for c in ws[1]]
            for key,value in update.items():
                if key not in headers:ws.cell(1,len(headers)+1,key);headers.append(key)
                ws.cell(2,headers.index(key)+1,value)
            book.save(self.file);book.close()
            with patch.object(main,'get_driver') as browser:main.cmd_submit(self.file);browser.assert_not_called()

    def test_progress_keeps_original_physical_row_after_blank_rows(self):
        self.workbook();book=load_workbook(self.file);book['Sheet2'].insert_rows(2);book.save(self.file);book.close()
        self.assertEqual(list(read_exam_dataframe(self.file).index),[1])

    def test_expired_and_closed_portals_are_blocked(self):
        for url,text in [('https://jkbose.jk.gov.in/utilities/auth_pages/login.aspx',''),('https://jkbose.jk.gov.in/utilities/auth_pages/renewalformxixii.aspx','RR Renewal Is Closed'),('https://example.org/renewalformxixii.aspx','')]:
            driver=Mock(current_url=url);driver.find_element.return_value.text=text
            with self.assertRaises(ExamPortalUnavailable):require_exam_form(driver)

    def test_inspect_is_read_only_by_default_and_redacts_credentials(self):
        driver=Mock(current_url='https://jkbose.jk.gov.in/utilities/auth_pages/renewalformxixii.aspx')
        node=Mock(text='secret-value')
        node.get_attribute.side_effect=lambda key:{'id':'__VIEWSTATE','type':'hidden','value':'secret-value'}.get(key,'')
        driver.find_elements.side_effect=lambda by,key:[node] if key=='input' else []
        with patch.object(inspect,'_inspect_all_streams') as cycle:
            report=inspect.run(driver,output_csv=str(Path(self.tmp.name)/'fields.csv'))
            cycle.assert_not_called();driver.get.assert_not_called()
        self.assertEqual(report['fields'][0]['value'],'[redacted]')
        self.assertNotIn('secret-value',(Path(self.tmp.name)/'fields.csv').read_text(encoding='utf-8-sig'))

    def test_instructional_body_is_never_a_completion_receipt(self):
        driver=Mock();driver._exam_lookup_context=('2301000000000037','12')
        type(driver.switch_to).alert=property(lambda _:(_ for _ in ()).throw(NoAlertPresentException()))
        driver.find_elements.return_value=[]
        driver.execute_script.return_value='If already submitted, do not submit again'
        with patch.object(submit_form,'require_exam_form'):
            self.assertEqual(submit_form._detect_already_submitted(driver,'2301000000000037','12'),(False,''))

    def test_navigation_is_not_submission_success(self):
        driver=Mock(current_url='https://jkbose.jk.gov.in/utilities/auth_pages/login.aspx')
        driver._exam_lookup_context=('2301000000000037','12')
        with patch.object(submit_form,'_ensure_driver_alive'),patch.object(submit_form,'_safe_element_value',return_value=None),patch.object(submit_form,'_safe_named_value',return_value=None),patch.object(submit_form,'_install_submit_click_marker'),patch.object(submit_form,'_manual_submit_clicked',return_value=True),patch.object(submit_form,'_detect_already_submitted',side_effect=ExamPortalUnavailable('Login expired')):
            with self.assertRaises(ExamPortalUnavailable):submit_form._wait_for_manual_submission(driver,'2301000000000037')

if __name__=='__main__':unittest.main()
