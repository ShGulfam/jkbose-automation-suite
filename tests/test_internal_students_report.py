import copy
import json
import unittest
from types import SimpleNamespace
from unittest.mock import Mock,patch
from openpyxl import load_workbook
import test_internal_batch as fixtures
from jkbose_tool.features import internal_assessment_browser as browser
from jkbose_tool.features import internal_assessment_batch as batch
from jkbose_tool.features import internal_assessment_students_report as report


class StudentsReportTests(unittest.TestCase):
    setUp=fixtures.ClassRunTests.setUp
    write_book=fixtures.ClassRunTests.write_book
    jobs=fixtures.ClassRunTests.jobs
    plan=fixtures.ClassRunTests.plan
    state=fixtures.ClassRunTests.state
    run_jobs=fixtures.ClassRunTests.run_jobs

    def snapshot(self):
        return {'account':'school-test','selected_class':'11','selected_session':'2026-2027','headers':list(report.HEADERS),
                'row_widths':[21],'has_pager':False,'data_table':False,
                'rows':[{'sequence':'1','exam_roll':'301000001','name':'Synthetic One','class_name':'XI','approval':'Approved',
                         'regno':'REG1','exam_fee':'Paid','subjects':['General English, Physics, Chemistry','Biology,Mathematics','','','']}],
                'workbook_session':'2025-26','report_session':'2026-2027','report_url':report.URL}

    def match(self,snapshot=None,prior=None,subject='Physics'):
        return report.match_student_report(self.plan(subject),snapshot or self.snapshot(),prior or {'verified':set(),'manual':set()})

    def test_report_session_is_separate_exact_and_never_inferred_for_another_school_year(self):
        options=['2024-2026','2025-2026','2026-2027']
        self.assertEqual(report.resolve_report_session('2025-26',options),'2026-2027')
        self.assertEqual(report.resolve_report_session('Annual Regular 2026',options),'2026-2027')
        with self.assertRaises(ValueError): report.resolve_report_session('2026-27',options)
        with self.assertRaises(ValueError): report.resolve_report_session('2025-26',options+['2027-2028'])
        with self.assertRaises(ValueError): report.resolve_report_session('2025-26',options[:-1])
        with self.assertRaises(ValueError): report.resolve_report_session('2025-26',options+['2026-2027'])
        self.assertEqual(report.resolve_report_session('2025-26',options,'2025-2026'),'2025-2026')

    def test_stale_class_wrong_account_changed_columns_and_partial_report_are_rejected(self):
        self.assertIsNotNone(report.validate_snapshot(self.snapshot(),'11','school-test','2026-2027'))
        variants=[]
        stale=self.snapshot(); stale['rows'][0]['class_name']='XII'; variants.append(stale)
        for key,value in [('account','another-school'),('selected_class','12'),('selected_session','2025-2026'),
                          ('headers',report.HEADERS[:-1]),('row_widths',[20]),('has_pager',True),('data_table',True),('rows',[])]:
            changed=self.snapshot(); changed[key]=value; variants.append(changed)
        numbering=self.snapshot(); numbering['rows'][0]['sequence']='2'; variants.append(numbering)
        for value in variants:
            with self.assertRaises(ValueError): report.validate_snapshot(value,'11','school-test','2026-2027')

    def test_subjects_come_from_exact_report_tokens_not_a_science_stream_assumption(self):
        row=self.snapshot()['rows'][0]
        for subject in ['Biology','Botany','Zoology','Mathematics','Physics']: self.assertTrue(report.enrolled_in_subject(row,subject))
        row['subjects']=['General English,Physics,Chemistry','Mathematics','','','']
        self.assertFalse(report.enrolled_in_subject(row,'Biology'))
        self.assertFalse(report.enrolled_in_subject(row,'Botany'))
        row['subjects'][1]='Botany;Zoology'
        self.assertTrue(report.enrolled_in_subject(row,'Biology'))
        row['subjects'][1]='Botany'
        self.assertFalse(report.enrolled_in_subject(row,'Biology'))

    def test_complete_report_reconciles_enrolment_without_inventing_attendance_or_submission(self):
        result=self.match()
        self.assertTrue(result['coverage_complete']); self.assertTrue(result['identity_roster_reconciled'])
        record=result['records'][0]
        self.assertEqual(record['match_status'],'matched_roster')
        self.assertEqual(record['jkbose_attendance'],''); self.assertEqual(record['official_maximum'],'')
        self.assertEqual(record['award_action'],'roster_check_only')
        self.assertIn('incomplete',result['coverage_reason'])

    def test_saved_pair_normally_remains_in_enrolment_and_missing_pair_is_not_saved_proof(self):
        result=self.match(prior={'verified':{'REG1'},'manual':set()})
        self.assertEqual(result['records'][0]['match_status'],'matched_roster')
        empty=self.snapshot(); empty['rows'][0]['subjects']=['Urdu','','','','']
        result=self.match(empty,{'verified':{'REG1'},'manual':set()})
        self.assertEqual(result['records'][0]['match_status'],'excel_only')
        self.assertFalse(result['identity_roster_reconciled'])

    def test_board_only_paired_identity_name_and_fee_exceptions_remain_visible(self):
        changed=self.snapshot(); changed['rows'][0]['exam_roll']='301000009'
        self.assertEqual(self.match(changed)['records'][0]['match_status'],'identifier_mismatch')
        changed=self.snapshot(); changed['rows'][0]['name']='Another Name'
        self.assertEqual(self.match(changed)['records'][0]['match_status'],'name_mismatch')
        changed=self.snapshot(); extra=copy.deepcopy(changed['rows'][0]); extra.update(sequence='2',regno='REG2',exam_roll='301000002')
        changed['rows'].append(extra)
        self.assertEqual(self.match(changed)['counts']['jkbose_only'],1)
        changed=self.snapshot(); changed['rows'][0].update(approval='Pending',exam_fee='Unpaid')
        record=self.match(changed)['records'][0]
        self.assertEqual(record['award_action'],'manual_only'); self.assertEqual(len(record['manual_issues']),2)

    def test_application_without_exam_roll_is_skipped_without_deduplicating_by_name(self):
        snapshot=self.snapshot(); blank=copy.deepcopy(snapshot['rows'][0]); blank.update(sequence='2',exam_roll='')
        snapshot['rows'].append(blank)
        result=self.match(snapshot)
        self.assertEqual(result['counts'],{'matched_roster':1}); self.assertTrue(result['identity_roster_reconciled'])
        # A second assigned roll remains visible even for the same name and registration.
        snapshot['rows'][1]['exam_roll']='301000002'
        self.assertFalse(self.match(snapshot)['identity_roster_reconciled'])
        snapshot['rows'][1]['exam_roll']='malformed'
        self.assertEqual(self.match(snapshot)['counts']['identifier_conflict'],2)

    def test_skipped_application_is_logged_in_export_but_never_becomes_an_absence_hold(self):
        snapshot=self.snapshot(); blank=copy.deepcopy(snapshot['rows'][0]); blank.update(sequence='2',exam_roll='')
        snapshot['rows'].append(blank); driver=Mock()
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=lambda d,p,a,**k:self.state(p)),patch.object(browser,'_evidence',return_value='synthetic'):
            result=self.run_jobs(self.jobs(subjects=['PH']),driver_factory=Mock(return_value=driver),login_callback=Mock(),report_session='Auto',report_reader=Mock(return_value=snapshot))
        metadata=result['student_report_matching']
        self.assertEqual(metadata['student_count'],1); self.assertEqual(metadata['raw_application_count'],2)
        self.assertEqual(metadata['skipped_no_exam_roll_count'],1); self.assertEqual(result['class_wide_holds'],[])
        book=load_workbook(metadata['exports']['excel']); ws=book['Skipped Applications']
        self.assertEqual(ws.max_row,2); self.assertEqual(ws['A2'].value,'REG1'); self.assertEqual(ws['A2'].data_type,'s'); book.close()

    def test_reader_always_searches_and_checks_returned_class_without_any_save_control(self):
        driver=Mock(); driver.current_url=report.URL
        driver.find_element.side_effect=lambda by,key:SimpleNamespace(id=key,click=Mock())
        classes=SimpleNamespace(options=[Mock()],select_by_value=Mock(),first_selected_option=Mock()); classes.options[0].get_attribute.return_value='11'; classes.first_selected_option.get_attribute.return_value='11'
        sessions=SimpleNamespace(options=[Mock()],select_by_value=Mock(),first_selected_option=Mock()); sessions.options[0].get_attribute.return_value='2026-2027'; sessions.first_selected_option.get_attribute.return_value='2026-2027'
        controls={}
        def element(by,key):
            controls.setdefault(key,SimpleNamespace(id=key,click=Mock(),get_attribute=Mock(return_value=''))); return controls[key]
        driver.find_element.side_effect=element; driver.execute_script.return_value=self.snapshot()
        with patch.object(report,'Select',side_effect=lambda node:classes if node.id=='bodypH_ddlclasssearch' else sessions),patch.object(browser,'_begin_read',return_value='before'),patch.object(browser,'_wait_page') as waited:
            result=report.read_student_report(driver,'11','school-test','2025-26')
        controls['bodypH_btnsearch'].click.assert_called_once(); waited.assert_called_once()
        classes.select_by_value.assert_called_once_with('11'); sessions.select_by_value.assert_called_once_with('2026-2027')
        self.assertEqual(result['workbook_session'],'2025-26'); self.assertEqual(result['report_session'],'2026-2027')
        self.assertNotIn('bodypH_btnSaveMarks',controls); driver.save_screenshot.assert_not_called()

    def test_class_postback_finishes_before_session_control_is_acquired(self):
        driver=Mock(); driver.current_url=report.URL; order=[]
        classes=SimpleNamespace(options=[Mock()],first_selected_option=Mock(),select_by_value=Mock(side_effect=lambda value:order.append('class_selected')))
        classes.options[0].get_attribute.return_value='11'; classes.first_selected_option.get_attribute.return_value='12'
        sessions=SimpleNamespace(options=[Mock()],first_selected_option=Mock(),select_by_value=Mock())
        sessions.options[0].get_attribute.return_value='2026-2027'; sessions.first_selected_option.get_attribute.return_value='2026-2027'
        def element(by,key):
            if key=='bodypH_ddlsessionsearch': self.assertEqual(order,['class_selected','postback_wait'])
            return SimpleNamespace(id=key,click=Mock(),get_attribute=Mock(return_value='__doPostBack' if key=='bodypH_ddlclasssearch' else ''))
        driver.find_element.side_effect=element; driver.execute_script.return_value=self.snapshot()
        with patch.object(report,'Select',side_effect=lambda node:classes if node.id=='bodypH_ddlclasssearch' else sessions),patch.object(browser,'_begin_read',return_value='before'),patch.object(browser,'_wait_page',side_effect=lambda *args:order.append('postback_wait')):
            report.read_student_report(driver,'11','school-test','2025-26')
        self.assertEqual(order,['class_selected','postback_wait','postback_wait'])

    def test_class_runner_reads_report_once_and_saves_distinct_roster_excel_pdf_and_sessions(self):
        jobs=self.jobs(subjects=['PH','MA']); driver=Mock(); reader=Mock(return_value=self.snapshot())
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=lambda d,p,a,**k:self.state(p)),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            result=self.run_jobs(jobs,driver_factory=Mock(return_value=driver),login_callback=Mock(),report_session='Auto',report_reader=reader)
        reader.assert_called_once(); fill.assert_not_called()
        metadata=result['student_report_matching']
        self.assertEqual(metadata['status'],'roster_reconciled'); self.assertEqual(metadata['report_session'],'2026-2027')
        self.assertEqual(metadata['workbook_session'],'2025-26'); self.assertEqual(result['manifest']['session'],'2025-26')
        folder=next((self.history/'References'/'_class_runs').iterdir())
        for path in [folder/'matching.xlsx',folder/'matching.pdf',folder/'students_report'/'matching.xlsx',folder/'students_report'/'matching.pdf']: self.assertTrue(path.exists())
        book=load_workbook(folder/'students_report'/'matching.xlsx')
        self.assertIn('2026-2027',book['Matching Summary']['B8'].value)
        self.assertEqual(book['Matching Summary']['E10'].value,'Complete report coverage')
        self.assertEqual(book['Paired Student Matching']['C2'].data_type,'s'); book.close()

    def test_report_matching_uses_assessment_only_for_interface_and_count(self):
        jobs=self.jobs(subjects=['MA']);driver=Mock()
        def state(d,plan,account,**kwargs):
            value=self.state(plan);value['total']=15
            value['available_subjects']=[{'value':'Mathematics','label':'Mathematics'}]
            return value
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=state),patch.object(browser,'_evidence',return_value='synthetic'),patch('jkbose_tool.features.internal_assessment_matching._read_only_next',side_effect=AssertionError('No assessment pagination in report matching')):
            result=self.run_jobs(jobs,driver_factory=Mock(return_value=driver),login_callback=Mock(),report_session='Auto',report_reader=Mock(return_value=self.snapshot()))
        self.assertEqual(result['student_report_matching']['status'],'roster_reconciled')
        checked=result['subjects'][0]
        self.assertEqual(checked['status'],'checked_read_only')
        self.assertEqual(checked['matching_summary']['expected_pending_count'],15)
        self.assertEqual(checked['matching_summary']['check_scope'],'interface_and_pending_count')
        self.assertFalse(checked['matching_summary']['coverage_complete'])
        self.assertEqual(len(checked['matching_summary']['available_subjects']),1)
        driver.find_element.assert_not_called()

    def test_report_submission_checks_only_selected_unfinished_subjects(self):
        jobs=self.jobs(subjects=['PH','MA'])
        fixtures.ClassRunTests.receipt(self,jobs[0]['plan'])
        seen=[]
        def open_subject(d,plan,stop):
            seen.append(plan.subject)
            if plan.subject!='Mathematics':raise ValueError('Completed/unselected subject no longer offered')
        with patch.object(browser,'open_subject',side_effect=open_subject),patch.object(browser,'snapshot',side_effect=lambda d,p,a,**k:self.state(p)),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            result=batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',execute=True,confirm_session=True,confirm_callback=Mock(return_value=False),driver_factory=Mock(return_value=Mock()),login_callback=Mock(),report_session='Auto',report_reader=Mock(return_value=self.snapshot()))
        self.assertEqual(seen,['Mathematics','Mathematics'])
        self.assertEqual(result['subjects'][0]['status'],'skipped_verified_history')
        self.assertEqual([row['subject'] for row in result['safety_preflight']],['Mathematics'])
        fill.assert_not_called()

    def test_selected_only_report_preflight_keeps_absence_in_unselected_source(self):
        self.write_book(physics='AB');jobs=self.jobs(subjects=['MA']);driver=Mock()
        with patch.object(browser,'open_subject') as opened,patch.object(browser,'snapshot',side_effect=lambda d,p,a,**k:self.state(p)),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            result=batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',execute=True,confirm_session=True,confirm_callback=Mock(return_value=False),driver_factory=Mock(return_value=driver),login_callback=Mock(),report_session='Auto',report_reader=Mock(return_value=self.snapshot()))
        self.assertTrue(all(call.args[1].subject=='Mathematics' for call in opened.call_args_list))
        self.assertEqual(result['subjects'][0]['status'],'cancelled')
        fill.assert_not_called();driver.find_element.assert_not_called()

    def test_report_selected_subject_missing_still_fails_before_saving(self):
        with patch.object(browser,'open_subject',side_effect=ValueError('The selected subject is missing')),patch.object(browser,'_fill') as fill:
            with self.assertRaisesRegex(ValueError,'selected subject is missing'):
                self.run_jobs(self.jobs(subjects=['MA']),execute=True,driver_factory=Mock(return_value=Mock()),login_callback=Mock(),report_session='Auto',report_reader=Mock(return_value=self.snapshot()))
        fill.assert_not_called()

    def test_full_enrolment_report_does_not_bypass_per_batch_confirmation(self):
        jobs=self.jobs(subjects=['MA']); driver=Mock(); driver.find_elements.return_value=[]
        def state(d,plan,account,**kwargs):
            value=self.state(plan); value['total']=2; return value
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=state),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            result=batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',execute=True,confirm_session=True,confirm_callback=Mock(return_value=False),driver_factory=Mock(return_value=driver),login_callback=Mock(),report_session='Auto',report_reader=Mock(return_value=self.snapshot()))
            self.assertEqual(result['status'],'cancelled')
        fill.assert_not_called(); driver.find_element.assert_not_called()

    def test_enrolment_mismatch_stops_before_practical_filling_or_submission(self):
        changed=self.snapshot(); changed['rows'][0]['exam_roll']='301000009'; driver=Mock()
        with patch.object(browser,'run') as runner:
            with self.assertRaisesRegex(ValueError,'enrolment/identity mismatch'):
                self.run_jobs(self.jobs(subjects=['PH']),execute=True,driver_factory=Mock(return_value=driver),login_callback=Mock(),report_session='Auto',report_reader=Mock(return_value=changed))
        runner.assert_not_called(); driver.find_element.assert_not_called()
        folder=next((self.history/'References'/'_class_runs').iterdir())
        self.assertTrue((folder/'students_report'/'matching.pdf').exists())

    def test_verified_completed_submission_skips_without_opening_report_browser(self):
        jobs=self.jobs(subjects=['PH'])
        fixtures.ClassRunTests.receipt(self,jobs[0]['plan'])
        factory=Mock(); reader=Mock()
        result=self.run_jobs(jobs,execute=True,driver_factory=factory,report_session='Auto',report_reader=reader)
        factory.assert_not_called(); reader.assert_not_called()
        self.assertEqual(result['subjects'][0]['status'],'skipped_verified_history')

if __name__=='__main__': unittest.main()


