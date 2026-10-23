import hashlib,json,unittest
from openpyxl import load_workbook
import test_internal_batch as fixtures
import test_internal_saved_report as saved_fixtures
from jkbose_tool.features.internal_assessment_tracking import tracking_payload,refresh_tracking_copy

class TrackingTests(unittest.TestCase):
    setUp=fixtures.ClassRunTests.setUp
    write_book=fixtures.ClassRunTests.write_book
    plan=fixtures.ClassRunTests.plan
    def report(self):
        folder=self.history/'_class_runs'/'synthetic';folder.mkdir(parents=True,exist_ok=True)
        data=saved_fixtures.SavedReportTests().report(self.plan('Mathematics'))
        (folder/'official_saved_assessment_report.json').write_text(json.dumps(data))
    def test_subject_and_student_flags_distinguish_requested_from_approved(self):
        self.report();payload=tracking_payload(self.file,self.history,'11','2025-26','school-test')
        maths=next(r for r in payload['summary'] if r[0]=='Mathematics')
        self.assertEqual(maths[1:],[1,1,1,0,0,0,0,0])
        self.assertEqual(payload['flags'][1][:3],['Submitted','Requested','Pending principal approval'])
    def test_class_wide_source_absence_blocks_other_unsaved_subjects(self):
        self.write_book(physics='AB');payload=tracking_payload(self.file,self.history,'11','2025-26','school-test')
        self.assertTrue(all(row[0]=='Pending' for row in payload['flags']))
        self.assertTrue(all(row[2]=='Manual Principal approval only' for row in payload['flags']))
    def test_export_copy_preserves_original_and_has_status_columns(self):
        self.report();before=self.file.read_bytes();output=refresh_tracking_copy(self.file,self.history,'11','2025-26','school-test')
        self.assertEqual(self.file.read_bytes(),before)
        original=load_workbook(self.file);copy=load_workbook(output)
        self.assertIn('Submission_History',copy.sheetnames)
        for row in original.active:
            for cell in row:self.assertEqual(copy[original.active.title].cell(cell.row,cell.column).value,cell.value)
        self.assertEqual(copy[original.active.title].cell(1,original.active.max_column+1).value,'Submission Status')
        original.close();copy.close()
    def test_rejected_awards_are_explicit_not_uncertain_or_submitted(self):
        folder=self.history/'_class_runs'/'rejected';folder.mkdir(parents=True)
        helper=saved_fixtures.SavedReportTests();data=helper.report(self.plan('Mathematics'))
        data['records'][0].update(status='Rejected',rejection_reason='test entry');data=helper.sign(data)
        (folder/'official_saved_assessment_report.json').write_text(json.dumps(data))
        payload=tracking_payload(self.file,self.history,'11','2025-26','school-test')
        self.assertEqual(payload['flags'][1][0],'Rejected')
        self.assertIn('test entry',payload['flags'][1][5])
        maths=next(r for r in payload['summary'] if r[0]=='Mathematics')
        self.assertEqual(maths[1:],[1,0,0,0,0,0,0,1])
