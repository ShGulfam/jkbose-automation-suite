import unittest,json
from openpyxl import load_workbook
import test_internal_batch as fixtures
from jkbose_tool.features.internal_assessment_policy import CohortAbsencePolicy

class SourceCorrectionTests(unittest.TestCase):
    setUp=fixtures.ClassRunTests.setUp
    write_book=fixtures.ClassRunTests.write_book
    def policy(self):return CohortAbsencePolicy(self.history,'school-test','11','2025-26')
    def hold(self,reg='REG1',exam='301000001',reason='Marked absent in consolidated source'):
        return dict(registration=reg,exam_roll=exam,subject='Physics',reason=reason)
    def test_corrected_present_handoff_releases_only_source_hold_and_preserves_chain(self):
        p=self.policy();p.add([self.hold()],origin='consolidated_source',evidence='old-source')
        before=p.path.read_bytes();self.assertEqual(len(p.reconcile_source(self.file)),1)
        self.assertTrue(p.path.read_bytes().startswith(before));self.assertEqual(self.policy().holds,[])
        self.assertEqual(p.reconcile_source(self.file),[])
    def test_current_absence_in_any_subject_keeps_classwide_hold(self):
        p=self.policy();p.add([self.hold()],origin='consolidated_source');self.write_book(maths='AB')
        self.assertEqual(p.reconcile_source(self.file),[]);self.assertEqual(len(self.policy().holds),1)
    def test_changed_identifier_or_missing_student_never_releases_hold(self):
        for reg,exam in [('REG2','301000001'),('REG1','301000009')]:
            p=self.policy();p.add([self.hold(reg,exam)],origin='consolidated_source')
        self.assertEqual(p.reconcile_source(self.file),[]);self.assertEqual(len(self.policy().holds),2)
    def test_live_origin_and_identity_conflict_are_never_automatically_released(self):
        p=self.policy();p.add([self.hold()],origin='consolidated_source')
        p.add([self.hold()],origin='live_jkbose_read_only')
        p.add([self.hold(reason='Registration has conflicting exam rolls across subjects')],origin='consolidated_source')
        self.assertEqual(p.reconcile_source(self.file),[]);self.assertEqual(len(self.policy().holds),2)
    def test_invalid_award_or_missing_explicit_attendance_does_not_clear(self):
        p=self.policy();p.add([self.hold()],origin='consolidated_source');self.write_book(maths=99)
        self.assertEqual(p.reconcile_source(self.file),[])
        self.write_book();book=load_workbook(self.file);book.active.cell(2,12,'');book.save(self.file);book.close()
        self.assertEqual(p.reconcile_source(self.file),[])
    def test_later_real_absence_can_add_a_new_hold_after_correction(self):
        p=self.policy();p.add([self.hold()],origin='consolidated_source');p.reconcile_source(self.file)
        p.add([self.hold()],origin='consolidated_source',evidence='new-absence')
        self.assertEqual(len(self.policy().holds),1)
