"""Synthetic explicit absence submission, approval exclusion and sidecar evidence."""
import copy,json,unittest,hashlib
from unittest.mock import Mock,patch
from openpyxl import load_workbook
import test_internal_batch as fixtures
from jkbose_tool.features import internal_assessment_browser as browser
from jkbose_tool.features import internal_assessment_batch as batch
from jkbose_tool.features.internal_assessment import load_plan,match_batch,submission_issues,Audit
from jkbose_tool.features.internal_assessment_policy import apply_holds,source_absences
from jkbose_tool.features.internal_assessment_saved_report import _award_matches
from jkbose_tool.features.internal_assessment_approval import match_requests,grid_row_matches
from jkbose_tool.features.internal_assessment_absence_register import refresh_absence_register
from jkbose_tool.features.internal_assessment_identity import WARNING

class AbsenceSubmissionTests(unittest.TestCase):
    setUp=fixtures.ClassRunTests.setUp
    write_book=fixtures.ClassRunTests.write_book
    plan=fixtures.ClassRunTests.plan
    state=fixtures.ClassRunTests.state
    def absent(self,second=False):
        self.write_book(physics='AB',second=second)
        if second:
            b=load_workbook(self.file);w=b.active;w.cell(3,10,8);w.cell(3,12,'Present');b.save(self.file);b.close()
        return self.plan()
    def live(self,p):
        d=self.state(p)
        for i,r in enumerate(d['rows']):r.update(absent_input_id=f'absent_{i}',absent_enabled=True)
        return d
    def test_explicit_absent_import_has_no_fake_marks_and_is_submission_only(self):
        p=self.absent();a=p.awards[0];self.assertEqual((a.attendance,a.marks,a.components),('A','',{'Physics':'AB'}))
        self.assertFalse(submission_issues(a));d=self.live(p)
        self.assertFalse(match_batch(p,d,allow_absent=True)[1]);self.assertTrue(match_batch(p,d)[1])
    def test_missing_unknown_conflicting_and_below_pass_never_become_absent(self):
        for marks,att in [('', 'Present'),(1,'Present'),('AB','Present'),('AB','Unknown'),(8,'Absent')]:
            self.write_book(physics=marks);b=load_workbook(self.file);b.active.cell(2,12,att);b.save(self.file);b.close();p=self.plan()
            self.assertTrue(match_batch(p,self.live(p),allow_absent=True)[1],(marks,att))
    def test_absent_requires_exact_identity_and_verified_radio_and_no_existing_mark(self):
        p=self.absent();original=self.live(p)
        for field,value in [('exam_roll','301000009'),('name','Different Student'),('absent_input_id',''),('absent_enabled',False),('value','8')]:
            d=copy.deepcopy(original);d['rows'][0][field]=value;self.assertTrue(match_batch(p,d,allow_absent=True)[1],field)
    def test_present_source_live_absent_is_conflict_never_changed_to_present(self):
        p=self.plan();d=self.live(p);d['rows'][0]['attendance']='A'
        self.assertTrue(match_batch(p,d,allow_absent=True)[1])
    def test_cross_subject_absence_excludes_principal_but_permits_valid_submission(self):
        self.absent();p=self.plan('Mathematics');apply_holds(p,source_absences(self.file));d=self.live(p)
        self.assertFalse(match_batch(p,d,allow_absent=True)[1]);self.assertTrue(match_batch(p,d)[1])
        r=dict(class_name='11',regno='REG1',exam_roll='301000001',name='Synthetic One',subject='Mathematics',attendance='P',maximum='20',marks='12',id='1',requested_on='date')
        self.assertFalse(match_requests([p],[r])[0]['eligible'])
    def test_identity_hold_still_blocks_submission(self):
        p=self.plan();apply_holds(p,[dict(registration='REG1',exam_roll='301000001',subject='identity review',reason='Conflicting registration')])
        self.assertTrue(match_batch(p,self.live(p),allow_absent=True)[1])
    def test_absent_official_receipts_require_absent_attendance_and_no_positive_marks(self):
        p=self.absent();a=p.awards[0];r=dict(name=a.name,attendance='A',status='Requested',rejection_reason='',marks='0',maximum='10')
        self.assertTrue(_award_matches(p,a,r));self.assertFalse(_award_matches(p,a,dict(r,attendance='P')));self.assertFalse(_award_matches(p,a,dict(r,marks='8')))
    def test_empty_and_zero_absent_grid_marks_can_be_matched_without_numeric_error(self):
        r=dict(class_name='11',regno='REG1',exam_roll='301000001',name='Synthetic One',subject='Physics',attendance='A',maximum='10',marks='AB',requested_on='date')
        self.assertTrue(grid_row_matches(dict(cells=['11','301000001','REG1','Synthetic One','Physics','AB','10','date']),r))
    def test_mixed_absent_present_batches_submit_and_resume_both_subjects(self):
        self.absent(second=True);jobs=batch.load_jobs(self.file,student_class='11',session='2025-26',subjects=['PH','MA']);original=self.file.read_bytes()
        states={};active={};saved=[];driver=Mock();driver.switch_to=Mock()
        type(driver.switch_to).alert=property(lambda _:(_ for _ in ()).throw(browser.NoAlertPresentException()))
        def open_subject(d,p,stop):
            active['subject']=p.subject
            if p.subject not in states:states[p.subject]=self.live(p)
        def snapshot(d,p,a,**k):return copy.deepcopy(states[p.subject])
        def find(by,identity):
            element=Mock();state=states[active['subject']]
            if identity=='bodypH_btnSaveMarks':
                def save():saved.append(copy.deepcopy(state));state.update(total=0,rows=[])
                element.click.side_effect=save;return element
            for i,row in enumerate(state['rows']):
                if identity==row['input_id']:
                    element.get_attribute.side_effect=lambda key,r=row:r['value'] if key=='value' else ''
                    element.clear.side_effect=lambda r=row:r.update(value='')
                    element.send_keys.side_effect=lambda v,r=row:r.update(value=v)
                    return element
                if identity==row['absent_input_id']:
                    element.is_enabled.return_value=True;element.get_attribute.return_value='A'
                    element.click.side_effect=lambda r=row:r.update(attendance='A',disabled=True)
                    return element
            raise AssertionError(identity)
        driver.find_element.side_effect=find
        with patch.object(browser,'open_subject',side_effect=open_subject),patch.object(browser,'snapshot',side_effect=snapshot),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'show_students'):
            result=batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',execute=True,confirm_session=True,confirm_callback=Mock(return_value=True),driver_factory=Mock(return_value=driver),login_callback=Mock())
            self.assertEqual(result['new_verified_saved_count'],4);self.assertEqual(len(saved),2)
            self.assertEqual([(r['attendance'],r['value']) for r in saved[0]['rows']],[('A',''),('P','8')])
            self.assertTrue(all(r['attendance']=='P' for r in saved[1]['rows']))
            factory=Mock();batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',execute=True,confirm_session=True,confirm_callback=Mock(),driver_factory=factory)
            factory.assert_not_called()
        self.assertEqual(self.file.read_bytes(),original)
        output=refresh_absence_register(self.file,self.history,'11','2025-26','school-test');b=load_workbook(output);r=list(b.active.values);b.close()
        self.assertEqual(len(r),6);self.assertEqual(r[-1][9],'Submitted as Absent - verified');self.assertIn('NEVER',r[-1][10])
    def test_uncertain_attempt_register_does_not_claim_submitted(self):
        p=self.absent();a=Audit(self.history,p);a.append('run_started',school_account='school-test',source=p.manifest(),source_preview=p.rows());a.append('submit_attempt',batch_id='one',registrations=['REG1'])
        output=refresh_absence_register(self.file,self.history,'11','2025-26','school-test');b=load_workbook(output);self.assertIn('Uncertain',b.active.cell(6,10).value);b.close()
    def test_present_plan_keeps_legacy_digest(self):
        p=self.plan();self.assertNotIn('attendance',p.payload()['awards'][0]);self.assertEqual(p.digest(),hashlib.sha256(json.dumps(p.payload(),sort_keys=True).encode()).hexdigest())
    def test_exact_user_confirmation_clears_only_specific_warning(self):
        b=load_workbook(self.file);b.active.cell(2,13,WARNING.removeprefix('Portal review required: '));b.save(self.file);b.close()
        record=dict(schema='jkbose-identity-confirmations-v1',confirmations=[dict(**{'class':'11'},session='2025-26',registration='REG1',exam_roll='301000001',name='Synthetic One',warning=WARNING,confirmed_by='user',evidence='Synthetic administrator and official paired-identity verification')])
        side=self.file.with_name('JKBOSE_Identity_Confirmations.json');side.write_text(json.dumps(record));self.assertFalse(self.plan().awards[0].issues)
        record['confirmations'][0]['exam_roll']='301000009';side.write_text(json.dumps(record));self.assertTrue(self.plan().awards[0].issues)

    def test_absent_source_identity_warning_is_never_filtered(self):
        self.absent();b=load_workbook(self.file);b.active.cell(2,13,'Registration conflict; admin review.');b.save(self.file);b.close();p=self.plan()
        self.assertTrue(match_batch(p,self.live(p),allow_absent=True)[1])
    def test_readback_attendance_change_prevents_save(self):
        p=self.absent();before=self.live(p);after=copy.deepcopy(before);driver=Mock();driver.find_element.return_value.get_attribute.return_value='A';driver.find_element.return_value.is_enabled.return_value=True
        with patch.object(browser,'snapshot',return_value=after):
            with self.assertRaises(ValueError):browser._fill(driver,p,before,p.awards,'school-test',None)
        self.assertTrue(all(call.args[1]!='bodypH_btnSaveMarks' for call in driver.find_element.call_args_list))

if __name__=='__main__':unittest.main()
