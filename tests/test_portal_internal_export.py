"""Offline portal handoff contract tests; synthetic files and no browser/network."""
import sys,unittest,tempfile
from pathlib import Path
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from jkbose_tool.features.internal_assessment import load_plan

class PortalInternalExportTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
  self.file=Path(self.tmp.name)/'portal.xlsx'
 def workbook(self,issue='',mark='4.25',maximum='10',schema='hss-jkbose-ia-v1',exam='301000001'):
  b=Workbook();b.active.title='Awards_Matrix';b.active.append(['Not an automation sheet'])
  w=b.create_sheet('JKBOSE_Internal_Assessment')
  w.append(['Class','Session','Evaluation Type','Subject Code','Subject Name','Board Reg. No.','Exam Roll No.','Class Roll No.','Student Name','Marks','Max Marks','Attendance','Source Issues','Export Schema','Source Award ID'])
  w.append(['12th','2025-26','Internal Assessment','PH','Physics','0000000000000001',exam,'1','Synthetic Student',mark,maximum,'Absent' if mark=='AB' else 'Present',issue,schema,'synthetic-award'])
  b.save(self.file);b.close()
 def plan(self,**kw):return load_plan(self.file,student_class='12',session='2025-26',subject='Physics',maximum='10',**kw)
 def test_auto_prefers_values_handoff_and_preserves_exact_decimal_and_identifier(self):
  self.workbook();before=self.file.read_bytes();p=self.plan()
  self.assertEqual(p.sheet,'JKBOSE_Internal_Assessment');self.assertEqual(p.awards[0].regno,'0000000000000001');self.assertEqual(p.awards[0].marks,'4.25');self.assertFalse(p.awards[0].issues);self.assertEqual(before,self.file.read_bytes())
 def test_source_conflicts_and_missing_or_wrong_maximum_are_reported(self):
  for kwargs in [{'issue':'Conflicting approved awards.'},{'maximum':''},{'maximum':'5'},{'exam':''}]:
   with self.subTest(kwargs=kwargs):self.workbook(**kwargs);self.assertTrue(self.plan().awards[0].issues)
 def test_absent_and_below_pass_are_blocked(self):
  for mark in ['AB','3.99','']:
   with self.subTest(mark=mark):self.workbook(mark=mark);self.assertTrue(self.plan().awards[0].issues)
 def test_unknown_schema_is_rejected(self):
  self.workbook(schema='unknown')
  with self.assertRaises(ValueError):self.plan()
 def test_explicit_worksheet_is_honored_and_wrong_sheet_does_not_fall_back(self):
  self.workbook()
  with self.assertRaises(ValueError):self.plan(sheet='Awards_Matrix')
 def test_removed_safety_columns_are_rejected(self):
  self.workbook()
  from openpyxl import load_workbook
  b=load_workbook(self.file);b['JKBOSE_Internal_Assessment'].delete_cols(13,2);b.save(self.file);b.close()
  with self.assertRaises(ValueError):self.plan()

if __name__=='__main__':unittest.main()
