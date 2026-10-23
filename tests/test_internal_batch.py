"""Synthetic multi-subject, verified-resume and paired-identity safety tests."""
import copy
import hashlib
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import Mock,patch
from openpyxl import Workbook,load_workbook
from jkbose_tool.features.internal_assessment import Audit,load_plan
from jkbose_tool.features import internal_assessment_browser as browser
from jkbose_tool.features import internal_assessment_batch as batch


class ClassRunTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name).resolve(); self.file=self.root/'awards.xlsx'; self.history=self.root/'history'
        self.write_book()

    def write_book(self,physics=8,maths=12,exam='301000001',second=False):
        book=Workbook(); ws=book.active; ws.title='JKBOSE_Internal_Assessment'
        ws.append(['Class','Session','Evaluation Type','Subject Code','Subject Name','Board Reg. No.','Exam Roll No.','Class Roll No.','Student Name','Marks','Max Marks','Attendance','Source Issues','Export Schema','Source Award ID'])
        for code,label,mark,maximum in [('PH','Physics',physics,10),('MA','Mathematics',maths,20),('BO','Botany',3,5),('ZO','Zoology',4,5)]:
            ws.append(['11th','2025-26','Internal Assessment',code,label,'REG1',exam,'1','Synthetic One',mark,maximum,'Absent' if mark=='AB' else 'Present','','hss-jkbose-ia-v1','synthetic'])
            if second: ws.append(['11th','2025-26','Internal Assessment',code,label,'REG2','301000002','2','Synthetic Two',mark,maximum,'Absent' if mark=='AB' else 'Present','','hss-jkbose-ia-v1','synthetic'])
        book.save(self.file); book.close()

    def jobs(self,**kwargs):
        return batch.load_jobs(self.file,student_class='11',session='2025-26',**kwargs)

    def plan(self,subject='Physics'):
        return load_plan(self.file,student_class='11',session='2025-26',subject=subject)

    def receipt(self,plan,regs=None,manual=False,account='school-test'):
        regs=regs or [award.regno for award in plan.awards]
        audit=Audit(self.history,plan)
        audit.append('run_started',school_account=account,source=plan.manifest(),source_preview=plan.rows())
        audit.append('submit_attempt',batch_id='synthetic-batch',registrations=regs)
        if manual:
            audit.reconcile('synthetic-batch','saved_manually_verified','Synthetic Admin','Checked the entire synthetic saved award batch.')
        else:
            audit.append('verified_saved',batch_id='synthetic-batch',registrations=regs)
        return audit

    def run_jobs(self,jobs,**kwargs):
        return batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',confirm_session=True,confirm_callback=Mock(return_value=True),**kwargs)

    def state(self,plan):
        return {'class':'11','subject':plan.subject,'total':len(plan.awards),'rows':[
            {'regno':a.regno,'name':a.name,'exam_roll':a.exam_roll,'subject':plan.subject,'attendance':'P','max':plan.maximum,
             'min':'0','step':'any','value':'','input_id':str(i),'disabled':False,'readonly':False}
            for i,a in enumerate(plan.awards)]}

    def test_all_subjects_discovers_workbook_and_combines_biology_once(self):
        jobs=self.jobs(all_subjects=True)
        self.assertEqual([job['subject'] for job in jobs],['Physics','Biology','Mathematics'])
        self.assertTrue(all(job['plan'] for job in jobs))
        self.assertEqual(self.jobs(subjects=['BO','ZO','BI','PH'])[0]['plan'].awards[0].marks,'7')
        self.assertEqual(len(self.jobs(subjects=['BO','ZO','BI','PH'])),2)

    def test_explicit_multi_selection_and_different_maxima(self):
        jobs=self.jobs(subjects=['MA','PH'])
        self.assertEqual([job['plan'].maximum for job in jobs],['20','10'])
        with self.assertRaises(ValueError): self.jobs(subjects=['MA','PH'],maximum='20')
        self.assertEqual(self.jobs(subjects=['Urdu'])[0]['plan'],None)

    def test_mixed_class_session_or_non_internal_workbook_rejected(self):
        for column,value in [(1,'12th'),(2,'2026-27'),(3,'External Practical')]:
            self.write_book(); book=load_workbook(self.file); book.active.cell(2,column,value); book.save(self.file); book.close()
            with self.assertRaises(ValueError): self.jobs(all_subjects=True)

    def test_one_subject_source_problem_keeps_other_jobs_reviewable(self):
        book=load_workbook(self.file); book.active.cell(2,11,''); book.save(self.file); book.close()
        jobs=self.jobs(all_subjects=True)
        self.assertIsNone(jobs[0]['plan']); self.assertTrue(jobs[0]['error'])
        self.assertTrue(jobs[1]['plan']); self.assertTrue(jobs[2]['plan'])

    def test_one_login_for_multiple_subjects_and_manual_batch_continues(self):
        jobs=self.jobs(subjects=['PH','MA']); driver=Mock(); factory=Mock(return_value=driver); login=Mock()
        def run(driver,plan,*args,**kwargs):
            self.assertTrue((self.history/'References'/'internal-assessment.lock').exists())
            self.assertTrue(kwargs['lock_held'])
            return {'subject':plan.subject,'status':'manual_required' if plan.subject=='Physics' else 'checked_read_only','new_verified_count':0}
        with patch.object(browser,'run',side_effect=run) as runner:
            report=self.run_jobs(jobs,driver_factory=factory,login_callback=login)
        factory.assert_called_once(); login.assert_called_once_with(driver); self.assertEqual(runner.call_count,2); driver.quit.assert_called_once()
        self.assertEqual(report['status'],'finished_with_manual_review')
        self.assertEqual(report['new_verified_saved_count'],0)
        self.assertFalse((self.history/'References'/'internal-assessment.lock').exists())

    def test_all_verified_subjects_skip_without_browser(self):
        jobs=self.jobs(subjects=['PH','MA'])
        for job in jobs: self.receipt(job['plan'])
        factory=Mock()
        report=self.run_jobs(jobs,execute=True,driver_factory=factory)
        factory.assert_not_called()
        self.assertTrue(all(result['status']=='skipped_verified_history' for result in report['subjects']))
        self.assertEqual(report['new_verified_saved_count'],0)

    def test_reexported_identical_awards_skip_but_changed_saved_marks_stop(self):
        old=self.plan(); self.receipt(old)
        book=load_workbook(self.file); book.create_sheet('Extra export metadata'); book.save(self.file); book.close()
        new=self.plan(); self.assertNotEqual(new.source_sha256,old.source_sha256)
        prior=batch.verified_history(self.history,new,'school-test')
        self.assertEqual(prior['verified'],{'REG1'})
        browser.skip_verified_subject(new,self.history,'school-test',prior)
        report=json.loads((Audit(self.history,new).folder/'report.json').read_text())
        self.assertEqual(report['records'][0]['status'],'previously_verified_saved')
        self.write_book(physics=9)
        factory=Mock()
        with self.assertRaises(ValueError): self.run_jobs(self.jobs(subjects=['MA','PH']),driver_factory=factory)
        factory.assert_not_called()

    def test_saved_history_never_crosses_school_class_session_subject_or_exam_roll(self):
        old=self.plan(); self.receipt(old)
        with self.assertRaises(ValueError): batch.verified_history(self.history,old,'another-school')
        other=copy.deepcopy(old); other.student_class='12'
        self.assertFalse(batch.verified_history(self.history,other,'school-test')['verified'])
        other=copy.deepcopy(old); other.session='2026-27'
        self.assertFalse(batch.verified_history(self.history,other,'school-test')['verified'])
        self.assertFalse(batch.verified_history(self.history,self.plan('Mathematics'),'school-test')['verified'])
        self.write_book(exam='301999999')
        with self.assertRaises(ValueError): batch.verified_history(self.history,self.plan(),'school-test')

    def test_manual_saved_receipt_is_distinct_and_uncertain_save_stops_all_preflight(self):
        plan=self.plan(); self.receipt(plan,manual=True)
        prior=batch.verified_history(self.history,plan,'school-test')
        self.assertEqual(prior['manual'],{'REG1'}); self.assertFalse(prior['verified'])
        maths=self.plan('Mathematics'); audit=Audit(self.history,maths)
        audit.append('run_started',school_account='school-test',source_preview=maths.rows())
        audit.append('submit_attempt',batch_id='uncertain',registrations=['REG1'])
        factory=Mock()
        with self.assertRaises(ValueError): self.run_jobs(self.jobs(subjects=['PH','MA']),driver_factory=factory)
        factory.assert_not_called()

    def test_empty_queue_or_readonly_history_is_not_saved_proof(self):
        jobs=self.jobs(subjects=['PH']); plan=jobs[0]['plan']; audit=Audit(self.history,plan)
        audit.append('run_started',school_account='school-test',source_preview=plan.rows())
        audit.append('pending_queue_empty',verified_saved_count=0)
        self.assertFalse(batch.verified_history(self.history,plan,'school-test')['verified'])
        factory=Mock(return_value=Mock())
        with patch.object(browser,'run',return_value={'subject':'Physics','status':'pending_empty','pending_count':0}):
            report=self.run_jobs(jobs,driver_factory=factory,login_callback=Mock())
        factory.assert_called_once(); self.assertEqual(report['subjects'][0]['status'],'pending_empty')

    def test_partial_resume_reappearing_saved_pair_blocks_before_fill(self):
        self.write_book(second=True); plan=self.plan(); self.receipt(plan,['REG1'])
        live=self.state(plan); live['rows']=live['rows'][:1]; live['total']=1
        driver=Mock()
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',return_value=live),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            result=browser.run(driver,plan,self.history,'school-test',Mock(return_value=True),submit=True,confirm_session=True)
        self.assertEqual(result['status'],'manual_required'); fill.assert_not_called(); driver.find_element.assert_not_called()

    def test_confirmation_decline_stops_remaining_subjects(self):
        jobs=self.jobs(subjects=['PH','MA']); driver=Mock()
        with patch.object(browser,'run',return_value={'subject':'Physics','status':'cancelled'}) as runner:
            report=self.run_jobs(jobs,driver_factory=Mock(return_value=driver),login_callback=Mock())
        self.assertEqual(runner.call_count,1); self.assertEqual(report['status'],'cancelled')
        self.assertEqual(report['subjects'][1]['status'],'not_started')

    def test_stop_tampered_history_and_changed_source_never_open_browser(self):
        jobs=self.jobs(subjects=['PH','MA']); stop=threading.Event(); stop.set(); factory=Mock()
        with self.assertRaises(browser.AssessmentStopped): self.run_jobs(jobs,stop_event=stop,driver_factory=factory)
        factory.assert_not_called()
        audit=self.receipt(jobs[0]['plan']); text=audit.path.read_text(); audit.path.write_text(text.replace('school-test','tampered'))
        with self.assertRaises(ValueError): self.run_jobs(jobs,driver_factory=factory)
        factory.assert_not_called()
        self.write_book(physics=9)
        with self.assertRaises(ValueError): self.run_jobs(jobs,driver_factory=factory)
        factory.assert_not_called()

    def test_group_summary_and_journal_chain_are_durable_and_offline(self):
        jobs=self.jobs(all_subjects=True); factory=Mock()
        report=self.run_jobs(jobs,preview_only=True,driver_factory=factory)
        factory.assert_not_called(); self.assertEqual(len(report['subjects']),3)
        folder=next((self.history/'References'/'_class_runs').iterdir())
        self.assertTrue((folder/'report.csv').is_file()); self.assertTrue((folder/'manifest.json').is_file())
        previous='0'*64
        for line in (folder/'events.jsonl').read_text().splitlines():
            event=json.loads(line); signature=event.pop('hash')
            self.assertEqual(event['previous'],previous); self.assertEqual(Audit._hash(event),signature); previous=signature

    def test_both_identifiers_mandatory_and_duplicate_or_swapped_exam_rolls_block(self):
        for exam in ['', '123', 'not-an-exam-roll']:
            self.write_book(exam=exam); plan=self.plan()
            self.assertTrue(plan.awards[0].issues)
            self.assertTrue(browser.match_batch(plan,self.state(plan))[1])
        self.write_book(second=True); plan=self.plan(); live=self.state(plan)
        live['rows'][0]['exam_roll'],live['rows'][1]['exam_roll']=live['rows'][1]['exam_roll'],live['rows'][0]['exam_roll']
        self.assertTrue(browser.match_batch(plan,live)[1])
        live=self.state(plan); live['rows'][1]['exam_roll']=live['rows'][0]['exam_roll']
        self.assertTrue(browser.match_batch(plan,live)[1])
        book=load_workbook(self.file); book.active.cell(3,7,'301000001'); book.save(self.file); book.close()
        self.assertTrue(self.plan().awards[0].issues); self.assertTrue(self.plan().awards[1].issues)

    def test_real_adapter_readonly_class_loop_holds_one_lock_and_never_fills(self):
        jobs=self.jobs(subjects=['PH','MA']); driver=Mock()
        with patch.object(browser,'open_subject') as open_subject,patch.object(browser,'snapshot',side_effect=lambda d,p,a,**k:self.state(p)),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            report=self.run_jobs(jobs,driver_factory=Mock(return_value=driver),login_callback=Mock())
        self.assertEqual(open_subject.call_count,2); fill.assert_not_called(); driver.find_element.assert_not_called()
        self.assertTrue(all(result['status']=='checked_read_only' for result in report['subjects']))

    def test_real_adapter_class_save_and_resume_filters_pending_rows_once(self):
        jobs=self.jobs(subjects=['PH','MA']); driver=Mock(); pending={'Physics':True,'Mathematics':True,'Biology':True}; active={}
        driver.switch_to=Mock()
        type(driver.switch_to).alert=property(lambda _: (_ for _ in ()).throw(browser.NoAlertPresentException()))
        def open_subject(d,plan,stop): active['subject']=plan.subject
        def snapshot(d,plan,account,**kwargs):
            state=self.state(plan)
            if not pending[plan.subject]: state['rows']=[]; state['total']=0
            return state
        driver.find_element.return_value.click.side_effect=lambda:pending.update({active['subject']:False})
        with patch.object(browser,'open_subject',side_effect=open_subject),patch.object(browser,'snapshot',side_effect=snapshot),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill',side_effect=lambda d,p,b,*args:b),patch.object(browser,'show_students'):
            report=self.run_jobs(jobs,execute=True,driver_factory=Mock(return_value=driver),login_callback=Mock())
        self.assertEqual(report['new_verified_saved_count'],2)
        self.assertEqual(driver.find_element.return_value.click.call_count,2)
        factory=Mock(); resumed=self.run_jobs(jobs,execute=True,driver_factory=factory)
        factory.assert_not_called(); self.assertEqual(resumed['new_verified_saved_count'],0)
        self.assertTrue(all(result['status']=='skipped_verified_history' for result in resumed['subjects']))

    def test_unknown_absent_control_blocks_only_its_subject_other_subject_submits(self):
        self.write_book(physics='AB'); jobs=self.jobs(subjects=['PH','MA']); driver=Mock(); pending={'Physics':True,'Mathematics':True,'Biology':True}; active={}
        driver.switch_to=Mock()
        type(driver.switch_to).alert=property(lambda _: (_ for _ in ()).throw(browser.NoAlertPresentException()))
        def open_subject(d,plan,stop): active['subject']=plan.subject
        def snapshot(d,plan,account,**kwargs):
            state=self.state(plan)
            if not pending[plan.subject]: state['rows']=[]; state['total']=0
            return state
        driver.find_element.return_value.click.side_effect=lambda:pending.update({active['subject']:False})
        with patch.object(browser,'open_subject',side_effect=open_subject),patch.object(browser,'snapshot',side_effect=snapshot),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill',side_effect=lambda d,p,b,*args:b) as fill,patch.object(browser,'show_students'):
            report=self.run_jobs(jobs,execute=True,driver_factory=Mock(return_value=driver),login_callback=Mock())
        self.assertEqual(report['subjects'][0]['status'],'manual_required')
        self.assertTrue(pending['Physics']); self.assertFalse(pending['Mathematics'])
        self.assertEqual(fill.call_count,1); self.assertEqual(report['new_verified_saved_count'],1)
        driver.find_element.return_value.click.assert_called_once()

if __name__=='__main__': unittest.main()

