"""Safety and reconciliation tests use fabricated students, never live submissions."""
import json
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock,patch
from openpyxl import load_workbook
import test_internal_batch as fixtures
from jkbose_tool.features import internal_assessment_batch as batch
from jkbose_tool.features import internal_assessment_browser as browser
from jkbose_tool.features import internal_assessment_matching as matching
from jkbose_tool.features.internal_assessment_policy import source_absences,apply_holds,CohortAbsencePolicy


class MatchingTests(unittest.TestCase):
    setUp=fixtures.ClassRunTests.setUp
    write_book=fixtures.ClassRunTests.write_book
    jobs=fixtures.ClassRunTests.jobs
    plan=fixtures.ClassRunTests.plan
    state=fixtures.ClassRunTests.state
    run_jobs=fixtures.ClassRunTests.run_jobs

    def compare(self,rows,complete=True,prior=None):
        return matching.compare_rosters(self.plan(),rows,complete=complete,official_maximum=Decimal('10'),
            prior=prior or {'verified':set(),'manual':set()},coverage_reason='Synthetic coverage')

    def test_unseen_rows_never_become_saved_or_missing_by_assumption(self):
        self.assertEqual(self.compare([],False)['records'][0]['match_status'],'not_seen_unread_pages')
        result=self.compare([])
        self.assertEqual(result['records'][0]['match_status'],'not_in_pending_unverified')
        self.assertFalse(result['identity_roster_reconciled'])
        result=self.compare([],prior={'verified':{'REG1'},'manual':set()})
        self.assertEqual(result['records'][0]['match_status'],'verified_submitted')
        self.assertTrue(result['identity_roster_reconciled'])
        unknown=matching.compare_rosters(self.plan(),[],complete=True,official_maximum=None,
            prior={'verified':set(),'manual':set()},coverage_reason='Empty pending queue')
        self.assertEqual(unknown['official_maximum'],'')

    def test_rosters_require_paired_ids_and_detect_board_only_and_names(self):
        rows=self.state(self.plan())['rows']
        self.assertTrue(self.compare(rows)['identity_roster_reconciled'])
        rows[0]['exam_roll']='301999999'
        self.assertEqual(self.compare(rows)['records'][0]['match_status'],'identifier_mismatch')
        rows[0]['regno']='OTHER'; rows[0]['exam_roll']='301888888'
        result=self.compare(rows)
        self.assertEqual(result['counts']['jkbose_only'],1)
        self.assertEqual(result['counts']['not_in_pending_unverified'],1)
        rows=self.state(self.plan())['rows']; rows[0]['name']='Different Synthetic Name'
        self.assertEqual(self.compare(rows)['records'][0]['match_status'],'name_mismatch')

    def test_live_max_fail_absent_and_over_max_are_manual_only(self):
        for marks,attendance,fragment in [(11,'P','Exceeds'),(3,'P','Below'),(8,'A','attendance')]:
            self.write_book(physics=marks); rows=self.state(self.plan())['rows']; rows[0]['attendance']=attendance
            record=self.compare(rows)['records'][0]
            self.assertEqual(record['award_action'],'manual_only')
            self.assertTrue(any(fragment.lower() in issue.lower() for issue in record['manual_issues']))

    def test_submit_and_next_is_never_a_read_only_pager(self):
        driver=Mock(); save=Mock()
        save.get_attribute.side_effect=lambda name:{'href':"javascript:__doPostBack('bodypH$btnSaveMarks','')",'onclick':'return validateMarks();'}.get(name)
        driver.find_elements.return_value=[save]
        self.assertIsNone(matching._read_only_next(driver,2)); save.click.assert_not_called()
        link=Mock(); link.get_attribute.side_effect=lambda name:{'href':"javascript:__doPostBack('bodypH$gridStudents','Page$2')",'onclick':''}.get(name)
        driver.find_elements.return_value=[link]
        self.assertIs(matching._read_only_next(driver,2),link)
        self.assertIsNone(matching._read_only_next(driver,3))

    def test_absence_in_unselected_subject_blocks_selected_subject(self):
        self.write_book(physics='AB')
        jobs=self.jobs(subjects=['MA'])
        self.assertTrue(any('Automatic Principal approval hold:' in reason for reason in jobs[0]['plan'].awards[0].issues))
        self.assertEqual(source_absences(self.file)[0]['subject'],'Physics')

    def test_absence_history_persists_deduplicates_and_is_scoped_and_validated(self):
        hold={'registration':'REG1','exam_roll':'301000001','subject':'Physics','reason':'Synthetic absence'}
        policy=CohortAbsencePolicy(self.history,'school-test','11','2025-26')
        policy.add([hold,hold],origin='synthetic',evidence='test')
        restored=CohortAbsencePolicy(self.history,'SCHOOL-TEST','11','2025-26')
        self.assertEqual(restored.holds,[hold])
        plan=self.plan('Mathematics'); apply_holds(plan,restored.holds)
        self.assertTrue(plan.awards[0].issues)
        for account,cls,session in [('another-school','11','2025-26'),('school-test','12','2025-26'),('school-test','11','2026-27')]:
            self.assertEqual(CohortAbsencePolicy(self.history,account,cls,session).holds,[])
        text=policy.path.read_text(); policy.path.write_text(text.replace('Synthetic absence','altered'))
        with self.assertRaises(ValueError): CohortAbsencePolicy(self.history,'school-test','11','2025-26')

    def test_formula_in_unselected_attendance_blocks_before_browser(self):
        book=load_workbook(self.file); book.active.cell(2,12,'=IF(1,"Present","Absent")'); book.save(self.file); book.close()
        with self.assertRaises(ValueError): self.jobs(subjects=['MA'])

    def test_live_absence_in_unselected_subject_blocks_before_any_fill(self):
        jobs=self.jobs(subjects=['MA']); driver=Mock(); seen=[]
        def state(d,plan,account,**kwargs):
            seen.append(plan.subject); value=self.state(plan)
            if plan.subject=='Physics': value['rows'][0]['attendance']='A'
            return value
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=state),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            report=batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',execute=True,confirm_session=True,confirm_callback=Mock(return_value=False),driver_factory=Mock(return_value=driver),login_callback=Mock())
        self.assertIn('Physics',seen); self.assertIn('Biology',seen)
        self.assertEqual(report['subjects'][0]['status'],'cancelled')
        fill.assert_not_called(); driver.find_element.assert_not_called()
        self.assertTrue(report['class_wide_holds'])

    def test_unread_pages_require_fresh_batch_confirmation_without_false_complete_coverage(self):
        jobs=self.jobs(subjects=['MA']); driver=Mock(); driver.find_elements.return_value=[]
        def state(d,plan,account,**kwargs):
            value=self.state(plan); value['total']=2; return value
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=state),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(browser,'_fill') as fill:
            result=batch.run_jobs(jobs,self.file,self.history,'school-test',student_class='11',session='2025-26',execute=True,confirm_session=True,confirm_callback=Mock(return_value=False),driver_factory=Mock(return_value=driver),login_callback=Mock())
        fill.assert_not_called(); driver.find_element.assert_not_called()
        folder=next((self.history/'References'/'_class_runs').iterdir())
        self.assertTrue((folder/'matching.pdf').exists()); self.assertTrue((folder/'matching.xlsx').exists())
        report=json.loads((folder/'report.json').read_text())
        self.assertEqual(report['status'],'cancelled'); self.assertEqual(len(report['safety_preflight']),3)
        self.assertTrue(all(not item['matching_summary']['coverage_complete'] for item in report['safety_preflight']))

    def test_exports_preserve_exact_text_identifiers_numeric_marks_and_review_notes(self):
        jobs=self.jobs(subjects=['PH']); driver=Mock()
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=lambda d,p,a,**k:self.state(p)),patch.object(browser,'_evidence',return_value='synthetic'):
            self.run_jobs(jobs,driver_factory=Mock(return_value=driver),login_callback=Mock())
        folder=next((self.history/'References'/'_class_runs').iterdir()); book=load_workbook(folder/'matching.xlsx')
        self.assertEqual(book.sheetnames,['Matching Summary','Paired Student Matching','Class-wide Holds'])
        ws=book['Paired Student Matching']
        for cell in ['B2','C2','E2','F2']: self.assertEqual(ws[cell].data_type,'s')
        self.assertEqual(ws['C2'].value,'301000001'); self.assertEqual(ws['J2'].value,8)
        self.assertEqual(ws['J2'].data_type,'n'); self.assertEqual(ws['L2'].value,10)
        self.assertIn('Saved awards disappear',book['Matching Summary']['B7'].value)
        book.close(); self.assertTrue((folder/'matching.pdf').read_bytes().startswith(b'%PDF-'))
        report=json.loads((folder/'report.json').read_text())
        self.assertEqual(len(report['matching_exports']['excel_sha256']),64)

    def test_report_export_failure_is_durable_and_never_reported_as_success(self):
        jobs=self.jobs(subjects=['PH']); driver=Mock()
        with patch.object(browser,'open_subject'),patch.object(browser,'snapshot',side_effect=lambda d,p,a,**k:self.state(p)),patch.object(browser,'_evidence',return_value='synthetic'),patch.object(matching,'export_matching_reports',side_effect=OSError('Synthetic unavailable report disk')):
            with self.assertRaises(OSError): self.run_jobs(jobs,driver_factory=Mock(return_value=driver),login_callback=Mock())
        folder=next((self.history/'References'/'_class_runs').iterdir()); report=json.loads((folder/'report.json').read_text())
        self.assertEqual(report['status'],'failed'); self.assertIn('unavailable report disk',report['matching_export_error'])
        driver.find_element.assert_not_called(); self.assertFalse((self.history/'References'/'internal-assessment.lock').exists())

if __name__=='__main__': unittest.main()

