"""Class-wide absence holds, including subjects outside the submission selection."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import os
import re
from .internal_assessment import Audit, SUBJECTS, norm, subject_name, identifier


def source_absences(excel_file,sheet='Auto'):
    from openpyxl import load_workbook
    book=load_workbook(excel_file,read_only=True,data_only=False,keep_links=False)
    try:
        if sheet in ('Auto','',None): sheet='JKBOSE_Internal_Assessment' if 'JKBOSE_Internal_Assessment' in book.sheetnames else 'Awards_Matrix'
        ws=book[sheet]
        if ws.max_row>20000 or ws.max_column>300: raise ValueError('Workbook dimensions exceed supported bounds.')
        rows=list(ws.iter_rows())
        header=next((i for i,row in enumerate(rows[:40]) if any(norm(c.value) in {'regno','boardregno','boardregistrationnumber','registrationno','registrationnumber'} for c in row)
                     and any(norm(c.value) in {'studentname','studentsname','name'} for c in row)),None)
        if header is None: raise ValueError('Cannot check class-wide absences without awards headers.')
        labels=[norm(c.value) for c in rows[header]]
        def column(options): return next((i for i,label in enumerate(labels) if label in options),None)
        reg_col=column({'regno','boardregno','boardregistrationnumber','registrationno','registrationnumber'})
        exam_col=column({'examrollno','boardrollno','rollno'})
        name_col=column({'studentname','studentsname','name'})
        subject_col=column({'subjectcode','subjectname','subject'})
        issue_col=column({'sourceissues'})
        mark_cols=[i for i,label in enumerate(labels) if label in {'marks','practicalmarks','vivamarks','totalmarks','attendance','presentabsent'}]
        subject_cols=[(i,subject_name(cell.value)) for i,cell in enumerate(rows[header]) if subject_name(cell.value) in SUBJECTS.values()]
        holds=[]; identities={}; exam_owners={}
        for row in rows[header+1:]:
            if row[reg_col].value is None and row[name_col].value is None: continue
            reg=identifier(row[reg_col].value)
            if not reg: continue  # The normal importer blocks this structurally invalid row.
            exam=str(row[exam_col].value or '').strip() if exam_col is not None else ''
            identities.setdefault(reg,set()).add(exam)
            if exam: exam_owners.setdefault(exam,set()).add(reg)
            entries=[(i,subject_name(row[subject_col].value)) for i in mark_cols] if subject_col is not None else subject_cols
            absent_subjects=set()
            for i,subject in entries:
                if row[i].data_type=='f': raise ValueError('A formula prevents the class-wide absence check; export verified values.')
                if re.fullmatch(r'(a|ab|abs|absent)',str(row[i].value or '').strip(),re.I): absent_subjects.add(subject)
            if issue_col is not None and re.search(r'\babsent\b',str(row[issue_col].value or ''),re.I):
                absent_subjects.add(subject_name(row[subject_col].value) if subject_col is not None else 'source attendance')
            for subject in sorted(absent_subjects): holds.append({'registration':reg,'exam_roll':exam,'subject':subject,'reason':'Marked absent in consolidated source'})
        for reg,exams in identities.items():
            if len(exams)>1:
                for exam in exams: holds.append({'registration':reg,'exam_roll':exam,'subject':'identity review','reason':'Registration has conflicting exam rolls across subjects'})
        for exam,regs in exam_owners.items():
            if len(regs)>1:
                for reg in regs: holds.append({'registration':reg,'exam_roll':exam,'subject':'identity review','reason':'Exam roll belongs to conflicting registrations across subjects'})
        return list({(hold['registration'],hold['exam_roll'],hold['subject'],hold['reason']):hold for hold in holds}.values())
    finally: book.close()


def apply_holds(plan,holds):
    for award in plan.awards:
        relevant=[hold for hold in holds if hold['registration']==award.regno or (hold['exam_roll'] and hold['exam_roll']==award.exam_roll)]
        if relevant:
            subjects=', '.join(sorted({hold['subject'] for hold in relevant}))
            identity_conflict=any(h['reason'] not in {'Marked absent in consolidated source','Marked absent on the live JKBOSE pending page'} for h in relevant)
            reason=('Class-wide manual-only hold: absent or identity conflict in '+subjects+'. Automation is blocked for every subject for this student.' if identity_conflict else 'Automatic Principal approval hold: absence in '+subjects+'. This student requires manual approval in every subject.')
            if reason not in award.issues: award.issues.append(reason)


class CohortAbsencePolicy:
    def __init__(self,root,account,student_class,session):
        scope={'school_account':account.lower(),'class':student_class,'session':session}
        key=hashlib.sha256(json.dumps(scope,sort_keys=True).encode()).hexdigest()[:24]
        self.folder=Path(root)/'_cohort_holds'/key; self.folder.mkdir(parents=True,exist_ok=True)
        self.path=self.folder/'events.jsonl'; self.previous='0'*64; self.holds=[]; self.keys=set();self.origins={}
        if self.path.exists():
            for line in self.path.read_text(encoding='utf-8').splitlines():
                event=json.loads(line); signature=event.pop('hash')
                if event.get('previous')!=self.previous or Audit._hash(event)!=signature: raise ValueError('Class-wide absence history is incomplete or changed; stop for manual review.')
                self.previous=signature
                if event['scope']!=scope: raise ValueError('Absence history scope does not match the school/class/session.')
                if event['kind']=='class_wide_manual_hold':
                    hold=event['hold'];key=self._key(hold)
                    if key not in self.keys:self.holds.append(hold)
                    self.keys.add(key);self.origins.setdefault(key,set()).add(event['origin'])
                elif event['kind']=='source_absence_correction':
                    for hold in event['released_holds']:
                        key=self._key(hold)
                        if key not in self.keys or self.origins[key]!={'consolidated_source'} or hold['reason']!='Marked absent in consolidated source':
                            raise ValueError('Invalid source-only absence correction; manual review required.')
                        self.holds=[h for h in self.holds if self._key(h)!=key];self.keys.remove(key);self.origins.pop(key)
                else:raise ValueError('Unknown absence history event; stop for manual review.')
        self.scope=scope

    @staticmethod
    def _key(hold): return tuple(hold.get(key,'') for key in ['registration','exam_roll','subject','reason'])

    def add(self,holds,*,origin,evidence=''):
        for hold in holds:
            if self._key(hold) in self.keys and origin in self.origins.get(self._key(hold),set()): continue
            event={'time':datetime.now(timezone.utc).isoformat(),'kind':'class_wide_manual_hold','scope':self.scope,
                   'hold':hold,'origin':origin,'evidence':evidence,'previous':self.previous}
            event['hash']=Audit._hash(event)
            with self.path.open('a',encoding='utf-8') as stream:
                stream.write(json.dumps(event)+'\n'); stream.flush(); os.fsync(stream.fileno())
            self.previous=event['hash']
            if self._key(hold) not in self.keys:self.holds.append(hold)
            self.keys.add(self._key(hold));self.origins.setdefault(self._key(hold),set()).add(origin)

    def reconcile_source(self,source,sheet='Auto',current_holds=None):
        """Latest explicit portal handoff supersedes source-only absence holds.

        Never clears live/identity holds, missing students, changed identifier
        pairs, invalid awards, or sources without explicit Present attendance.
        Preserve the original events and append the exact correction evidence.
        Caller must hold the assessment RunLock.
        """
        from .internal_assessment_batch import load_jobs
        from openpyxl import load_workbook
        current_holds=source_absences(source,sheet) if current_holds is None else current_holds
        blocked_regs={h['registration'] for h in current_holds};blocked_exams={h['exam_roll'] for h in current_holds if h['exam_roll']}
        candidates=[h for h in self.holds if self.origins[self._key(h)]=={'consolidated_source'}
                    and h['reason']=='Marked absent in consolidated source'
                    and h['registration'] not in blocked_regs and h['exam_roll'] not in blocked_exams]
        # Most runs have no outdated hold. Reuse the caller's all-source scan
        # and avoid reparsing every subject just to retain current absences.
        if not candidates:return []
        source=Path(source).resolve(strict=True);digest=hashlib.sha256(source.read_bytes()).hexdigest()
        jobs=load_jobs(source,sheet,self.scope['class'],self.scope['session'],all_subjects=True)
        if any(j['plan'] is None for j in jobs):return []
        plans=[j['plan'] for j in jobs];chosen=plans[0].sheet
        if chosen!='JKBOSE_Internal_Assessment':return []
        book=load_workbook(source,read_only=True,data_only=False)
        try:
            rows=list(book[chosen].values);headers=[norm(v) for v in rows[0]]
            aliases=[{'boardregno','registrationno','registrationnumber'},{'examroll','examrollno'},{'attendance'},{'exportschema'}]
            columns=[[i for i,key in enumerate(headers) if key in options] for options in aliases]
            if any(len(matches)!=1 for matches in columns):return []
            regcol,examcol,attcol,schemacol=[matches[0] for matches in columns]
            explicit={};invalid=set()
            for row in rows[1:]:
                reg=str(row[regcol] or '').strip();exam=str(row[examcol] or '').strip()
                if not reg:continue
                explicit.setdefault(reg,set()).add(exam)
                if norm(row[attcol]) not in {'p','present'} or row[schemacol]!='hss-jkbose-ia-v1':invalid.add(reg)
        finally:book.close()
        for plan in plans:
            invalid.update(a.regno for a in plan.awards if a.issues)
        released=[hold for hold in candidates if hold['registration'] not in invalid
                  and explicit.get(hold['registration'])=={hold['exam_roll']}]
        if not released:return []
        if hashlib.sha256(source.read_bytes()).hexdigest()!=digest:raise ValueError('Source changed during absence correction; stop and reopen the workbook.')
        event=dict(time=datetime.now(timezone.utc).isoformat(),kind='source_absence_correction',scope=self.scope,
                   released_holds=released,source=str(source),source_sha256=digest,
                   reason='Latest consolidated handoff explicitly marks every award Present with valid unchanged paired identifiers.',previous=self.previous)
        event['hash']=Audit._hash(event)
        with self.path.open('a',encoding='utf-8') as stream:
            stream.write(json.dumps(event)+'\n');stream.flush();os.fsync(stream.fileno())
        self.previous=event['hash']
        for hold in released:
            key=self._key(hold);self.keys.remove(key);self.origins.pop(key)
        self.holds=[hold for hold in self.holds if self._key(hold) in self.keys]
        return released
