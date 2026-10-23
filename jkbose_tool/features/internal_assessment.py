"""Private, fail-closed Internal Assessment import, preview and audit model.

No browser or network access in this module. The source workbook is never saved.
"""
from dataclasses import dataclass, asdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import csv
import json
import os
import re
import unicodedata
import shutil
from .internal_assessment_history import storage_root, prepare_layout

DEFAULT_HISTORY_ROOT = Path(__file__).resolve().parents[2] / 'history' / 'internal_assessment'

def prepare_history_root(root=None):
    """Copy legacy evidence intact before using the new local default; never overwrite it."""
    target = Path(root or DEFAULT_HISTORY_ROOT).resolve()
    if target == DEFAULT_HISTORY_ROOT.resolve():
        legacy = Path.home() / '.jkbose_tool_private' / 'internal_assessment'
        marker = storage_root(target) / '.legacy-history-imported'
        if legacy.is_dir() and legacy.resolve() != target and not marker.exists():
            if (legacy / 'internal-assessment.lock').exists():
                raise ValueError('Legacy history is locked. Reconcile the interrupted run before migrating history.')
            # Validate all collisions before copying any evidence.
            for source in legacy.rglob('*'):
                destination = storage_root(target) / source.relative_to(legacy)
                if source.is_file() and destination.exists() and (not destination.is_file() or source.read_bytes() != destination.read_bytes()):
                    raise ValueError('Legacy/local history conflict. Reconcile both histories; no evidence was overwritten.')
            for source in legacy.rglob('*'):
                destination = storage_root(target) / source.relative_to(legacy)
                if source.is_file() and not destination.exists():
                    destination.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(source, destination)
            marker.parent.mkdir(parents=True, exist_ok=True)
            marker.write_text('Legacy evidence copied intact; original retained at ' + str(legacy), encoding='utf-8')
    target.mkdir(parents=True, exist_ok=True)
    return prepare_layout(target)

SUBJECTS = {
    'EN': 'General English', 'PH': 'Physics', 'CH': 'Chemistry',
    'BI': 'Biology', 'BO': 'Botany', 'ZO': 'Zoology', 'MA': 'Mathematics',
    'UR': 'Urdu', 'ED': 'Education', 'HT': 'History', 'PS': 'Political Science',
    'EC': 'Economics', 'ES': 'Environmental Science', 'PD': 'Physical Education',
    'SC': 'Science', 'SS': 'Social Science',
}

def norm(value):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKC', str(value or '')).lower())

def subject_name(value):
    key = norm(value)
    for code, label in SUBJECTS.items():
        if key in {norm(code), norm(label), norm(f'{label} ({code})')}:
            return label
    if key in {'english', 'generalenglishen'}: return 'General English'
    if key in {'socialstudies', 'socialstudiesss'}: return 'Social Science'
    if key in {'biologybotanyzoologybi', 'biologybotanyzoology'}: return 'Biology'
    return str(value or '').strip()

def session_key(value):
    text = re.sub(r'[\u2010-\u2015\u2212]', '-', str(value or ''))
    match = re.search(r'\b(20\d{2})\s*-\s*(20\d{2}|\d{2})\b', text)
    if match:
        start = int(match[1]); end = int(match[2]) if len(match[2]) == 4 else start // 100 * 100 + int(match[2])
        if end != start + 1: raise ValueError('Academic session must contain consecutive years (YYYY-YY).')
        return f'{start}-{str(end)[-2:]}'
    match = re.search(r'annual\s+regular\s+(20\d{2})', text, re.I)
    if match: return f'{int(match[1])-1}-{match[1][-2:]}'
    raise ValueError('Session must be YYYY-YY or Annual Regular YYYY.')

def class_key(value):
    match = re.fullmatch(r'(?:class\s*)?(10|11|12)(?:th)?', str(value or '').strip(), re.I)
    if not match: raise ValueError('Class must be 10, 11 or 12.')
    return match[1]

def number(value, label='Marks'):
    text = str(value if value is not None else '').strip()
    if not re.fullmatch(r'\d+(?:\.\d+)?', text):
        raise ValueError(f'{label}: missing, absent or invalid value; manual handling required.')
    result = Decimal(text)
    if not result.is_finite(): raise ValueError(f'{label}: invalid number.')
    return result

def identifier(value):
    if value is None: return ''
    if isinstance(value, bool): raise ValueError('Invalid registration identifier.')
    if isinstance(value, (int, float)):
        if value != int(value) or len(str(int(value))) > 15:
            raise ValueError('Registration number is numeric/rounded in Excel; export it as Text.')
        value = str(int(value))
    text = unicodedata.normalize('NFKC', str(value)).strip()
    if text and not re.fullmatch(r'[A-Za-z0-9-]+', text):
        raise ValueError('Invalid registration number; preserve the original Text identifier.')
    return text

@dataclass
class Award:
    regno: str
    name: str
    exam_roll: str
    marks: str
    issues: list
    components: dict
    source_rows: list
    attendance: str = "P"

@dataclass
class Plan:
    source: str
    source_sha256: str
    sheet: str
    student_class: str
    session: str
    subject: str
    maximum: str
    pass_percent: str
    pass_marks: str
    biology_component_max: str
    awards: list
    component_maxima: dict = None
    combine_biology: bool = True

    def payload(self):
        result = asdict(self)
        for award in result["awards"]:
            if award.get("attendance") == "P": award.pop("attendance")
        # Preserve fingerprints when reading histories created by older versions.
        if self.component_maxima is None: result.pop('component_maxima')
        if self.combine_biology: result.pop('combine_biology')
        return result

    def manifest(self):
        result = self.payload()
        return {key: value for key, value in result.items() if key != 'awards'}

    def rows(self): return [asdict(award) for award in self.awards]

    def digest(self):
        # Includes the interpreted awards as well as the exact source file.
        return hashlib.sha256(json.dumps(self.payload(), sort_keys=True).encode()).hexdigest()

def load_plan(excel_file, sheet='Auto', student_class='11', session='',
              subject='Biology', maximum='Auto', pass_percent='40', biology_component_max=None, combine_biology=True):
    from openpyxl import load_workbook
    path = Path(excel_file).resolve(strict=True)
    if path.suffix.lower() != '.xlsx': raise ValueError('Choose an .xlsx awards workbook.')
    if path.stat().st_size > 20 * 1024 * 1024: raise ValueError('Workbook exceeds 20 MB.')
    cls, session, subject = class_key(student_class), session_key(session), subject_name(subject)
    if subject not in SUBJECTS.values(): raise ValueError('Choose a supported subject.')
    if combine_biology and subject in {'Botany', 'Zoology'}:
        raise ValueError('Select Biology to combine both parts, or uncheck Combine Biology for separate awards.')
    combining = subject == 'Biology' and combine_biology
    wanted = ['Botany', 'Zoology'] if combining else [subject]
    maximum = None if norm(maximum) in {'', 'auto'} else number(maximum, 'Maximum')
    percent = number(pass_percent, 'Pass percentage')
    if (maximum is not None and maximum <= 0) or percent != Decimal('40'):
        raise ValueError('Maximum must be positive; the owner requires a 40% passing threshold.')
    component_max = number(biology_component_max, 'Biology component maximum') if biology_component_max not in (None, '') else None
    if component_max is not None and component_max <= 0: raise ValueError('Biology component maximum must be positive.')
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet in ('Auto', '', None):
            sheet = 'JKBOSE_Internal_Assessment' if 'JKBOSE_Internal_Assessment' in book.sheetnames else 'Awards_Matrix'
        if sheet not in book.sheetnames: raise ValueError('Worksheet not found. Select the completed awards sheet.')
        ws = book[sheet]
        if ws.max_row > 20000 or ws.max_column > 300: raise ValueError('Workbook dimensions exceed the supported bounds.')
        rows = list(ws.iter_rows())
        aliases = {
            'regno': ['RegNo', 'Board Reg. No.', 'Board Registration Number', 'Registration No.', 'Registration Number'],
            'name': ['Student Name', "Student's Name", 'Name'],
            'exam': ['Exam Roll No.', 'Exam Roll No', 'Board Roll No', 'Roll No.'],
            'class': ['Class'], 'session': ['Session'], 'evaluation': ['Evaluation Type'],
            'subject': ['Subject Name', 'Subject', 'Subject Code'],
            'marks': ['Practical Marks', 'Marks', 'Total Marks'], 'maximum': ['Max Marks', 'Maximum Marks'],
            'attendance': ['Attendance', 'Present / Absent'],
            'source_issues': ['Source Issues'], 'schema': ['Export Schema'],
            'class_roll': ['Class Roll No.', 'Class Roll No'],
        }
        header_index, columns, headers = None, {}, []
        for index, row in enumerate(rows[:40]):
            labels = [norm(cell.value) for cell in row]
            found = {}
            for key, options in aliases.items():
                matches = [i for i, label in enumerate(labels) if label and label in {norm(option) for option in options}]
                if len(matches) > 1:
                    if key == 'subject' and len(matches) == 2 and {labels[i] for i in matches} == {'subjectname', 'subjectcode'}:
                        found['subject_secondary'] = matches[1]
                    else: raise ValueError(f'Duplicate {key} columns; use one authoritative column.')
                if matches: found[key] = matches[0]
            if 'regno' in found and 'name' in found:
                header_index, columns, headers = index, found, [str(cell.value or '').strip() for cell in row]
                break
        if header_index is None: raise ValueError('No registration-number and student-name header was found.')
        if sheet == 'JKBOSE_Internal_Assessment' and not {'schema', 'source_issues', 'class_roll', 'exam', 'attendance'}.issubset(columns):
            raise ValueError('Portal handoff is missing safety columns; export the workbook again.')
        matrix = 'marks' not in columns
        preamble = '\n'.join(' '.join(str(cell.value or '') for cell in row) for row in rows[:header_index])
        if matrix:
            if not re.search(r'internal\s+assessment', preamble, re.I):
                raise ValueError('Awards matrix must explicitly identify Internal Assessment; external/custom awards are blocked.')
            class_match = re.search(r'class\s*(10|11|12)(?:th)?', preamble, re.I)
            if not class_match or class_match[1] != cls: raise ValueError('Workbook class does not match the selected class.')
            if session_key(preamble) != session: raise ValueError('Workbook session does not match the selected session.')
            subject_columns = {}
            for i, header in enumerate(headers):
                label = subject_name(header)
                if label in SUBJECTS.values():
                    if label in subject_columns: raise ValueError(f'Duplicate {label} columns.')
                    subject_columns[label] = i
            if any(label not in subject_columns for label in wanted):
                raise ValueError('Required subject columns are missing. Combined Biology requires Botany and Zoology; uncheck Combine for a completed Biology total.')
        else:
            required = {'class', 'session', 'evaluation', 'subject', 'maximum'}
            if not required.issubset(columns): raise ValueError('Row-based awards require Class, Session, Evaluation Type, Subject, Marks and Max Marks.')

        def cell_value(row, index):
            cell = row[index]
            if cell.data_type == 'f': raise ValueError(f'Formula in required source cell {cell.coordinate}; export verified values.')
            return cell.value

        source_limits = {}
        if matrix:
            metadata = [row for row in rows[:header_index] if norm(row[0].value) in {'maxmarks', 'maximummarks'}]
            if len(metadata) > 1: raise ValueError('Duplicate maximum-marks metadata rows.')
            for label in wanted:
                raw = cell_value(metadata[0], subject_columns[label]) if metadata else None
                fallback = component_max if combining else maximum
                limit = number(raw, f'{label} maximum from Excel') if raw not in (None, '') else fallback
                if limit is None:
                    raise ValueError('Maximum marks are missing from Excel. Re-export consolidated awards or explicitly supply verified legacy maxima.')
                if limit <= 0: raise ValueError('Source maximum must be positive.')
                if fallback is not None and limit != fallback: raise ValueError('Excel maximum conflicts with the manually supplied maximum.')
                source_limits[label] = limit
        else:
            for label in wanted:
                limits = set()
                for row in rows[header_index+1:]:
                    if subject_name(cell_value(row, columns['subject'])) != label: continue
                    raw = cell_value(row, columns['maximum'])
                    try:
                        limit = number(raw, f'{label} maximum from Excel')
                        if limit <= 0: raise ValueError('Source maximum must be positive.')
                        limits.add(limit)
                    except ValueError:
                        if maximum is None: raise ValueError(f'{label}: missing/invalid maximum in Excel; re-export approved awards.')
                if len(limits) > 1: raise ValueError(f'{label}: conflicting maxima in Excel; reconcile before proceeding.')
                if limits: source_limits[label] = next(iter(limits))
                elif maximum is not None and (not combining or component_max is not None):
                    source_limits[label] = component_max if combining else maximum
                else: raise ValueError(f'{label}: source maximum is missing; both Biology parts are required when combining.')
        inferred_maximum = sum(source_limits.values(), Decimal(0))
        if maximum is None: maximum = inferred_maximum
        if combining and inferred_maximum != maximum:
            raise ValueError('The Biology component maxima do not add up to the confirmed subject maximum.')

        grouped = {}
        for index, row in enumerate(rows[header_index+1:], header_index+2):
            reg_value = cell_value(row, columns['regno'])
            name = str(cell_value(row, columns['name']) or '').strip()
            if reg_value is None and not name: continue
            reg = identifier(reg_value)
            if not reg or not name: raise ValueError(f'Row {index}: registration number and name are both required.')
            issues = []; attendance = ''
            exam = str(cell_value(row, columns['exam']) or '').strip() if 'exam' in columns else ''
            if not re.fullmatch(r'\d{9}',exam): issues.append('Missing/invalid exam roll; both registration and current exam roll are mandatory.')
            if matrix:
                entries = [(label, cell_value(row, subject_columns[label]), source_limits[label]) for label in wanted]
                if all(str(raw or '').strip().lower() == 'x' for _, raw, _ in entries): continue
            else:
                if class_key(cell_value(row, columns['class'])) != cls or session_key(cell_value(row, columns['session'])) != session:
                    raise ValueError(f'Row {index}: class/session mismatch. Export one cohort at a time.')
                evaluation = norm(cell_value(row, columns['evaluation']))
                if evaluation not in {'internal', 'internalassessment', 'ia'}: raise ValueError(f'Row {index}: not an Internal Assessment award.')
                label = subject_name(cell_value(row, columns['subject']))
                if 'subject_secondary' in columns and subject_name(cell_value(row, columns['subject_secondary'])) != label:
                    raise ValueError(f'Row {index}: subject name and code conflict.')
                if label not in wanted: continue
                expected_limit = source_limits[label] if combining else maximum
                if 'schema' in columns:
                    if cell_value(row, columns['schema']) != 'hss-jkbose-ia-v1':
                        raise ValueError(f'Row {index}: unsupported portal export schema.')
                    if not {'source_issues', 'class_roll', 'exam', 'attendance'}.issubset(columns):
                        raise ValueError('Portal handoff is missing safety columns; export the workbook again.')
                    source_issue = str(cell_value(row, columns['source_issues']) or '').strip()
                    if source_issue and not (source_issue == "Absent: admin must handle manually." and (norm(cell_value(row,columns['attendance'])) in {'a','ab','abs','absent'} or norm(cell_value(row,columns['marks'])) in {'a','ab','abs','absent'})): issues.append('Portal review required: ' + source_issue)
                    if not re.fullmatch(r'\d{1,3}', str(cell_value(row, columns['class_roll']) or '').strip()):
                        issues.append('Missing/invalid assigned class roll; admin must review.')
                    if not re.fullmatch(r'\d{9}', exam): issues.append('Missing/invalid current exam roll; admin must review.')
                try:
                    limit = number(cell_value(row, columns['maximum']), 'Source maximum')
                except ValueError as error:
                    issues.append(str(error))
                    limit = expected_limit  # This row stays blocked; no missing maximum is accepted.
                if limit != expected_limit: issues.append('Source maximum does not match the confirmed maximum.')
                entries = [(label, cell_value(row, columns['marks']), limit)]
                if 'attendance' in columns:
                    attendance = norm(cell_value(row, columns['attendance']))
                    if attendance not in {'', 'p', 'present', 'a', 'ab', 'abs', 'absent'}: issues.append('Unknown attendance: manual handling required.')
            item = grouped.setdefault(reg, {'name': name, 'exam': exam, 'values': {}, 'issues': [], 'rows': []})
            if item['name'] != name or item['exam'] != exam:
                raise ValueError(f'Row {index}: conflicting identities for the same registration number.')
            item['issues'].extend(issues); item['rows'].append(index)
            for label, raw, limit in entries:
                if label in item['values']: raise ValueError(f'Row {index}: duplicate student/subject award.')
                try:
                    absent_mark = norm(raw) in {'a','ab','abs','absent'}
                    if absent_mark or attendance in {'a','ab','abs','absent'}:
                        if absent_mark and attendance in {'p','present'}: item['issues'].append('Present attendance conflicts with absent mark.')
                        if not absent_mark and raw not in (None,'',0,'0'): item['issues'].append('Absent attendance conflicts with numeric marks.')
                        item['values'][label] = 'AB'; continue
                    value = number(raw, label)
                    if value > limit: raise ValueError(f'{label}: exceeds source maximum.')
                    item['values'][label] = str(value)
                except ValueError as error:
                    item['values'][label] = ''; item['issues'].append(str(error))
        awards = []
        for reg, item in grouped.items():
            if any(label not in item['values'] or item['values'][label] == '' for label in wanted):
                item['issues'].append('Required mark/component is missing; never substitute zero.')
            mark = ''
            is_absent = 'AB' in item['values'].values()
            if is_absent: item['issues'].append(ABSENT_APPROVAL_HOLD)
            if not is_absent and all(item['values'].get(label) for label in wanted):
                value = sum((Decimal(item['values'][label]) for label in wanted), Decimal(0))
                mark = str(value)
                if value > maximum: item['issues'].append('Combined marks exceed the official maximum.')
                if value < maximum * percent / 100: item['issues'].append('Below 40% pass marks: admin must handle manually.')
            awards.append(Award(reg, item['name'], item['exam'], mark, list(dict.fromkeys(item['issues'])), item['values'], item['rows'], 'A' if is_absent else 'P'))
        if not awards: raise ValueError('No awards for this class/session/subject were found.')
        exam_owners={}
        for award in awards:
            if award.exam_roll: exam_owners.setdefault(award.exam_roll,set()).add(award.regno)
        for award in awards:
            if len(exam_owners.get(award.exam_roll,set()))>1:
                award.issues.append('Duplicate exam roll across registration numbers; manual identity reconciliation required.')
        if hashlib.sha256(path.read_bytes()).hexdigest() != source_hash: raise ValueError('Workbook changed during import; reopen the preview.')
        plan = Plan(str(path), source_hash, sheet, cls, session, subject, str(maximum), str(percent),
                    str(maximum * percent / 100), str(component_max or ''), awards,
                    {label: str(limit) for label, limit in source_limits.items()}, combine_biology)
        from .internal_assessment_identity import apply_identity_confirmations
        apply_identity_confirmations(plan)
        return plan
    finally:
        book.close()

def live_maximum(live):
    """The official page, validation field and every input must agree."""
    values = [row.get('max') for row in live['rows']]
    values.extend(live[key] for key in ['maximum', 'hidden_maximum'] if key in live)
    limits = {number(value, 'Official JKBOSE maximum') for value in values}
    if len(limits) != 1 or next(iter(limits)) <= 0:
        raise ValueError('JKBOSE displayed, validation and row maxima are missing or conflicting.')
    return next(iter(limits))

ABSENT_APPROVAL_HOLD = 'Absent award; automatic Principal approval prohibited.'
APPROVAL_HOLD_PREFIX = 'Automatic Principal approval hold: '

def submission_issues(award):
    return [issue for issue in award.issues if issue != ABSENT_APPROVAL_HOLD and not issue.startswith(APPROVAL_HOLD_PREFIX)]

def rejected_correction(plan,award,row):
    if award is None:return False
    proof=getattr(plan,'rejected_receipts',{}).get((award.regno,award.exam_roll))
    return bool(proof and proof.get('source_sha256')==plan.source_sha256 and proof.get('status')=='Rejected'
                and proof.get('subject')==plan.subject and norm(proof.get('name'))==norm(award.name)
                and row.get('regno')==award.regno and row.get('exam_roll')==award.exam_roll and norm(row.get('name'))==norm(award.name))


def rejected_existing_value(plan,award,row):
    if not rejected_correction(plan,award,row):return False
    proof=plan.rejected_receipts[(award.regno,award.exam_roll)]
    try:
        if proof.get('detail')=='status_only':return row.get('value') in ('',None,'A','AB') or Decimal(0)<=number(row.get('value'))<=number(row.get('max'))
        return number(row.get('value'))==(Decimal(0) if proof['attendance']=='A' and str(proof['marks']).strip().upper() in {'','A','AB','ABS','ABSENT'} else number(proof['marks']))
    except ValueError:return False


def match_batch(plan, live, previously_verified=(), *, allow_absent=False):
    """Return a complete batch or problems. Never return a partial safe subset."""
    if Decimal(plan.pass_percent) != 40 or Decimal(plan.maximum) <= 0:
        raise ValueError('Invalid passing policy; no marks may be entered.')
    by_reg = {award.regno: award for award in plan.awards}
    problems, matched, seen, seen_exams = [], [], set(), set()
    try: official = live_maximum(live) if live['rows'] else Decimal(plan.maximum)
    except ValueError as error: return [], [str(error)]
    for row in live['rows']:
        reg = identifier(row.get('regno'))
        if not reg or reg in seen: problems.append('Missing/duplicate registration number on JKBOSE.'); continue
        seen.add(reg)
        award = by_reg.get(reg)
        if not award: problems.append(f'{reg}: no matching source award.'); continue
        if reg in previously_verified: problems.append(f'{reg}: previously verified award reappeared; reconcile manually.')
        issues = submission_issues(award) if allow_absent else award.issues
        if issues: problems.append(f'{reg}: ' + '; '.join(issues))
        if norm(row.get('name')) != norm(award.name): problems.append(f'{reg}: student name differs; review identity manually.')
        live_exam=str(row.get('exam_roll','')).strip()
        if not re.fullmatch(r'\d{9}',award.exam_roll) or not re.fullmatch(r'\d{9}',live_exam):
            problems.append(f'{reg}: valid registration AND exam roll are both mandatory; manual review required.')
        if live_exam in seen_exams: problems.append(f'{reg}: duplicate exam roll in the displayed JKBOSE batch.')
        seen_exams.add(live_exam)
        if live_exam != award.exam_roll:
            problems.append(f'{reg}: registration/exam-roll pair differs; class roll or name cannot replace either identifier.')
        if subject_name(row.get('subject')) != plan.subject: problems.append(f'{reg}: live subject mismatch.')
        correcting=allow_absent and rejected_correction(plan,award,row)
        if row.get('attendance') not in ({'P','A'} if allow_absent and (award.attendance=='A' or correcting) else {'P'}): problems.append(f'{reg}: attendance conflicts with the source; manual handling required.')
        if row.get('attendance_count', 1) != 1: problems.append(f'{reg}: attendance selection is missing or ambiguous.')
        if (row.get('disabled') or row.get('readonly')) and not (allow_absent and row.get('attendance')=='A' and (award.attendance=='A' or correcting)): problems.append(f'{reg}: marks are locked.')
        if correcting and award.attendance=='P' and row.get('attendance')=='A' and (not row.get('present_input_id') or not row.get('present_enabled')):problems.append(f'{reg}: verified Present radio control is unavailable.')
        if allow_absent and award.attendance=='A':
            if row.get('attendance')!='A' and (not row.get('absent_input_id') or not row.get('absent_enabled')): problems.append(f'{reg}: verified Absent radio control is unavailable.')
            if official != Decimal(plan.maximum): problems.append(f'{reg}: Excel maximum differs from official JKBOSE maximum.')
            if row.get('value') not in (('',None,'0','0.0') if row.get('attendance')=='A' else ('',None)) and not (correcting and rejected_existing_value(plan,award,row)): problems.append(f'{reg}: existing marks conflict with Absent; never overwrite.')
            matched.append(award); continue
        try:
            mark = number(award.marks)
            if mark < official * Decimal('0.4'):
                problems.append(f'{reg}: below 40% of the official JKBOSE maximum; manual submission only.')
            if mark > official:
                problems.append(f'{reg}: exceeds the official JKBOSE maximum; manual submission only.')
            if plan.subject == 'Biology' and plan.combine_biology:
                parts = [number(award.components.get(key), key) for key in ['Botany', 'Zoology']]
                limits = [number((plan.component_maxima or {}).get(key, plan.biology_component_max), 'Component maximum') for key in ['Botany', 'Zoology']]
                if sum(limits) != Decimal(plan.maximum) or any(part > limit for part, limit in zip(parts, limits)) or sum(parts) != mark:
                    problems.append(f'{reg}: Biology components do not validate against the approved total.')
            if official != Decimal(plan.maximum): problems.append(f'{reg}: Excel maximum differs from official JKBOSE maximum; manual review required.')
            minimum = number(row.get('min', '0'), 'Live minimum')
            if award.marks and Decimal(award.marks) < minimum: problems.append(f'{reg}: marks are below the live minimum.')
            if row.get('step') not in ('any', '', None) and award.marks:
                step = number(row['step'], 'Live step')
                if step <= 0 or (Decimal(award.marks) - minimum) % step: problems.append(f'{reg}: marks do not match the permitted decimal step.')
            if row.get('value') not in ('', None) and number(row['value']) != Decimal(award.marks) and not (correcting and rejected_existing_value(plan,award,row)): problems.append(f'{reg}: existing marks differ; never overwrite.')
        except ValueError as error: problems.append(f'{reg}: {error}')
        matched.append(award)
    if not live['rows']: problems.append('No pending rows; a separate completion check is required.')
    return matched, problems

def verify_transition(before, after):
    before_regs = {row['regno'] for row in before['rows']}
    after_regs = {row['regno'] for row in after['rows']}
    expected = before['total'] - len(before['rows'])
    return (bool(before_regs) and expected >= 0 and after['total'] == expected
            and before_regs.isdisjoint(after_regs)
            and ((expected == 0 and not after['rows']) or (expected > 0 and bool(after['rows']))))

def saved_outcomes(events):
    """Latest durable save/rejection outcome for each student, in event order."""
    outcomes={}
    for event in events:
        if event['kind'] in {'verified_saved','portal_award_rejected','official_recorded_award'} or (event['kind']=='manual_reconciliation' and event['disposition']=='saved_manually_verified'):
            for reg in event.get('registrations') or ([event['regno']] if event.get('regno') else []):outcomes[reg]=event
    return outcomes

class Audit:
    """Durable hash-chained journal. A partial write or uncertain submit blocks resume."""
    def __init__(self, root, plan):
        self.folder = storage_root(root) / plan.digest()[:24]
        self.folder.mkdir(parents=True, exist_ok=True)
        self.path = self.folder / 'events.jsonl'
        self.events = []; self.previous = '0' * 64
        if self.path.exists():
            for line in self.path.read_text(encoding='utf-8').splitlines():
                event = json.loads(line); signature = event.pop('hash')
                if event.get('previous') != self.previous or self._hash(event) != signature:
                    raise ValueError('Audit history is incomplete or modified. Reconcile manually before resuming.')
                event['hash'] = signature; self.events.append(event); self.previous = signature
        manifest = self.folder / 'manifest.json'
        content = {**plan.manifest(), 'plan_digest': plan.digest()}
        if manifest.exists():
            if json.loads(manifest.read_text(encoding='utf-8')) != content: raise ValueError('Audit manifest mismatch.')
        else:
            with manifest.open('x', encoding='utf-8') as stream:
                json.dump(content, stream, indent=2); stream.flush(); os.fsync(stream.fileno())

    @staticmethod
    def _hash(event): return hashlib.sha256(json.dumps(event, sort_keys=True, ensure_ascii=True).encode()).hexdigest()

    def append(self, kind, **data):
        event = {'time': datetime.now(timezone.utc).isoformat(), 'kind': kind, 'previous': self.previous, **data}
        event['hash'] = self._hash(event)
        with self.path.open('a', encoding='utf-8') as stream:
            stream.write(json.dumps(event, ensure_ascii=True) + '\n'); stream.flush(); os.fsync(stream.fileno())
        self.events.append(event); self.previous = event['hash']

    def rejected(self):
        return {reg for reg,event in saved_outcomes(self.events).items() if event['kind']=='portal_award_rejected'}

    def verified(self):
        return {reg for reg,event in saved_outcomes(self.events).items() if event['kind']=='verified_saved'}

    def manual_saved(self):
        return {reg for reg,event in saved_outcomes(self.events).items() if event['kind']=='manual_reconciliation'}

    def record_portal_rejection(self, registration, administrator='Principal/Administrator', note='Rejected on portal for resubmission'):
        regs = [registration] if isinstance(registration, str) else list(registration)
        self.append('portal_award_rejected', registrations=regs, regno=regs[0] if len(regs)==1 else '',
                    administrator=administrator.strip(), note=note.strip(),
                    disposition='returned_to_pending',
                    verification='Portal rejection acknowledged; award returned to pending pool.')

    def unresolved(self):
        attempted = {event['batch_id'] for event in self.events if event['kind'] == 'submit_attempt'}
        settled = {event['batch_id'] for event in self.events if event['kind'] in {'verified_saved', 'manual_reconciliation'}}
        return attempted - settled

    @classmethod
    def existing(cls, folder):
        folder = Path(folder).resolve(strict=True)
        manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
        expected = manifest.pop('plan_digest')
        events = [json.loads(line) for line in (folder / 'events.jsonl').read_text(encoding='utf-8').splitlines()]
        preview = next((event.get('source_preview', event.get('awards')) for event in events
                        if event['kind'] in {'run_started', 'offline_preview'}), None)
        if preview is None: raise ValueError('No source preview exists; this journal cannot be reconciled here.')
        plan = Plan(**manifest, awards=[Award(**award) for award in preview])
        if plan.digest() != expected or folder.name != expected[:24]: raise ValueError('History source/manifest does not match.')
        return cls(folder.parent, plan), plan

    def reconcile(self, batch_id, disposition, administrator, note):
        if batch_id not in self.unresolved(): raise ValueError('Select an unresolved batch; settled history cannot be rewritten.')
        if disposition not in {'saved_manually_verified', 'not_saved_manually_verified'}:
            raise ValueError('Confirm explicitly whether the awards were saved or not saved on JKBOSE.')
        if len(administrator.strip()) < 2 or len(note.strip()) < 10:
            raise ValueError('Administrator name and a specific manual verification note are required.')
        attempt = next(event for event in self.events if event['kind'] == 'submit_attempt' and event['batch_id'] == batch_id)
        self.append('manual_reconciliation', batch_id=batch_id, disposition=disposition,
                    administrator=administrator.strip(), note=note.strip(), registrations=attempt['registrations'],
                    verification='Human attestation after checking JKBOSE. This is not automated verification.')

    def report(self, plan):
        saved, manual_saved = self.verified(), self.manual_saved()
        recorded={reg:e for reg,e in saved_outcomes(self.events).items() if e['kind']=='official_recorded_award'}
        rejected = self.rejected()
        prior_saved={reg for event in self.events if event['kind']=='history_resume' for reg in event['verified_registrations']} - rejected
        prior_manual={reg for event in self.events if event['kind']=='history_resume' for reg in event['manual_registrations']} - rejected
        official = next((event for event in reversed(self.events) if event['kind'] == 'official_maximum_checked'), None)
        blocked, preview_regs = set(), []
        for event in self.events:
            if event['kind'] == 'batch_preview': preview_regs = event['registrations']
            if event['kind'] == 'batch_blocked': blocked.update(event.get('registrations', preview_regs))
        def record(award):
            issues = list(award.issues)
            if official:
                if official['mismatch']: issues.append('Excel maximum differs from official JKBOSE maximum; manual review required.')
                try:
                    mark = number(award.marks)
                    if mark > number(official['maximum']): issues.append('Exceeds official JKBOSE maximum; manual submission only.')
                    if mark < number(official['pass_marks']): issues.append('Below official 40% pass threshold; manual submission only.')
                except ValueError: issues.append('Absent/missing/invalid marks; manual submission only.')
            if award.regno in recorded:issues.append(('Already recorded status on JKBOSE; saved marks/attendance are not exposed. Report only, never retry. Status evidence: ' if recorded[award.regno]['saved_records'][award.regno].get('detail')=='status_only' else 'Already recorded on JKBOSE with different values; report only, never retry. Saved: ')+json.dumps(recorded[award.regno]['saved_records'][award.regno],sort_keys=True))
            return {**asdict(award), 'issues': list(dict.fromkeys(issues)),
                    'status': 'already_recorded_status_only' if award.regno in recorded and recorded[award.regno]['saved_records'][award.regno].get('detail')=='status_only' else 'already_recorded_source_changed' if award.regno in recorded else 'portal_rejected' if award.regno in rejected else 'verified_saved' if award.regno in saved else 'saved_manually_verified' if award.regno in manual_saved else 'previously_verified_saved' if award.regno in prior_saved else 'previously_manually_verified_saved' if award.regno in prior_manual else 'manual_required' if issues else 'batch_blocked' if award.regno in blocked else 'not_verified_by_this_tool'}
        content = {'manifest': plan.manifest(), 'verified_saved_count': len(saved),
                   'previously_verified_saved_count':len(prior_saved-saved),
                   'previously_manually_verified_saved_count':len(prior_manual-manual_saved),
                   'official_jkbose_maximum': official['maximum'] if official else None,
                   'official_pass_marks': official['pass_marks'] if official else None,
                   'unresolved_batches': sorted(self.unresolved()),
                   'batch_issues': [event['reasons'] for event in self.events if event['kind'] == 'batch_blocked'],
                   'records': [record(award) for award in plan.awards]}
        temporary = self.folder / 'report.tmp'
        with temporary.open('w', encoding='utf-8') as stream:
            json.dump(content, stream, indent=2); stream.flush(); os.fsync(stream.fileno())
        os.replace(temporary, self.folder / 'report.json')
        # Readable Excel-compatible report; quote formula-like strings as text.
        def safe(value):
            text = str(value if value is not None else '')
            return "'" + text if text.startswith(('=', '+', '-', '@', '\t', '\r')) else text
        temporary = self.folder / 'report-csv.tmp'
        with temporary.open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(['Registration No', 'Student Name', 'Exam Roll', 'Subject', 'Marks', 'Excel Maximum', 'Official JKBOSE Maximum', 'Official Pass Marks', 'Status', 'Manual review reason'])
            for record in content['records']:
                writer.writerow([safe(value) for value in [record['regno'], record['name'], record['exam_roll'], plan.subject,
                    record['marks'], plan.maximum, content['official_jkbose_maximum'], content['official_pass_marks'], record['status'], '; '.join(record['issues'])]])
            for reasons in content['batch_issues']:
                for reason in reasons: writer.writerow(['', '', '', plan.subject, '', plan.maximum, content['official_jkbose_maximum'], content['official_pass_marks'], 'batch_blocked', safe(reason)])
            stream.flush(); os.fsync(stream.fileno())
        os.replace(temporary, self.folder / 'report.csv')

class RunLock:
    def __init__(self, root): self.path = storage_root(root) / 'internal-assessment.lock'
    def __enter__(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        try:
            descriptor = os.open(self.path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
        except FileExistsError:
            raise ValueError('Another run or interrupted run holds the history lock. Reconcile it manually; do not run concurrently.')
        with os.fdopen(descriptor, 'w') as stream:
            stream.write(str(os.getpid())); stream.flush(); os.fsync(stream.fileno())
        return self
    def __exit__(self, *_): self.path.unlink(missing_ok=True)
