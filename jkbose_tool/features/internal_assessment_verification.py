"""Offline Internal Assessment verification; never signs in or changes an award."""
from pathlib import Path
from datetime import datetime,timezone,timedelta
from decimal import Decimal,InvalidOperation
from collections import defaultdict
from contextvars import ContextVar
from html.parser import HTMLParser
from xml.sax.saxutils import escape
import hashlib,json,re,os,uuid,shutil

SUBJECTS={'EN':'General English','PH':'Physics','CH':'Chemistry','BI':'Biology','ES':'Environmental Science','PD':'Physical Education','MA':'Mathematics','UR':'Urdu','ED':'Education','HT':'History','PS':'Political Science','EC':'Economics','BO':'Botany','ZO':'Zoology','SC':'Science','SS':'Social Science','HI':'Hindi','KA':'Kashmiri','AR':'Arabic','SA':'Sanskrit','GE':'Geography'}
ALIASES={re.sub('[^a-z0-9]','',v.lower()):k for k,v in SUBJECTS.items()};ALIASES.update({k.lower():k for k in SUBJECTS});ALIASES.update(english='EN',botanyzoology='BI',biologybotanyzoology='BI')
VOC={'healthcare','healthcarehtc','htc','itandites','itites','ite','informationtechnology'}
ID=r'(?:\d{8,24}|[A-Za-z]{1,4}\d{7,23})'
RULES=ContextVar('verification_subject_rules',default={})

def subject_catalog():return dict(SUBJECTS,**{k:r['label']for k,r in RULES.get().items()})

def parse_rules(text):
    rules={}
    for entry in re.split(r'[;\n]+',text or ''):
        if not entry.strip():continue
        if '='not in entry:raise ValueError('Subject maxima use Subject=Maximum, separated by semicolons.')
        name,value=entry.rsplit('=',1);name=name.strip();maximum_value=numeric(value)
        if not name or not 0<Decimal(maximum_value)<=1000:raise ValueError('Enter a subject and a positive maximum up to 1,000.')
        code=ALIASES.get(norm(name),norm(name).upper())
        if code in rules:raise ValueError('Duplicate subject maximum: '+name)
        rules[code]=dict(label=SUBJECTS.get(code,name),maximum=maximum_value)
    return rules

def norm(v):return re.sub('[^a-z0-9]','',str(v or '').lower())
def subject(v):
    text=re.sub(r'\s*\([A-Z]+\)\s*$','',str(v or '')).strip();key=norm(text)
    if key in VOC:return None
    if key in ALIASES:return ALIASES[key]
    for code,rule in RULES.get().items():
        if key in{norm(code),norm(rule['label'])}:return code
    raise ValueError('Unrecognised subject: '+str(v)+'. Add its name and maximum under Subject maxima.')
def numeric(v):
    try:n=Decimal(str(v).strip())
    except InvalidOperation:raise ValueError('Invalid numeric marks: '+str(v))
    if not n.is_finite():raise ValueError('Non-finite marks are not supported.')
    return format(n.normalize(),'f')
def maximum(code,cls):
    if code in RULES.get():return Decimal(RULES.get()[code]['maximum'])
    if cls not in{'11','12'}or code not in{'EN','PH','CH','BI','ES','PD','MA','UR','ED','HT','PS','EC','BO','ZO'}:return None
    return 5 if code in {'BO','ZO'} else (15 if cls=='12'else 10)if code=='PD'else 10 if code in {'PH','CH','BI','ES'}else 20
def session_end(session):
    if not re.fullmatch(r'20\d{2}-\d{2}',session)or int(session[-2:])!=(int(session[:4])+1)%100:raise ValueError('Use academic session YYYY-YY with consecutive years.')
    return int(session[:4])+1

def record(roll,reg,name,code,mark,mx,attendance,location):
    roll=str(roll).strip();reg=str(reg).strip()
    if not re.fullmatch(r'\d{9}',roll)or not re.fullmatch(r'[A-Za-z0-9-]+',reg):raise ValueError('Invalid exam-roll / registration pair at '+location)
    absent=norm(attendance)in{'a','ab','abs','absent'}or norm(mark)in{'a','ab','abs','absent'}
    present=norm(attendance)in{'p','present'}
    return dict(roll=roll,reg=reg,name=' '.join(str(name).split()),subject=code,marks=None if absent or mark in(None,'')else numeric(mark),maximum=None if mx in(None,'')else numeric(mx),attendance='A'if absent else'P'if present else'?',location=location)

def pdf_pages(path):
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError
    try:reader=PdfReader(path)
    except PyPdfError as exc:raise ValueError('PDF could not be read: '+str(exc))from exc
    if reader.is_encrypted:raise ValueError('Encrypted PDF: provide an unlocked copy.')
    if len(reader.pages)>2000:raise ValueError('PDF exceeds 2,000-page limit.')
    result=[]
    for n,page in enumerate(reader.pages,1):
        text=page.extract_text(extraction_mode='layout')or''
        if not text.strip():raise ValueError(f'Page {n} has no readable text. Supply a searchable PDF; OCR is not performed.')
        result.append(text.replace('\u2212','-').replace('\u2013','-').replace('\u2014','-'))
    return result

def read_admits(pages,cls,session):
    cards=[]
    for pn,text in enumerate(pages,1):
        blocks=re.split(r'ADMIT\s+CARD',text,flags=re.I)[1:]
        if not blocks:raise ValueError(f'No admit-card blocks identified on page {pn}.')
        for block in blocks:
            cohort=re.search(r'Class\s+(\d+)',block,re.I);year=re.search(r'Annual.*?\b(20\d{2})\b',block,re.I)
            if not cohort or cohort[1]!=cls or not year or int(year[1])!=session_end(session):raise ValueError(f'Admit-card class/exam year differs from selected cohort on page {pn}.')
            reg=re.search(r'R\.\s*R\s*No\.?\s*('+ID+')',block,re.I);name=re.search(r'(?m)^\s*Name\s+(.+?)\s+Gender',block)
            identity_text=block[:reg.start(1)]+' '+block[reg.end(1):]if reg else block
            rolls=re.findall(r'(?<!\d)\d{9}(?!\d)',identity_text)
            if len(rolls)!=1 or not reg or not name:raise ValueError(f'Incomplete admit-card identifiers/name on page {pn}.')
            offered=re.search(r'Compulsory\s+([\s\S]*?)(?:Candidate.?s Signature|Note:)',block,re.I)
            if not offered:raise ValueError('Admit-card subjects are unreadable.')
            body=re.split(r'Vocational',offered[1],flags=re.I)[0];body=re.sub(r'\bAdditional\b',' ',body,flags=re.I)
            parts=[re.sub(r'\s+',' ',v).strip(' ,;')for v in re.split(r'\d+\.\s*',body)[1:]]
            codes=[subject(v)for v in parts];codes=[c for c in codes if c]
            if not codes or len(codes)!=len(set(codes)):raise ValueError('Incomplete or duplicate admit-card subjects.')
            cards.append(dict(roll=rolls[0],reg=reg[1],name=' '.join(name[1].split()),subjects=codes,location=f'p{pn}'))
    return cards

def combine_components(rows):
    groups=defaultdict(list);ordinary=[]
    for row in rows:
        if row['subject']in{'BO','ZO'}:groups[(row['roll'],row['reg'])].append(row)
        else:ordinary.append(row)
    for pair,parts in groups.items():
        if sorted(r['subject']for r in parts)!=['BO','ZO']:raise ValueError('Biology requires one Botany and one Zoology component for '+pair[0])
        if any(x['maximum'] is None or Decimal(x['maximum'])!=5 or (x['marks'] is not None and not 0<=Decimal(x['marks'])<=5)for x in parts):raise ValueError('Invalid Biology component maximum/marks for '+pair[0])
        if len({norm(r['name'])for r in parts})!=1:raise ValueError('Biology component names conflict.')
        r=dict(parts[0],subject='BI',location=' + '.join(x['location']for x in parts));r['maximum']=numeric(sum(Decimal(x['maximum'])for x in parts))if all(x['maximum']is not None for x in parts)else None
        r['attendance']='A'if any(x['attendance']=='A'for x in parts)else'P'if all(x['attendance']=='P'for x in parts)else'?';r['marks']=numeric(sum(Decimal(x['marks'])for x in parts))if all(x['marks']is not None for x in parts)else None;ordinary.append(r)
    return ordinary

def read_matrix(pages,cls):
    rows=[];order=None;serials=[]
    for pn,text in enumerate(pages,1):
        lines=text.splitlines()
        for line in lines:
            codes=re.findall(r'\b(?:'+'|'.join(re.escape(c)for c in subject_catalog())+r')\b',line)
            if codes and 'REG' in line.upper() and re.search(r'NAME',line,re.I):
                if len(codes)!=len(set(codes)):raise ValueError('Duplicate matrix subject columns.')
                if order and order!=codes:raise ValueError('Matrix subject columns changed between pages.')
                order=codes
        for i,line in enumerate(lines):
            m=re.match(r'^\s*('+str(len(serials)+1)+r')\s*(\d{9})\s*('+ID+r')\s*(.*)$',line)
            if not m:
                if re.match(r'^\s*\d{1,4}\s*\d{9}',line):raise ValueError(f'Matrix row serial/identifiers are incomplete or out of order on page {pn}.')
                continue
            if not order:raise ValueError('Matrix subject columns were not identified.')
            seq,roll,reg,rest=m.groups();tokens=rest.split();values=tokens[-len(order)-1:];name=' '.join(tokens[:-len(order)-1])
            if len(values)!=len(order)+1 or any(not re.fullmatch(r'x|a|ab|absent|\d+(?:\.\d+)?',v,re.I)for v in values):raise ValueError(f'Malformed matrix row on page {pn}.')
            if not name:
                adjacent=[lines[i-1].strip(),lines[i+1].strip()if i+1<len(lines)else''];name=' '.join(v for v in adjacent if re.fullmatch(r"[A-Za-z .'-]+",v))
            elif i+1<len(lines) and re.fullmatch(r"\s*[A-Za-z][A-Za-z .'-]*",lines[i+1])and not re.search(r'Certificate|Signature|Principal',lines[i+1],re.I):name+=' '+lines[i+1].strip()
            serials.append(int(seq));numbers=[Decimal(v)for v in values[:-1]if re.fullmatch(r'\d+(?:\.\d+)?',v)]
            if numeric(sum(numbers))!=numeric(values[-1]):raise ValueError(f'Matrix hash total differs on page {pn}, exam roll {roll}.')
            for code,value in zip(order,values[:-1]):
                if value.lower()!='x':rows.append(record(roll,reg,name,code,value,maximum(code,cls),'A'if norm(value)in{'a','ab','absent'}else'P',f'p{pn}, row {seq}'))
    if not rows or sorted(serials)!=list(range(1,len(serials)+1)):raise ValueError('Matrix serial rows are incomplete or duplicated.')
    return combine_components(rows)

def parse_subjects(text,roll,reg,name,location):
    rows=[];segments=re.split(r',\s*(?=[A-Za-z])',re.sub(r'\s+',' ',text.strip()))
    for segment in segments:
        m=re.fullmatch(r'([^:]+):\s*\(([^/]+)/([^)]*)\)\s*',segment)
        if not m:
            m=re.fullmatch(r'([^:]+):\s*(Pending|Rejected|Absent|AB|A)\s*',segment,re.I)
            if not m:raise ValueError('Unrecognised subject marks/status at '+location+': '+segment[:100])
            code=subject(m[1]);
            if code:rows.append(record(roll,reg,name,code,None,None,'A'if norm(m[2])in{'a','ab','absent'}else'?',location))
        else:
            code=subject(m[1])
            if code:rows.append(record(roll,reg,name,code,m[2].strip(),m[3].strip(),'A'if norm(m[2])in{'a','ab','absent'}else'P',location))
    return rows

def read_grouped(pages,cls):
    chunks=[];current=None;serials=[]
    for pn,text in enumerate(pages,1):
        for line in text.splitlines():
            m=re.match(r'^\s*(\d+)\s+(\d{1,2})\s+('+ID+r')\s+(\d{9})\s+(.+)$',line)
            if m:
                if current:chunks.append(current)
                seq,cohort,reg,roll,tail=m.groups()
                if cohort!=cls:raise ValueError('Grouped report contains a different class.')
                current=[int(seq),reg,roll,tail,pn];serials.append(int(seq))
            elif current and line.strip()and not re.search(r'OFFICE OF|GOVT |INTERNAL ASSESSMENT|School Code|Name of the Exam|Session:|Class:|Registration No|Subject-wise|Signature|Page \d|^\s*(?:Sr\.|No:|Dated:)',line,re.I):current[3]+=' '+line.strip()
    if current:chunks.append(current)
    if not chunks or sorted(serials)!=list(range(1,len(serials)+1)):raise ValueError('Grouped student serials are incomplete or duplicated.')
    rows=[]
    for seq,reg,roll,tail,pn in chunks:
        first=re.search(r'([A-Za-z][A-Za-z &/.-]*?):\s*\(',tail)
        # A wide space between the name and first subject is retained by layout extraction.
        split=re.split(r'\s{2,}(?=[A-Za-z][A-Za-z &/.-]*?:)',tail,maxsplit=1)
        if len(split)!=2:raise ValueError('Cannot separate student name and grouped subjects at page '+str(pn))
        rows.extend(parse_subjects(split[1],roll,reg,split[0],f'p{pn}, row {seq}'))
    return combine_components(rows)

def read_flat_pdf(pages,cls):
    rows=[];subnames='|'.join(re.escape(v)for v in sorted(subject_catalog().values(),key=len,reverse=True))
    pat=re.compile(r'^\s*(\d+)\s+('+cls+r')\s*(\d{9})\s+('+ID+r')\s+(.+?)\s+('+subnames+r')\s+(P|A|Present|Absent)\s+(\S+)\s+(\d+(?:\.\d+)?)\s+(.+)$',re.I)
    for pn,text in enumerate(pages,1):
        for line in text.splitlines():
            m=pat.match(line)
            if m:
                _,_,roll,reg,name,label,attendance,marks,mx,_=m.groups();rows.append(record(roll,reg,name,subject(label),marks,mx,attendance,f'p{pn}'))
            elif re.match(r'^\s*\d+\s+\d{1,2}\s*\d{9}\s+'+ID,line):raise ValueError(f'Malformed or different-class award row on page {pn}.')
    if not rows:raise ValueError('No individual JKBOSE award rows were identified.')
    return combine_components(rows)

class HtmlRows(HTMLParser):
    def __init__(self):super().__init__(convert_charrefs=True);self.rows=[];self.row=None;self.cell=None
    def handle_starttag(self,tag,attrs):
        if tag=='tr':self.row=[]
        if tag in{'td','th'}:self.cell=[]
    def handle_data(self,data):
        if self.cell is not None:self.cell.append(data)
    def handle_endtag(self,tag):
        if tag in{'td','th'}and self.cell is not None:
            if self.row is not None:self.row.append(''.join(self.cell).strip())
            self.cell=None
        if tag=='tr'and self.row is not None:
            if self.row:self.rows.append(self.row)
            self.row=None

def read_table(values,cls,session,sheet):
    header=next((i for i,r in enumerate(values[:40])if any(norm(v)in{'boardregno','registrationnumber','registrationno'}for v in r)),None)
    if header is None:raise ValueError('Excel registration headers were not identified.')
    labels=[norm(v)for v in values[header]]
    def col(*names):return next((i for i,v in enumerate(labels)if v in names),None)
    reg=col('boardregno','registrationnumber','registrationno');roll=col('examrollno','rollnumber','rollno');name=col('studentname','name');cohort=col('class');sess=col('session');sub=col('subjectcode','subjectname','subject');mark=col('marks','givenmarks');mx=col('maxmarks','maximummarks');att=col('attendance','presentabsent');grouped=col('subjectwisemarksstatus','subjectwisemarks')
    if None in(reg,roll,name):raise ValueError('Excel identity headers are incomplete.')
    rows=[]
    for n,row in enumerate(values[header+1:],header+2):
        if all(v is None or v==''for v in row):continue
        if row[reg]in(None,'')and row[name]in(None,''):continue
        if cohort is not None and re.sub('[^0-9]','',str(row[cohort]))!=cls:raise ValueError('Excel row belongs to another class.')
        if sess is not None and str(row[sess]).strip()!=session:raise ValueError('Excel academic session differs.')
        loc=f'{sheet}, row {n}'
        if grouped is not None:rows.extend(parse_subjects(str(row[grouped]),row[roll],row[reg],row[name],loc));continue
        if sub is not None:
            code=subject(row[sub])
            if not code:continue
            if mark is None:raise ValueError('Excel award marks column is missing.')
            rows.append(record(row[roll],row[reg],row[name],code,row[mark],row[mx]if mx is not None else None,row[att]if att is not None else'?',loc))
        else:
            found=False
            for i,label in enumerate(values[header]):
                code_match=re.search(r'\(([A-Z]+)\)',str(label))
                if not code_match or code_match[1]not in subject_catalog():continue
                value=row[i]
                if str(value).lower()in{'x','none',''}:continue
                found=True;code=code_match[1];rows.append(record(row[roll],row[reg],row[name],code,value,maximum(code,cls),'A'if norm(value)in{'a','ab','absent'}else'P',loc))
            if not found:raise ValueError('Excel matrix has no recognised award values at '+loc)
    if not rows:raise ValueError('No award rows were identified in Excel.')
    return combine_components(rows)

def read_awards(path,role,cls,session):
    path=Path(path);internal=[]
    if path.suffix.lower()=='.pdf':
        pages=pdf_pages(path);text='\n'.join(pages)
        if 'Subject-wise Marks' in text:rows=read_grouped(pages,cls);kind='JKBOSE grouped marks'
        elif 'Requested On'in text and 'Registration Number'in text:rows=read_flat_pdf(pages,cls);kind='JKBOSE individual requests'
        else:rows=read_matrix(pages,cls);kind='School consolidated matrix'
        cohort=re.search(r'\bClass\s*:?\s*(\d+)',text,re.I)
        if cohort and cohort[1]!=cls:raise ValueError('PDF class differs from selected cohort.')
        source_session=re.search(r'Session\s*(?:&\s*Year)?\s*:\s*[^0-9\n]*(20\d{2})\s*-\s*(\d{2,4})',text,re.I)
        if source_session:
            start=int(source_session[1]);end=int(source_session[2]);end=end if end>99 else 2000+end
            allowed={(int(session[:4]),session_end(session))}
            if session=='2025-26':allowed.add((2026,2027)) # Existing owner-confirmed school/board report mapping.
            if(start,end)not in allowed:raise ValueError('Printed PDF session differs from the selected school session/report mapping.')
    elif path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            ws=book['JKBOSE_Internal_Assessment']if'JKBOSE_Internal_Assessment'in book.sheetnames else book.active
            if ws.max_row>20000 or ws.max_column>300:raise ValueError('Excel dimensions exceed supported limits.')
            values=list(ws.values)
            if any(isinstance(v,str)and v.startswith('=')for row in values for v in row):raise ValueError('Excel contains formulas; provide an exported values workbook.')
            rows=read_table(values,cls,session,ws.title);kind='Excel awards'
            if ws.title=='JKBOSE_Internal_Assessment'and'Awards_Matrix'in book.sheetnames:
                internal=read_table(list(book['Awards_Matrix'].values),cls,session,'Awards_Matrix')
        finally:book.close()
    elif path.suffix.lower()in{'.xls','.html'}:
        parser=HtmlRows();parser.feed(path.read_text(encoding='utf-8-sig'));rows=read_table(parser.rows,cls,session,'export');kind='JKBOSE HTML Excel export'
    else:raise ValueError('Unsupported award file; choose PDF, XLSX or the JKBOSE HTML XLS export.')
    return dict(role=role,path=str(path),label=path.name,kind=kind,rows=rows,internal_rows=internal,sha256=hashlib.sha256(path.read_bytes()).hexdigest())

def issue(kind,source='',row=None,detail='',severity='error'):
    r=row or{};return dict(kind=kind,source=source,roll=r.get('roll',''),reg=r.get('reg',''),name=r.get('name',''),subject=r.get('subject',''),detail=detail,severity=severity)
def key(row):return row['roll'],row['reg'],row['subject']
def values(row):return row['marks'],row['maximum'],row['attendance']

def analyze(cards,datasets,cls,session,errors=None):
    issues=list(errors or[]);students={};expected=set();reg_rolls=defaultdict(set);roll_regs=defaultdict(set)
    for c in cards:
        pair=c['roll'],c['reg']
        if pair in students:
            if set(students[pair]['subjects'])!=set(c['subjects']):issues.append(issue('Conflicting admit cards',row=c,detail='Subject lists differ.'))
            else:issues.append(issue('Repeated admit card',row=c,detail='Counted once.',severity='warning'))
        students[pair]=c;expected.update((c['roll'],c['reg'],s)for s in c['subjects']);reg_rolls[c['reg']].add(c['roll']);roll_regs[c['roll']].add(c['reg'])
    summaries=[];indexes=[]
    for ds in datasets:
        label=ds['label'];index=defaultdict(list)
        for row in ds['rows']:
            index[key(row)].append(row);reg_rolls[row['reg']].add(row['roll']);roll_regs[row['roll']].add(row['reg'])
            mx=maximum(row['subject'],cls)
            if mx is None:issues.append(issue('Subject maximum not configured',label,row,'Enter this subject maximum for the selected class; no passing result is assumed.'))
            elif row['maximum']is None or Decimal(row['maximum'])!=mx:issues.append(issue('Maximum differs',label,row,f'Expected {mx}; file shows {row["maximum"] or "not supplied"}.'))
            if row['attendance']!='P'or row['marks']is None:issues.append(issue('Absent / incomplete award',label,row,'Present attendance and numeric marks are required for a passing award.'))
            elif mx is not None and (Decimal(row['marks'])<Decimal(mx)*Decimal('0.4')or Decimal(row['marks'])>mx):issues.append(issue('Outside passing range',label,row,f'Marks {row["marks"]}; pass {Decimal(mx)*Decimal("0.4"):g}, maximum {mx}.'))
            pair=row['roll'],row['reg']
            if pair in students and norm(students[pair]['name'])!=norm(row['name']):issues.append(issue('Name differs (secondary check)',label,row,'Admit card: '+students[pair]['name'],severity='warning'))
        for k,rr in index.items():
            if len(rr)>1:issues.append(issue('Duplicate award',label,rr[0],f'{len(rr)} entries for the same paired IDs and subject.'))
        if cards:
            absent=defaultdict(list)
            for k in expected-index.keys():absent[k[:2]].append(k[2])
            for pair,codes in absent.items():issues.append(issue('Missing student/subjects',label,dict(students[pair],subject=', '.join(sorted(codes))),'No exact paired-identifier award for these enrolled subjects.'))
            for k in index.keys()-expected:issues.append(issue('Unexpected student/subject',label,index[k][0],'Not an exact admit-card student/subject combination.'))
        if ds.get('internal_rows'):
            other={key(r):values(r)for r in ds['internal_rows']};current={k:values(v[0])for k,v in index.items()if len(v)==1}
            for k in other.keys()|current.keys():
                if other.get(k)!=current.get(k):issues.append(issue('Workbook views differ',label,dict(roll=k[0],reg=k[1],subject=k[2]),'Awards_Matrix and automation handoff do not agree.'))
        summaries.append(dict(source=label,kind=ds['kind'],students=len({k[:2]for k in index}),awards=sum(len(v)for v in index.values()),unique_awards=len(index),missing=len(expected-index.keys())if cards else None,extra=len(index.keys()-expected)if cards else None,subjects={s:sum(k[2]==s for k in index)for s in subject_catalog() if s not in{'BO','ZO'}}));indexes.append(index)
    for i,left in enumerate(indexes):
        for j in range(i+1,len(indexes)):
            right=indexes[j];label=datasets[i]['label']+' vs '+datasets[j]['label']
            for k in left.keys()|right.keys():
                if k not in left or k not in right:
                    if not cards:issues.append(issue('Source coverage differs',label,dict(roll=k[0],reg=k[1],subject=k[2]),'Award appears in only one supplied source.'))
                elif len(left[k])==len(right[k])==1 and values(left[k][0])!=values(right[k][0]):issues.append(issue('Award values differ',label,left[k][0],f'{values(left[k][0])} vs {values(right[k][0])}.'))
    for ds in datasets:
        for row in ds['rows']:
            if len(reg_rolls[row['reg']])>1 or len(roll_regs[row['roll']])>1:issues.append(issue('Conflicting identifier ownership',ds['label'],row,'Registration or exam roll is linked to a different partner. No name-only match used.'))
    unique={json.dumps(x,sort_keys=True):x for x in issues};issues=list(unique.values());limited=not cards
    return dict(schema=1,created_at=datetime.now(timezone.utc).isoformat(),student_class=cls,session=session,admit_students=len(students),expected_awards=len(expected)if cards else None,expected_subjects={s:sum(k[2]==s for k in expected)for s in subject_catalog()},subject_labels=subject_catalog(),subject_maxima={s:str(maximum(s,cls))if maximum(s,cls)is not None else None for s in subject_catalog()},sources=summaries,issues=issues,error_count=sum(x['severity']=='error'for x in issues),warning_count=sum(x['severity']=='warning'for x in issues),completeness_checked=bool(cards),scope_note='All supplied admit-card student/subject pairs checked.'if cards else'No admit cards supplied: enrolled-student/subject completeness is not verified.',outcome='Needs review'if any(x['severity']=='error'for x in issues)else'Review name/duplicate-card warnings'if issues else'All supplied evidence matches'if cards and datasets else'Limited check - no admit-card completeness verification')

def verify_files(admit_files,school_files,jkbose_files,cls,session,stop_event=None,subject_maxima=''):
    token=RULES.set(parse_rules(subject_maxima))
    try:return _verify_files(admit_files,school_files,jkbose_files,cls,session,stop_event)
    finally:RULES.reset(token)

def _verify_files(admit_files,school_files,jkbose_files,cls,session,stop_event=None):
    cls=str(cls).replace('th','');session_end(session)
    if not cls.isdigit()or not 1<=int(cls)<=12:raise ValueError('Enter a school class from 1 to 12.')
    if not school_files and not jkbose_files:raise ValueError('Select at least one award source.')
    cards=[];datasets=[];errors=[];inputs=[];seen=set()
    for role,files in [('admit',admit_files),('school',school_files),('jkbose',jkbose_files)]:
        for file in files:
            if stop_event and stop_event.is_set():raise ValueError('Verification stopped; no completed history was recorded.')
            path=Path(file).resolve(strict=True)
            if path in seen:raise ValueError('The same file is selected more than once.')
            seen.add(path)
            if path.stat().st_size>250*1024*1024:raise ValueError('Input exceeds 250 MB.')
            digest=hashlib.sha256(path.read_bytes()).hexdigest();inputs.append(dict(role=role,path=str(path),sha256=digest));print('[VERIFY FILE] '+path.name)
            try:
                if role=='admit':cards.extend(read_admits(pdf_pages(path),cls,session))
                else:datasets.append(read_awards(path,role,cls,session))
            except (ValueError,KeyError,IndexError,TypeError,InvalidOperation)as exc:errors.append(issue('Unreadable / unsupported evidence',path.name,detail=str(exc)))
            if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:raise ValueError('Input file changed during verification: '+path.name)
    if admit_files and not cards:errors.append(issue('Admit-card completeness unavailable',detail='No admit cards could be parsed.'))
    report=analyze(cards,datasets,cls,session,errors);report['inputs']=inputs
    failed_admits={x['source']for x in errors}&{Path(p).name for p in admit_files}
    if failed_admits:
        report['completeness_checked']=False;report['scope_note']='Incomplete admit-card evidence: comparisons cover only readable cards; full enrolled-student/subject completeness is not verified.'
    if not datasets:report['outcome']='Incomplete - no readable award source';report['error_count']=max(report['error_count'],1)
    return report

def local_time(value):
    return datetime.fromisoformat(value).astimezone(timezone(timedelta(hours=5,minutes=30))).strftime('%d %b %Y %H:%M IST')

def export_pdf(report,path):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='AuditSmall',fontSize=8,leading=10));styles.add(ParagraphStyle(name='AuditCell',fontSize=8,leading=10))
    def P(v,style='AuditCell'):return Paragraph(escape(str(v)).replace('\n','<br/>'),styles[style])
    story=[P('Internal Assessment - Verification Report','Title'),P(f'Class {report["student_class"]} | School session {report["session"]} | {local_time(report["created_at"])}','Normal'),Spacer(1,9),P(report['outcome'],'Heading2'),P(f'{report["admit_students"]} admit-card students; {report["expected_awards"] if report["expected_awards"]is not None else "unverified"} expected awards; {report["error_count"]} issue(s), {report["warning_count"]} warning(s).','Normal'),P(report['scope_note'],'AuditSmall'),Spacer(1,9)]
    def table(rows,widths):
        t=Table([[P(v)for v in row]for row in rows],colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDEBEF')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F7F9')]),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]));return t
    admit_names=[Path(x['path']).name for x in report.get('inputs',[])if x.get('role')=='admit']
    if admit_names:story.extend([P('Admit cards: '+'; '.join(admit_names),'AuditSmall'),Spacer(1,6)])
    rows=[['Evidence','Students','Awards','Missing','Extra']]
    for s in report['sources']:rows.append([s['source'],s['students'],s['awards'],s['missing']if s['missing']is not None else'n/a',s['extra']if s['extra']is not None else'n/a'])
    story.extend([table(rows,[285,55,55,55,55]),Spacer(1,10)])
    if report['sources']:
        subject_labels=report.get('subject_labels',SUBJECTS);subjects=[s for s in subject_labels if s not in{'BO','ZO'}];rows=[['Subject / Maximum','Expected from admit cards']+[str(i+1)for i in range(len(report['sources']))]]
        # Counts are per independent selected source, never summed across snapshots.
        for code in subjects:
            counts=[s['subjects'].get(code,0)for s in report['sources']]
            if any(counts):rows.append([subject_labels[code]+' / '+str(report.get('subject_maxima',{}).get(code,'unverified')),report.get('expected_subjects',{}).get(code,0)if report['completeness_checked']else'n/a']+counts)
        if len(report['sources'])<=5:story.extend([table(rows,[135,150]+[220/len(report['sources'])]*len(report['sources'])),P('Numbered count columns follow the evidence order above.','AuditSmall')])
    if report['issues']:
        story.append(P('Exceptions and warnings','Heading2'));rows=[['Finding / Source','Student identifiers / Subject','Details']]
        for item in report['issues']:rows.append([item['kind']+'\n'+item['source'],item['name']+'\n'+item['roll']+' / '+item['reg']+'\n'+item['subject'],item['detail']])
        story.append(table(rows,[155,165,185]))
    story.extend([Spacer(1,10),P('Matching uses exact registration + exam roll + subject; names are secondary. Pass is 40% of the configured subject maximum. Classes 11/12 use established Internal Assessment limits unless explicitly overridden; other classes require entered subject maxima. Botany and Zoology are combined once. Healthcare and IT/ITES are not required Internal Assessment subjects.','AuditSmall'),P('This is an offline file comparison, not live approval verification. No portal data or source file was changed. A matching old export does not establish current portal status.','AuditSmall')])
    def footer(canvas,doc):canvas.setFont('Helvetica',8);canvas.drawString(36,20,'Confidential examination verification');canvas.drawRightString(A4[0]-36,20,str(doc.page))
    path=Path(path);temp=path.with_name(path.stem+'.tmp.pdf');SimpleDocTemplate(str(temp),pagesize=A4,leftMargin=45,rightMargin=45,topMargin=35,bottomMargin=40).build(story,onFirstPage=footer,onLaterPages=footer);os.replace(temp,path);return str(path)

def history_root(root):
    from .internal_assessment_history import storage_root
    return storage_root(root)/'Verification'
def save_history(report,root):
    base=history_root(root);folder=base/f'Class {report["student_class"]} - {report["session"]}'/(datetime.now(timezone.utc).strftime('%Y-%m-%d_%H-%M-%S')+'-'+uuid.uuid4().hex[:8]);folder.mkdir(parents=True)
    data=folder/'analysis.json';data.write_text(json.dumps(report,indent=2),encoding='utf-8');pdf=folder/'Verification_Report.pdf';export_pdf(report,pdf)
    manifest=dict(schema=1,created_at=report['created_at'],student_class=report['student_class'],session=report['session'],outcome=report['outcome'],files={f.name:hashlib.sha256(f.read_bytes()).hexdigest()for f in[data,pdf]})
    m=folder/'manifest.json';m.write_text(json.dumps(manifest,indent=2),encoding='utf-8');latest=base/'LATEST.json';temp=base/('LATEST-'+uuid.uuid4().hex+'.tmp');temp.write_text(json.dumps(dict(manifest=str(m.relative_to(base)),sha256=hashlib.sha256(m.read_bytes()).hexdigest())),encoding='utf-8');os.replace(temp,latest);return str(folder)
def load_history(folder,root):
    folder=Path(folder).resolve(strict=True)
    if not folder.is_relative_to(history_root(root).resolve()):raise ValueError('Verification history must remain in the configured private folder.')
    data=json.loads((folder/'manifest.json').read_text(encoding='utf-8'))
    if set(data['files'])!={'analysis.json','Verification_Report.pdf'}:raise ValueError('History manifest changed.')
    for name,digest in data['files'].items():
        if hashlib.sha256((folder/name).read_bytes()).hexdigest()!=digest:raise ValueError('Verification history integrity check failed.')
    return json.loads((folder/'analysis.json').read_text(encoding='utf-8'))
def list_history(root):
    base=history_root(root);rows=[]
    if base.exists():
        for m in base.glob('Class */*/manifest.json'):
            data=json.loads(m.read_text(encoding='utf-8'));rows.append(dict(data,folder=str(m.parent)))
    return sorted(rows,key=lambda r:r['created_at'],reverse=True)
def save_report_copy(folder,destination,root):
    report=load_history(folder,root);source=Path(folder)/'Verification_Report.pdf';dest=Path(destination).resolve()
    if dest in {Path(item['path']).resolve()for item in report.get('inputs',[])}:raise ValueError('Choose a different filename: a verification report cannot overwrite source evidence.')
    if dest==source.resolve():return str(dest)
    if dest.is_relative_to(history_root(root).resolve()):raise ValueError('Save the copy outside the immutable verification history.')
    temp=dest.with_name(dest.stem+'.verification.tmp.pdf');shutil.copyfile(source,temp);os.replace(temp,dest);return str(dest)
