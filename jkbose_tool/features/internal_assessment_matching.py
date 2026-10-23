"""Read-only paired-identity reconciliation; incomplete pages are never called matched."""
from collections import Counter
from datetime import datetime,timezone,timedelta
from pathlib import Path
from decimal import Decimal
import hashlib
import json
import os
import re
from xml.sax.saxutils import escape
from .internal_assessment import norm,number,live_maximum,submission_issues,rejected_correction


def compare_rosters(plan,rows,*,complete,official_maximum,prior,coverage_reason,source_kind='pending'):
    roster_mode=source_kind=='students_report'
    official_text=str(official_maximum) if official_maximum is not None else ''
    def manual_issues(award):
        issues=submission_issues(award)
        if official_maximum is None:
            if not roster_mode: issues.append('Official JKBOSE maximum is unavailable on an empty pending queue; submission is not verified.')
            return issues
        try:
            if award.attendance=='A':
                if number(plan.maximum)!=official_maximum: issues.append('Excel and official maxima differ; manual review required.')
                return issues
            mark=number(award.marks)
            if mark>official_maximum: issues.append('Exceeds official JKBOSE maximum; manual only.')
            if mark<official_maximum*Decimal('0.4'): issues.append('Below official 40% pass threshold; manual only.')
            if number(plan.maximum)!=official_maximum: issues.append('Excel and official maxima differ; manual review required.')
        except ValueError: issues.append('Absent/missing/invalid marks; manual only.')
        return issues
    source={award.regno:award for award in plan.awards}
    source_exams={award.exam_roll:award for award in plan.awards}
    reg_counts=Counter(row.get('regno','') for row in rows)
    exam_counts=Counter(str(row.get('exam_roll','')).strip() for row in rows)
    seen_source=set(); records=[]
    for row in rows:
        reg=str(row.get('regno','')).strip(); exam=str(row.get('exam_roll','')).strip()
        award=source.get(reg) or source_exams.get(exam)
        if award: seen_source.add(award.regno)
        issues=(manual_issues(award) if award else [])+row.get('report_issues',[])
        reason=''
        if not re.fullmatch(r'[A-Za-z0-9-]+',reg) or not re.fullmatch(r'\d{9}',exam) or reg_counts[reg]!=1 or exam_counts[exam]!=1:
            status='identifier_conflict'; reason='Missing, invalid or duplicated live registration/exam roll.'
        elif award is None:
            status='jkbose_only'; reason=('This Students report registration/exam-roll pair is absent from the selected Excel subject.' if roster_mode else 'This displayed JKBOSE registration/exam-roll pair is absent from the selected Excel subject.')
        elif award.regno!=reg or award.exam_roll!=exam:
            status='identifier_mismatch'; reason='Registration and exam roll do not match as the same pair.'
        elif norm(award.name)!=norm(row.get('name','')):
            status='name_mismatch'; reason='Paired identifiers match but student names need manual review.'
        elif not roster_mode and reg in prior['verified']|prior['manual']:
            status='saved_pair_reappeared'; reason='A previously verified saved pair reappeared in the pending queue.'
        else: status='matched_roster' if roster_mode else 'matched_pending'
        if not roster_mode and (row.get('attendance') not in {'P','A'} or (row.get('attendance')=='A' and (award is None or award.attendance!='A') and not rejected_correction(plan,award,row))): issues.append('Live attendance conflicts with the source; manual review.')
        if reason: issues.append(reason)
        records.append({'subject':plan.subject,'excel_registration':award.regno if award else '',
                        'excel_exam_roll':award.exam_roll if award else '','excel_name':award.name if award else '',
                        'jkbose_registration':reg,'jkbose_exam_roll':exam,'jkbose_name':str(row.get('name','')),
                        'match_status':status,'reason':reason,'marks':award.marks if award else '',
                        'excel_maximum':plan.maximum,'official_maximum':official_text,
                        'jkbose_attendance':row.get('attendance',''),'manual_issues':list(dict.fromkeys(issues)),
                        'award_action':'manual_only' if issues else 'roster_check_only' if roster_mode else 'eligible_after_batch_validation'})
    for award in plan.awards:
        if award.regno in seen_source: continue
        if roster_mode:
            status='excel_only'; reason='This Excel registration/exam pair is absent from the complete Students report for this subject and selected report session. Manual review required.'
        elif award.regno in prior['verified']|prior['manual']:
            status='verified_submitted'; reason='Unchanged paired award has an intact account-scoped saved receipt in history.'
        elif complete:
            status='not_in_pending_unverified'; reason='Not in the complete pending queue. It may have been submitted outside this tool; verify the official saved list.'
        else:
            status='not_seen_unread_pages'; reason='Not seen on the scanned pages; later pending pages are unread. This does not prove the student is missing.'
        issues=manual_issues(award)
        records.append({'subject':plan.subject,'excel_registration':award.regno,'excel_exam_roll':award.exam_roll,
                        'excel_name':award.name,'jkbose_registration':'','jkbose_exam_roll':'','jkbose_name':'',
                        'match_status':status,'reason':reason,'marks':award.marks,'excel_maximum':plan.maximum,
                        'official_maximum':official_text,'jkbose_attendance':'','manual_issues':issues,
                        'award_action':'already_verified_saved' if status=='verified_submitted' else 'manual_only' if issues else 'not_verified_for_submission'})
    counts=Counter(record['match_status'] for record in records)
    reconciled=complete and all(record['match_status'] in ({'matched_roster'} if roster_mode else {'matched_pending','verified_submitted'}) for record in records)
    return {'subject':plan.subject,'coverage_complete':complete,'coverage_reason':coverage_reason,
            'identity_roster_reconciled':reconciled,'excel_count':len(plan.awards),'scanned_jkbose_count':len(rows),
            'official_maximum':official_text,'counts':dict(counts),'records':records}


def _read_only_next(driver,next_page):
    from selenium.webdriver.common.by import By
    # Only an actual ASP.NET grid Page$ navigation anchor is accepted. Never a submit/button.
    candidates=[]
    for link in driver.find_elements(By.CSS_SELECTOR,'a[href]'):
        href=link.get_attribute('href') or ''
        match=re.fullmatch(r"javascript:__doPostBack\(['\"]([^'\"]*gridStudents[^'\"]*)['\"],\s*['\"]Page\$(Next|\d+)['\"]\);?",href,re.I)
        if match and match[2] in {'Next',str(next_page)} and not link.get_attribute('onclick') and link.is_displayed() and link.is_enabled(): candidates.append(link)
    if len(candidates)>1: raise ValueError('Multiple independent pending-page controls; stop for inspection.')
    return candidates[0] if candidates else None


def scan_subject(driver,plan,audit,school_account,prior,stop_event=None,absence_callback=None,interface_only=False):
    from . import internal_assessment_browser as browser
    from ..browser import page_change_token
    rows=[]; snapshots=[]; complete=False; reason='Not scanned'; first=None; expected_page=1
    def persist(failure=''):
        official=live_maximum(first) if first and first['rows'] else None
        data=compare_rosters(plan,rows,complete=complete,official_maximum=official,prior=prior,coverage_reason=reason)
        data.update(expected_pending_count=first['total'] if first else None,pages_scanned=len(snapshots),
                    check_scope='interface_and_pending_count' if interface_only else 'pending_roster',
                    available_subjects=first.get('available_subjects',[]) if first else [],
                    page_evidence=snapshots,source_sha256=plan.source_sha256,source=plan.source,
                    class_name=plan.student_class,session=plan.session,school_account=school_account,failure=failure,
                    session_note='The marks page has no academic-session selector; operator confirmation is required.')
        path=audit.folder/'matching.json'; temporary=audit.folder/'matching.tmp'
        with temporary.open('w',encoding='utf-8') as stream:
            json.dump(data,stream,indent=2); stream.flush(); os.fsync(stream.fileno())
        os.replace(temporary,path)
        summary={key:value for key,value in data.items() if key not in {'records','page_evidence','source'}}
        return path,data,summary
    try:
        while True:
            browser._stop(stop_event)
            state=browser.snapshot(driver,plan,school_account,require_first=expected_page==1)
            if first is None: first=state
            elif state['total']!=first['total'] or live_maximum(state)!=live_maximum(first):
                raise ValueError('Pending roster or maximum changed during read-only matching. Coverage cannot be certified.')
            if state.get('page',1) not in (expected_page,None) and state['total']:
                raise ValueError('Read-only navigation did not reach the next consecutive pending page.')
            evidence=browser._evidence(driver,audit,'matching-pending',state)
            snapshots.append({'page':state.get('page',expected_page),'rows':len(state['rows']),'evidence':evidence})
            rows.extend(state['rows'])
            if absence_callback:
                for row in state['rows']:
                    if row.get('attendance')=='A' and not rejected_correction(plan,next((a for a in plan.awards if a.regno==row['regno']),None),row):
                        absence_callback({'registration':row['regno'],'exam_roll':row.get('exam_roll',''),'subject':plan.subject,
                                          'reason':'Marked absent on the live JKBOSE pending page'},evidence)
            if state['total']==0 or len(rows)==state['total']:
                complete=True; reason='Every currently pending JKBOSE row was read without entering or submitting marks.'; break
            if len(rows)>state['total']: raise ValueError('Read-only pending rows exceed the declared total.')
            if interface_only:
                reason='Full identity matching uses the Students report. This assessment check records interface, maximum and pending count only; later batches are checked before saving.'; break
            link=_read_only_next(driver,expected_page+1)
            if link is None:
                reason='Incomplete: JKBOSE exposes only the displayed pending rows; Submit & Next writes awards and was never used to browse.'; break
            before=browser._begin_read(driver,link.get_attribute('id') or '', 'Pending roster next page'); browser._stop(stop_event); link.click()
            browser._wait_page(driver,before,stop_event); expected_page+=1
        path,data,summary=persist()
        if first and first['rows']:
            official=live_maximum(first)
            audit.append('official_maximum_checked',maximum=str(official),pass_marks=str(official*Decimal('0.4')),
                         excel_maximum=plan.maximum,mismatch=official!=Decimal(plan.maximum),evidence=snapshots[0]['evidence'])
        audit.append('read_only_matching',report=str(path),summary=summary)
        print(f'[MATCHING] {plan.subject}: {len(rows)}/{first["total"]} pending rows read; '
              f'{"complete pending coverage" if complete else "INCOMPLETE coverage"}. No marks entered or submitted.')
        return {'subject':plan.subject,'status':'checked_read_only' if interface_only or data['identity_roster_reconciled'] else 'matching_review_required',
                'new_verified_count':0,'known_saved_count':len(prior['verified']|prior['manual']),
                'pending_count':first['total'],'history':str(audit.folder),'matching_report':str(path),'matching_summary':summary}
    except Exception as error:
        complete=False; reason='Read-only matching interrupted: '+str(error)
        path,data,summary=persist(str(error)); error.matching_report=str(path); error.matching_summary=summary
        audit.append('read_only_matching_interrupted',report=str(path),reason=str(error)[:1000]); raise


def export_matching_reports(folder,class_report,holds=()):
    """Application exports: exact text identifiers, complete coverage notes and manual holds."""
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,PageBreak
    folder=Path(folder); manifest=class_report['manifest']; details=[]; summaries=[]
    roster_mode=manifest.get('matching_source')=='students_report'
    selected=class_report['subjects']
    export_results=selected+[result for result in class_report.get('safety_preflight',[]) if result['subject'] not in {item['subject'] for item in selected}]
    for result in export_results:
        path=result.get('matching_report')
        from .internal_assessment_history import resolve_reference
        # Old signed reports keep their original absolute paths after reorganization.
        history_root = next((parent for parent in folder.parents if (parent / '.internal-assessment-references').is_file()), folder)
        data=json.loads(resolve_reference(path,history_root).read_text(encoding='utf-8')) if path else None
        summary=data or {'subject':result['subject'],'coverage_complete':False,'identity_roster_reconciled':False,
                         'excel_count':0,'scanned_jkbose_count':0,'expected_pending_count':None,'counts':{},
                         'coverage_reason':result.get('reason','Subject was not scanned.'),'records':[]}
        summaries.append(summary)
        for record in summary['records']:
            record=dict(record); record['manual_issues']=list(record['manual_issues'])
            reg=record['excel_registration'] or record['jkbose_registration']; exam=record['excel_exam_roll'] or record['jkbose_exam_roll']
            affected=[hold['subject'] for hold in holds if hold['registration']==reg or (hold['exam_roll'] and hold['exam_roll']==exam)]
            if affected:
                record['manual_issues'].append('Class-wide absence/identity hold from '+', '.join(sorted(set(affected)))+'. All subjects manual-only.')
                record['award_action']='manual_only'
            details.append(record)
    generated=datetime.now(timezone(timedelta(hours=5,minutes=30))).strftime('%d %b %Y %H:%M IST')
    book=Workbook(); summary_sheet=book.active; summary_sheet.title='Matching Summary'
    summary_sheet.append(['INTERNAL ASSESSMENT - CONFIDENTIAL MATCHING REPORT'])
    for row in [['Class',manifest['class'],'Session',manifest['session']],['School account',manifest['school_account'],'Generated',generated],
                ['Source SHA-256',manifest['source_sha256']],['Source workbook',Path(manifest['source']).name],
                ['Run status',class_report['status']],['Important','Saved awards disappear from pending lists. Unseen records are not automatically missing or submitted.'],
                ['Session check',f'Workbook {manifest["session"]} -> Students report {manifest["report_session"]}. Enrolment only; report contains no practical attendance/marks.' if roster_mode else 'JKBOSE marks page exposes no academic-session selector; operator must verify active examination cycle.'],[]]: summary_sheet.append(row)
    summary_sheet.append(['Subject','Excel awards','Report subject total' if roster_mode else 'JKBOSE pending total','Report rows read' if roster_mode else 'Pending rows read','Complete report coverage' if roster_mode else 'Complete pending coverage','Identity roster reconciled','Matched enrolment' if roster_mode else 'Matched pending','Verified submitted','JKBOSE only','Mismatches','Excel only' if roster_mode else 'Unverified Excel','Coverage / review note'])
    for data in summaries:
        counts=data['counts']; mismatch=sum(counts.get(key,0) for key in ['identifier_conflict','identifier_mismatch','name_mismatch','saved_pair_reappeared'])
        summary_sheet.append([data['subject'],data['excel_count'],data.get('expected_pending_count'),data['scanned_jkbose_count'],
                              'Yes' if data['coverage_complete'] else 'NO - INCOMPLETE','Yes' if data['identity_roster_reconciled'] else 'Needs review',
                              counts.get('matched_roster' if roster_mode else 'matched_pending',0),counts.get('verified_submitted',0),counts.get('jkbose_only',0),mismatch,
                              counts.get('excel_only',0) if roster_mode else counts.get('not_in_pending_unverified',0)+counts.get('not_seen_unread_pages',0),data['coverage_reason']])
        if not data['coverage_complete']:
            summary_sheet.cell(summary_sheet.max_row,5).font=Font(color='B91C1C',bold=True)
            summary_sheet.cell(summary_sheet.max_row,5).fill=PatternFill('solid',fgColor='FFF0F0')
    detail_sheet=book.create_sheet('Paired Student Matching')
    keys=['subject','excel_registration','excel_exam_roll','excel_name','jkbose_registration','jkbose_exam_roll','jkbose_name',
          'match_status','award_action','marks','excel_maximum','official_maximum','jkbose_attendance','reason','manual_issues']
    headers=['Subject','Excel Registration','Excel Exam Roll','Excel Name','JKBOSE Registration','JKBOSE Exam Roll','JKBOSE Name',
             'Matching status','Award action','Source marks','Excel maximum','Official maximum','JKBOSE attendance','Matching reason','Manual-only reasons']
    if roster_mode:
        keys+=['pending_observation','report_subjects','report_approval','report_exam_fee']
        headers+=['Practical pending observation','Students report subjects','Board report approval','Exam fee']
    detail_sheet.append(headers)
    for record in details: detail_sheet.append(['; '.join(record.get(key,[])) if isinstance(record.get(key),list) else record.get(key,'') for key in keys])
    holds_sheet=book.create_sheet('Class-wide Holds'); holds_sheet.append(['Registration','Exam roll','Absent / conflicting subject','Reason'])
    for hold in holds: holds_sheet.append([hold['registration'],hold['exam_roll'],hold['subject'],hold['reason']])
    skipped=class_report.get('skipped_applications',[])
    if roster_mode:
        skipped_sheet=book.create_sheet('Skipped Applications')
        skipped_sheet.append(['Registration','Exam roll','Student name','Report class','Reason'])
        for row in skipped: skipped_sheet.append([row['regno'],row['exam_roll'],row['name'],row['class_name'],row['skip_reason']])
    for sheet in book:
        sheet.freeze_panes='A11' if sheet is summary_sheet else 'A2'
        header_row=10 if sheet is summary_sheet else 1
        for cell in sheet[header_row]: cell.fill=PatternFill('solid',fgColor='16324F'); cell.font=Font(color='FFFFFF',bold=True)
        sheet.auto_filter.ref=f'A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        for row in sheet:
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,str): cell.data_type='s'; cell.number_format='@'
        for column in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(column)].width=24
        sheet.row_dimensions[header_row].height=44
    summary_sheet.column_dimensions['L'].width=85; summary_sheet.column_dimensions['B'].width=26
    summary_sheet.row_dimensions[7].height=42; summary_sheet.row_dimensions[8].height=42
    for column in ['N','O']: detail_sheet.column_dimensions[column].width=75
    if roster_mode:
        detail_sheet.column_dimensions['P'].width=75; detail_sheet.column_dimensions['Q'].width=85
    for column in ['D','G']: detail_sheet.column_dimensions[column].width=28
    holds_sheet.column_dimensions['D'].width=75
    if roster_mode: skipped_sheet.column_dimensions['E'].width=75
    for row_number in [4,5,6,7,8]:
        summary_sheet.merge_cells(start_row=row_number,start_column=2,end_row=row_number,end_column=12)
        summary_sheet.row_dimensions[row_number].height=30
    for sheet,header_row in [(sheet,10 if sheet is summary_sheet else 1) for sheet in book]:
        for row in sheet.iter_rows(min_row=header_row+1):
            lines=max((sum(max(1,(len(part)+int(sheet.column_dimensions[cell.column_letter].width)-1)//int(sheet.column_dimensions[cell.column_letter].width)) for part in str(cell.value or '').split('\n')) for cell in row),default=1)
            sheet.row_dimensions[row[0].row].height=max(28,min(220,lines*14+8))
    for row in detail_sheet.iter_rows(min_row=2):
        for column in [9,10,11]:
            cell=row[column]
            if re.fullmatch(r'\d+(?:\.\d+)?',str(cell.value or '')):
                cell.value=float(cell.value); cell.number_format='0.##########'
    xlsx=folder/'matching.xlsx'; temporary=folder/'matching.tmp.xlsx'; book.save(temporary); os.replace(temporary,xlsx)
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='MatchSmall',fontName='Helvetica',fontSize=7.2,leading=9))
    styles.add(ParagraphStyle(name='MatchTiny',fontName='Helvetica',fontSize=6.5,leading=8))
    styles.add(ParagraphStyle(name='MatchHeader',fontName='Helvetica-Bold',fontSize=7,leading=9,textColor=colors.white))
    def para(value,style='MatchSmall'): return Paragraph(escape(str(value if value is not None else '')).replace('\n','<br/>'),styles[style])
    def table(data,widths,header=True):
        item=LongTable(data,colWidths=widths,repeatRows=1 if header else 0,hAlign='LEFT')
        commands=[('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#CCD6E0')),
                  ('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]
        if header: commands.append(('BACKGROUND',(0,0),(-1,0),colors.HexColor('#16324F')))
        item.setStyle(TableStyle(commands)); return item
    story=[para('Internal Assessment - Confidential Matching Report','Title'),Spacer(1,8),
           para(f'Class {manifest["class"]} | Session {manifest["session"]} | School {manifest["school_account"]} | {generated}'),
           para('Run status: '+class_report['status']+' | Mode: '+manifest['mode']),
           *([para(f'Students report session: {manifest["report_session"]} (separate from workbook session {manifest["session"]})'),
              para('Complete enrolment matching does not certify the practical pending queue. This report has no practical attendance, marks or saved-award status; existing submission and absence safeguards still apply.')] if roster_mode else []),
           para('Source: '+Path(manifest['source']).name),para('Source SHA-256: '+manifest['source_sha256']),Spacer(1,8),
           para('Coverage must be complete before a full roster reconciliation can be claimed. Submit & Next was never used to browse. Saved awards disappear from pending lists; an unseen student may already have been submitted outside this tool.'),
           para('The marks page has no academic-session selector. The operator must verify the active examination cycle. Matching identity does not authorize absent, failed or over-maximum awards.'),Spacer(1,10)]
    header=['Subject','Excel','Report' if roster_mode else 'Pending','Read','Report coverage' if roster_mode else 'Coverage','Identity result','Review note']
    summary_rows=[[para(value,'MatchHeader') for value in header]]
    for data in summaries:
        summary_rows.append([para(value) for value in [data['subject'],data['excel_count'],data.get('expected_pending_count',''),data['scanned_jkbose_count'],
                            'COMPLETE' if data['coverage_complete'] else 'INCOMPLETE','Reconciled' if data['identity_roster_reconciled'] else 'Needs review',data['coverage_reason']]])
    story.append(table(summary_rows,[90,35,45,35,75,80,433.89]))
    story.extend([Spacer(1,10),para(f'Class-wide holds: {len(holds)}. A known absence in any subject blocks automated submission for that student in every subject.')])
    if roster_mode:
        story.extend([Spacer(1,8),para(f'Applications skipped because no exam roll is assigned: {len(skipped)}. These are excluded from matching; another valid application is still checked using both identifiers.')])
        if skipped:
            skipped_rows=[[para(value,'MatchHeader') for value in ['Registration','Exam roll','Student name','Why excluded from matching']]]
            skipped_rows.extend([[para(row['regno']),para(row['exam_roll']),para(row['name']),para(row['skip_reason'])] for row in skipped])
            story.extend([Spacer(1,6),table(skipped_rows,[110,85,160,438.89])])
    if holds:
        story.extend([PageBreak(),para('Class-wide absence and identity holds','Heading2')])
        hold_rows=[[para(value,'MatchHeader') for value in ['Registration','Exam roll','Subject requiring manual handling','Reason - every subject is blocked for automation']]]
        hold_rows.extend([[para(hold['registration']),para(hold['exam_roll']),para(hold['subject']),para(hold['reason'])] for hold in holds])
        story.append(table(hold_rows,[110,85,160,438.89]))
    if details:
        story.extend([PageBreak(),para('Paired registration and exam-roll matching','Heading2')])
        rows=[[para(value,'MatchHeader') for value in ['Subject','Excel Reg / Exam','JKBOSE Reg / Exam','Student name','Matching / action','Reason and manual-only flags']]]
        for record in details:
            rows.append([para(record['subject']),para(record['excel_registration']+'\n'+record['excel_exam_roll'],'MatchTiny'),
                         para(record['jkbose_registration']+'\n'+record['jkbose_exam_roll'],'MatchTiny'),
                         para(record['excel_name']+'\n'+record['jkbose_name'],'MatchTiny'),
                         para(record['match_status']+'\n'+record['award_action'],'MatchTiny'),
                         para(record['reason']+'\n'+'; '.join(record['manual_issues'])+('\n'+record.get('pending_observation','') if roster_mode else ''),'MatchTiny')])
        story.append(table(rows,[75,110,110,120,125,253.89]))
    def footer(canvas,document):
        canvas.saveState(); canvas.setFont('Helvetica',7); canvas.setFillColor(colors.HexColor('#66788A'))
        canvas.drawString(24,12,'Confidential institutional matching report - not a saved-marks receipt')
        canvas.drawRightString(817.89,12,f'Page {document.page}'); canvas.restoreState()
    pdf=folder/'matching.pdf'; temporary=folder/'matching.tmp.pdf'
    # Platypus adds 6 pt of frame padding on each side: 18 + 6 = 24 pt.
    SimpleDocTemplate(str(temporary),pagesize=landscape(A4),leftMargin=18,rightMargin=18,topMargin=24,bottomMargin=26,
                      title='Internal Assessment - Confidential Matching Report',author='School Examination Administration').build(story,onFirstPage=footer,onLaterPages=footer)
    os.replace(temporary,pdf)
    return {'excel':str(xlsx),'pdf':str(pdf),
            'excel_sha256':hashlib.sha256(xlsx.read_bytes()).hexdigest(),
            'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest()}
