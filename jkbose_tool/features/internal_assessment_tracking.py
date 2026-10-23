"""Student/subject status projection from source awards and verified local evidence."""
from pathlib import Path
from datetime import datetime,timezone,timedelta
import copy,hashlib,json
from .internal_assessment import Audit,number,norm,subject_name,submission_issues,APPROVAL_HOLD_PREFIX
from .internal_assessment_batch import load_jobs
from .internal_assessment_history import storage_root,manifest_paths
from .internal_assessment_saved_report import _report_index,_award_matches

def tracking_payload(source,history_root,student_class,session,account):
    from openpyxl import load_workbook
    source=Path(source).resolve(strict=True);root=storage_root(history_root)
    jobs=load_jobs(source,student_class=student_class,session=session,all_subjects=True)
    plans={job['subject']:job['plan'] for job in jobs if job['plan']}
    from .internal_assessment_policy import CohortAbsencePolicy,apply_holds
    scope={'school_account':account.lower(),'class':student_class,'session':session}
    key=hashlib.sha256(json.dumps(scope,sort_keys=True).encode()).hexdigest()[:24]
    if (root/'_cohort_holds'/key/'events.jsonl').is_file():
        holds=CohortAbsencePolicy(root,account,student_class,session).holds
        for plan in plans.values():apply_holds(plan,holds)
    candidates=[]
    for path in (root/'_class_runs').rglob('official_saved*.json'):
        data=json.loads(path.read_text(encoding='utf-8'))
        if data.get('student_class')==student_class and data.get('workbook_session')==session and data.get('account','').lower()==account.lower():
            _report_index(data,account);candidates.append((path.stat().st_mtime,path,data))
    latest=max(candidates,key=lambda item:item[0]) if candidates else None
    from .internal_assessment_approval import tracking_receipts
    approved_requests,uncertain_approvals=tracking_receipts(root,account,student_class,session)
    official=_report_index(latest[2],account)[0] if latest else {}
    checked_at=datetime.fromtimestamp(latest[0],timezone(timedelta(hours=5,minutes=30))).strftime('%d %b %Y %H:%M') if latest else ''
    uncertain={};saved={}
    for path in manifest_paths(root):
        manifest=json.loads(path.read_text(encoding='utf-8'))
        if manifest.get('student_class')!=student_class or manifest.get('session')!=session:continue
        audit,plan=Audit.existing(path.parent)
        accounts={e['school_account'].lower() for e in audit.events if e['kind']=='run_started' and e.get('school_account')}
        if accounts!={account.lower()}:continue
        for event in audit.events:
            if event['kind']=='submit_attempt' and event['batch_id'] in audit.unresolved():uncertain.setdefault(plan.subject,set()).update(event['registrations'])
        for reg in audit.verified()|audit.manual_saved():
            old=next((a for a in plan.awards if a.regno==reg),None)
            if old:saved[(plan.subject,reg)]=(plan,old)
    statuses={};summary=[];confirmed_rows=[]
    for subject,plan in plans.items():
        counts={key:0 for key in ['Submitted','Requested','Approved','Pending','Blocked','Uncertain','Rejected']}
        for award in plan.awards:
            award_key=(award.regno,award.exam_roll,subject)
            row=official.get(award_key);note='; '.join(award.issues);row_checked_at=checked_at
            if getattr(award,'identity_confirmation',None):
                note=('Identity warning superseded by recorded user confirmation; paired IDs verified. '+note).strip()
                confirmed_rows.extend(award.source_rows)
            receipt=approved_requests.get(award_key)
            if receipt and latest and datetime.fromisoformat(receipt['verified_time']).timestamp()<latest[0]:receipt=None
            comparison=copy.copy(award);comparison.issues=[]
            if receipt and _award_matches(plan,comparison,dict(receipt,status='Approved',rejection_reason='')):
                status='Submitted';board='Approved';approval='Approved (manually verified)' if receipt['verified_method']=='manual' else 'Approved (tool verified)'
                evidence='Principal administrator verification' if receipt['verified_method']=='manual' else 'Principal approval history; exact request removal';counts['Approved']+=1
                row_checked_at=datetime.fromisoformat(receipt['verified_time']).astimezone(timezone(timedelta(hours=5,minutes=30))).strftime('%d %b %Y %H:%M')
            elif row and _award_matches(plan,comparison,row):
                status='Submitted';board=row['status'];approval='Pending principal approval' if board=='Requested' else 'Approved';evidence='Official JKBOSE saved report'
                counts[board]+=1
            elif row and row['status'] in {'Requested','Approved'} and not row.get('rejection_reason') and norm(row['name'])==norm(award.name):
                status='Submitted';board=row['status'];approval='Source differs; review only';evidence='Official JKBOSE saved report';counts[board]+=1
                note=(('Already recorded; no retry. Status-only report does not expose saved marks/attendance; current source values remain separate. ' if row.get('detail')=='status_only' else 'Already recorded; no retry. JKBOSE: '+str(row['marks'])+'/'+str(row['maximum'])+' '+row['attendance']+'; current export: '+str(award.marks or 'AB')+'/'+str(plan.maximum)+' '+award.attendance+'. ')+note).strip()
                if row.get('detail')=='status_only':approval='Status confirmed; saved values not exposed'
            elif row and row['status']=='Rejected':
                status='Rejected';board='Rejected';approval='Not approved';evidence='Official JKBOSE saved report'
                note='JKBOSE rejection: '+(row['rejection_reason'] or 'No reason supplied')+'. '+note
            elif row:
                status='Uncertain';board=row['status'];approval='';evidence='Official JKBOSE saved report';note='Saved record differs from source or attendance; review. '+note
            elif award.regno in uncertain.get(subject,set()):
                status='Uncertain';board='';approval='';evidence='Unresolved submission journal'
            elif submission_issues(award):
                status='Blocked';board='';approval='';evidence='Source validation / class-wide hold'
            elif (subject,award.regno) in saved:
                old_plan,old=saved[(subject,award.regno)]
                if (old.exam_roll==award.exam_roll and norm(old.name)==norm(award.name)
                        and old.attendance==award.attendance and (old.marks==award.marks if award.attendance=='A' else number(old.marks)==number(award.marks)) and number(old_plan.maximum)==number(plan.maximum)):
                    status='Submitted';board='Not refreshed';approval='Not refreshed';evidence='Verified submission history'
                else:status='Uncertain';board='';approval='';evidence='Changed source versus saved history'
            else:status='Pending';board='';approval='';evidence='Not found in verified saved evidence'
            if award.attendance=='A' or any(issue.startswith(APPROVAL_HOLD_PREFIX) for issue in award.issues):
                approval='Manual Principal approval only'
            counts[status]+=1
            if award_key in uncertain_approvals:
                approval='Uncertain approval; verify JKBOSE';note=('Principal approval attempt requires official reconciliation. '+note).strip()
            statuses[(subject,award.regno)]=[status,board,approval,row_checked_at,evidence,note]
        summary.append([subject,len(plan.awards),*[counts[k] for k in ['Submitted','Requested','Approved','Pending','Blocked','Uncertain','Rejected']]])
    book=load_workbook(source,read_only=True,data_only=False,keep_links=True)
    ws=book['JKBOSE_Internal_Assessment'];values=list(ws.values);book.close()
    headers=list(values[0]);flags=[]
    for row in values[1:]:
        subject=subject_name(row[4]);subject='Biology' if subject in {'Botany','Zoology'} else subject
        flags.append(statuses.get((subject,str(row[5] or '').strip()),['Uncertain','','',checked_at,'No safe status match','Source row identity/subject could not be matched.']))
    return dict(source=str(source),source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),student_class=student_class,session=session,account=account,
                status_rows=[dict(subject=k[0],registration=k[1],values=v) for k,v in statuses.items()],confirmed_rows=sorted(set(confirmed_rows)),checked_at=checked_at,checked_iso=datetime.fromtimestamp(latest[0],timezone.utc).isoformat() if latest else '',official_report=str(latest[1]) if latest else '',summary=summary,headers=headers,flags=flags,
                output=str(source.with_name(source.stem+'_Submission_Tracking.xlsx')))

def refresh_tracking_copy(source,history_root,student_class,session,account):
    """Application export. Keep the selected source immutable; replace only its sidecar."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from .internal_assessment_source_versions import record_source_version,add_source_changes_sheet
    version=record_source_version(source,history_root,student_class,session,account)
    payload=tracking_payload(source,history_root,student_class,session,account)
    book=load_workbook(source,keep_links=True)
    ws=book['JKBOSE_Internal_Assessment'];start=ws.max_column+1
    labels=['Submission Status','JKBOSE Status','Principal Approval','Verified At','Evidence','Review Note']
    for offset,label in enumerate(labels):
        cell=ws.cell(1,start+offset,label);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='334155');cell.alignment=Alignment(wrap_text=True)
        ws.column_dimensions[cell.column_letter].width=[22,18,28,23,42,75][offset]
    for index,flags in enumerate(payload['flags'],2):
        for offset,value in enumerate(flags):ws.cell(index,start+offset,value)
    issue_column=next((c.column for c in ws[1] if c.value=='Source Issues'),None)
    if issue_column:
        for row_number in payload['confirmed_rows']:
            cell=ws.cell(row_number,issue_column)
            if cell.value=='Award registration differs from roster; review identity.':cell.value=''
    overview=book.create_sheet('Submission_History');overview.append(['Internal assessment submission history'])
    overview.append([f'Class {student_class}, {session}. Verified at {payload["checked_at"]}. Requested is submitted, awaiting separate principal approval.'])
    overview.append(['Biology counts one award per student; Botany and Zoology source rows share its status.'])
    overview.append(['Subject','Awards','Submitted','Requested','Approved','Pending','Blocked','Uncertain','Rejected'])
    for row in payload['summary']:overview.append(row)
    for cell in overview[4]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='334155')
    overview.column_dimensions['A'].width=30
    for col in 'BCDEFGHI':overview.column_dimensions[col].width=14
    overview.freeze_panes='B5';overview.auto_filter.ref=f'A4:I{overview.max_row}'
    add_source_changes_sheet(book,version,{(r['subject'],r['registration']):r['values'] for r in payload['status_rows']})
    dest=Path(payload['output']);temp=dest.with_name(dest.stem+'.tmp.xlsx')
    book.save(temp);book.close();temp.replace(dest)
    if hashlib.sha256(Path(source).read_bytes()).hexdigest()!=payload['source_sha256']:raise ValueError('Source workbook changed while writing the separate tracker.')
    return str(dest)
