"""Class-wide Internal Assessment selection, verified resume and private run summaries."""
from pathlib import Path
from datetime import datetime, timezone
from decimal import Decimal
import csv
import hashlib
import json
import os
import uuid
from .internal_assessment import (
    SUBJECTS, Audit, RunLock, load_plan, class_key, session_key, subject_name,
    norm, number, prepare_history_root,
)


def workbook_subjects(excel_file, sheet='Auto', student_class='11', session='', combine_biology=True):
    """Discover only subjects actually present in this cohort's completed workbook."""
    from openpyxl import load_workbook
    source = Path(excel_file).resolve(strict=True)
    if source.suffix.lower() != '.xlsx' or source.stat().st_size > 20 * 1024 * 1024:
        raise ValueError('Choose an .xlsx awards workbook no larger than 20 MB.')
    cls, session = class_key(student_class), session_key(session)
    book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        chosen = sheet
        if chosen in ('Auto', '', None):
            chosen = 'JKBOSE_Internal_Assessment' if 'JKBOSE_Internal_Assessment' in book.sheetnames else 'Awards_Matrix'
        if chosen not in book.sheetnames: raise ValueError('Worksheet not found.')
        ws = book[chosen]
        if ws.max_row > 20000 or ws.max_column > 300: raise ValueError('Workbook dimensions exceed supported bounds.')
        rows = list(ws.iter_rows())
        header = next((i for i, row in enumerate(rows[:40])
                       if any(norm(c.value) in {'regno','boardregno','boardregistrationnumber','registrationno','registrationnumber'} for c in row)
                       and any(norm(c.value) in {'studentname','studentsname','name'} for c in row)), None)
        if header is None: raise ValueError('Completed awards headers were not found.')
        labels = [norm(c.value) for c in rows[header]]
        def column(options, required=True):
            matches = [i for i, label in enumerate(labels) if label in options]
            if len(matches) > 1 or (required and not matches): raise ValueError('Missing or ambiguous awards metadata columns.')
            return matches[0] if matches else None
        subject_column = column({'subjectcode'}, False)
        if subject_column is None: subject_column = column({'subject','subjectname'}, False)
        found = set()
        if subject_column is not None:
            class_column, session_column = column({'class'}), column({'session'})
            evaluation_column = column({'evaluationtype'})
            reg_column = column({'regno','boardregno','boardregistrationnumber','registrationno','registrationnumber'})
            name_column = column({'studentname','studentsname','name'})
            for row in rows[header+1:]:
                if row[reg_column].value is None and row[name_column].value is None: continue
                for i in [subject_column,class_column,session_column,evaluation_column]:
                    if row[i].data_type == 'f': raise ValueError('Formula in required cohort/subject metadata.')
                if class_key(row[class_column].value) != cls or session_key(row[session_column].value) != session:
                    raise ValueError('Workbook class/session mismatch; select one cohort at a time.')
                if norm(row[evaluation_column].value) not in {'internal','internalassessment','ia'}:
                    raise ValueError('Only Internal Assessment awards may be selected.')
                label = subject_name(row[subject_column].value)
                if label not in SUBJECTS.values(): raise ValueError('Unsupported source subject: ' + label)
                found.add(label)
        else:
            reg_column = column({'regno','boardregno','boardregistrationnumber','registrationno','registrationnumber'})
            name_column = column({'studentname','studentsname','name'})
            preamble = ' '.join(str(c.value or '') for row in rows[:header] for c in row)
            import re
            match = re.search(r'class\s*(10|11|12)(?:th)?', preamble, re.I)
            if not match or match[1] != cls or session_key(preamble) != session or not re.search(r'internal\s+assessment',preamble,re.I):
                raise ValueError('Matrix must identify Internal Assessment and the selected class/session.')
            for i, cell in enumerate(rows[header]):
                label = subject_name(cell.value)
                if label in SUBJECTS.values() and any(str(row[i].value or '').strip().lower() not in {'','x'}
                                                     for row in rows[header+1:] if row[reg_column].value is not None or row[name_column].value is not None):
                    found.add(label)
            # A matrix's BI column is a derived display total when split parts exist.
            if found & {'Botany','Zoology'}: found.discard('Biology')
        if combine_biology and found & {'Botany','Zoology'}:
            found.difference_update({'Botany','Zoology'}); found.add('Biology')
        if not found: raise ValueError('No completed Internal Assessment subjects were found for this cohort.')
        return [label for label in SUBJECTS.values() if label in found]
    finally:
        book.close()


def load_jobs(excel_file, sheet='Auto', student_class='11', session='', subjects=None,
              all_subjects=False, maximum='Auto', biology_component_max=None, combine_biology=True):
    available = workbook_subjects(excel_file,sheet,student_class,session,combine_biology)
    requested = available if all_subjects else [subject_name(value) for value in (subjects or [])]
    if combine_biology: requested = ['Biology' if value in {'Botany','Zoology'} else value for value in requested]
    requested = list(dict.fromkeys(requested))
    if not requested: raise ValueError('Select at least one subject, or choose All workbook subjects.')
    if len(requested) > 1 and norm(maximum) not in {'','auto'}:
        raise ValueError('Multiple subjects require Auto maxima; their official maxima can differ.')
    from .internal_assessment_policy import source_absences,apply_holds
    holds=source_absences(excel_file,sheet)
    jobs = []
    for subject in requested:
        plan, error = None, ''
        try:
            if subject not in available: raise ValueError('Subject is not present in the selected cohort workbook.')
            plan = load_plan(excel_file,sheet,student_class,session,subject,maximum,
                             biology_component_max=biology_component_max,combine_biology=combine_biology)
            apply_holds(plan,holds)
        except ValueError as exc: error = str(exc)
        jobs.append({'subject':subject,'plan':plan,'error':error,'sheet':plan.sheet if plan else sheet})
    return jobs


def _award_signature(plan, award):
    from .internal_assessment import submission_issues
    def numeric(value): return 'AB' if str(value).upper()=='AB' else str(number(value).normalize())
    maxima = plan.component_maxima or ({label:plan.biology_component_max for label in ['Botany','Zoology']}
              if plan.subject == 'Biology' and plan.combine_biology else {plan.subject:plan.maximum})
    return {'regno':award.regno,'name':norm(award.name),'exam_roll':award.exam_roll,
            'marks':'AB' if award.attendance=='A' else numeric(award.marks),'attendance':award.attendance,'maximum':numeric(plan.maximum),'issues':submission_issues(award),
            'components':{k:numeric(v) for k,v in award.components.items()},
            'component_maxima':{k:numeric(v) for k,v in maxima.items()}}


def verified_history(root, plan, school_account):
    """Only the latest intact same-account outcome may suppress a new save."""
    from .internal_assessment_browser import _receipt_history, ChangedSavedAwardError
    from .internal_assessment import submission_issues
    result={'verified':set(),'manual':set(),'proofs':[]};current={a.regno:a for a in plan.awards};cache={}
    for reg,(folder,event) in _receipt_history(root,plan,school_account).items():
        if reg not in current or event['kind']=='portal_award_rejected':continue
        if folder not in cache:
            audit,old_plan=Audit.existing(folder);cache[folder]=(old_plan,{a.regno:a for a in old_plan.awards})
        old_plan,awards=cache[folder]
        if event['kind']=='official_recorded_award':
            row=event['saved_records'][reg]
            if plan.source_sha256!=old_plan.source_sha256 or row['regno']!=current[reg].regno or row['exam_roll']!=current[reg].exam_roll or norm(row['name'])!=norm(current[reg].name):raise ChangedSavedAwardError('A previously recorded award needs fresh official-report comparison; never retry it blindly.')
            result['verified'].add(reg);result['proofs'].append(dict(registration=reg,history=str(folder),event_hash=event['hash'],kind='recorded_different',saved_record=row));continue
        try:matches=not submission_issues(current[reg]) and _award_signature(plan,current[reg])==_award_signature(old_plan,awards[reg])
        except (ValueError,KeyError):matches=False
        if not matches:raise ChangedSavedAwardError('A previously saved award changed in the source. Manual reconciliation is required; it will not be resubmitted.')
        kind='manual' if event['kind']=='manual_reconciliation' else 'verified';result[kind].add(reg)
        result['proofs'].append(dict(registration=reg,history=str(folder),event_hash=event['hash'],kind=kind))
    return result


class ClassRunAudit:
    def __init__(self, root, *, source, source_sha256, student_class, session, subjects, school_account, mode):
        self.folder = Path(root)/'_class_runs'/(datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'-'+uuid.uuid4().hex[:12])
        self.folder.mkdir(parents=True,exist_ok=False)
        self.previous='0'*64
        self.manifest={'source':source,'source_sha256':source_sha256,'class':student_class,'session':session,
                       'subjects':subjects,'school_account':school_account,'mode':mode}
        self._write_json('manifest.json',self.manifest)
        self.results=[{'subject':subject,'status':'not_started'} for subject in subjects]
        self.status='running'; self.append('class_run_started',**self.manifest); self.report()

    def _write_json(self,filename,value):
        temporary=self.folder/(filename+'.tmp')
        with temporary.open('w',encoding='utf-8') as stream:
            json.dump(value,stream,indent=2); stream.flush(); os.fsync(stream.fileno())
        os.replace(temporary,self.folder/filename)

    def append(self,kind,**data):
        event={'time':datetime.now(timezone.utc).isoformat(),'kind':kind,'previous':self.previous,**data}
        event['hash']=Audit._hash(event)
        with (self.folder/'events.jsonl').open('a',encoding='utf-8') as stream:
            stream.write(json.dumps(event)+'\n'); stream.flush(); os.fsync(stream.fileno())
        self.previous=event['hash']

    def update(self,index,result):
        self.results[index]=result; self.append('subject_outcome',**result); self.report()

    def report(self):
        report={'manifest':self.manifest,'status':self.status,'subjects':self.results,
                'class_wide_holds':getattr(self,'holds',[]),
                'safety_preflight':getattr(self,'preflight',[]),
                'matching_exports':getattr(self,'exports',{}),
                'matching_export_error':getattr(self,'export_error',''),
                'student_report_matching':getattr(self,'student_report_matching',{}),
                'tracking_workbook':getattr(self,'tracking_workbook',''),
                'tracking_workbook_error':getattr(self,'tracking_error',''),
                'new_verified_saved_count':sum(result.get('new_verified_count',0) for result in self.results),
                'note':'Read-only/empty queues never prove that source records were submitted. See each subject journal for evidence.'}
        self._write_json('report.json',report)
        temporary=self.folder/'report.csv.tmp'
        with temporary.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream); writer.writerow(['Subject','Status','New verified saves','Previously verified awards','Pending count','Subject history','Review reason'])
            for result in self.results:
                values=[result.get(key,'') for key in ['subject','status','new_verified_count','known_saved_count','pending_count','history','reason']]
                writer.writerow(["'"+str(value) if str(value).startswith(('=','+','-','@','\t','\r')) else value for value in values])
            stream.flush(); os.fsync(stream.fileno())
        os.replace(temporary,self.folder/'report.csv')
        return report


def run_jobs(jobs, excel_file, history_root, school_account, *, student_class, session,
             execute=False, preview_only=False, confirm_session=False, confirm_callback=None,
             on_driver_ready=None, stop_event=None, driver_factory=None, login_callback=None,
             report_session=None, report_reader=None, verify_saved=False, saved_report_reader=None, automatic_submit=False):
    from . import internal_assessment_browser as browser
    from ..browser import get_driver, login_only
    if not jobs or len({job['subject'] for job in jobs}) != len(jobs): raise ValueError('Choose distinct subjects for this class run.')
    cls, session = class_key(student_class), session_key(session)
    source=Path(excel_file).resolve(strict=True); source_hash=hashlib.sha256(source.read_bytes()).hexdigest()
    for job in jobs:
        plan=job['plan']
        if plan and (plan.student_class != cls or plan.session != session or Path(plan.source) != source or plan.source_sha256 != source_hash or plan.subject != job['subject']):
            raise ValueError('Selected plans no longer match this class/session/source. Reopen the preview.')
    if not preview_only and (not confirm_session or not browser.re_account(school_account)):
        raise ValueError('Confirm the active session and school account before a class run.')
    history_root=prepare_history_root(history_root)
    driver=None
    with RunLock(history_root):
        from .internal_assessment_policy import source_absences,apply_holds,CohortAbsencePolicy
        valid_plans=[job['plan'] for job in jobs if job['plan']]
        source_sheet=valid_plans[0].sheet if valid_plans else jobs[0].get('sheet','Auto')
        raw_holds=source_absences(source,source_sheet)
        policy=CohortAbsencePolicy(history_root,school_account,cls,session) if not preview_only else None
        if policy:
            policy.add(raw_holds,origin='consolidated_source',evidence=source_hash)
            corrections=policy.reconcile_source(source,source_sheet,current_holds=raw_holds)
            if corrections:print(f'[SOURCE CORRECTED] Released older source-only absence holds for {len({h["registration"] for h in corrections})} students using the current explicit Present handoff. Original history retained.')
        holds=policy.holds if policy else raw_holds
        for plan in valid_plans: apply_holds(plan,holds)
        match_check=None
        if automatic_submit and execute:
            if not verify_saved or report_session is None or len(valid_plans)!=len(jobs):
                raise ValueError('Automatic Submit requires valid selected subjects and complete live report preflight.')
            from .internal_assessment_match_check import load_match_check
            try:match_check=load_match_check(history_root,valid_plans,school_account,report_session)
            except ValueError:
                print('[SUBMIT PREFLIGHT] No reusable Match: complete report and identity checks will run before any Save.')
        audit=ClassRunAudit(history_root,source=str(source),source_sha256=source_hash,student_class=cls,session=session,
                            subjects=[job['subject'] for job in jobs],school_account=school_account,
                            mode='offline_preview' if preview_only else 'submit' if execute else 'read_only')
        audit.holds=holds
        student_report=None; prior_by_subject={}
        def ensure_driver():
            nonlocal driver
            if driver is None:
                driver=(driver_factory or get_driver)()
                if on_driver_ready: on_driver_ready(driver)
                (login_callback or login_only)(driver)
        def live_absence(hold,evidence):
            policy.add([hold],origin='live_jkbose_read_only',evidence=evidence)
            for selected_plan in valid_plans: apply_holds(selected_plan,policy.holds)
            audit.append('class_wide_live_hold',hold=hold,evidence=evidence)
        def final_batch_receipt(plan,awards,batch_id):
            from .internal_assessment_saved_report import read_saved_report,verify_saved_batch
            saved_report=(saved_report_reader or read_saved_report)(driver,cls,school_account,session,report_session or 'Auto',stop_event)
            receipt_path=audit.folder/('official_saved_after_'+batch_id+'.json')
            receipt_path.write_text(json.dumps(saved_report,indent=2),encoding='utf-8')
            statuses=verify_saved_batch(plan,awards,school_account,saved_report)
            audit.append('final_batch_receipt_checked',subject=plan.subject,batch_id=batch_id,evidence=str(receipt_path),sha256=saved_report['sha256'])
            return dict(evidence=str(receipt_path),official_report_sha256=saved_report['sha256'],official_statuses=statuses,receipt_detail=saved_report.get('receipt_detail','detailed'))
        try:
            if not preview_only:
                if verify_saved and valid_plans:
                    needs_receipts=not execute
                    for selected_plan in valid_plans:
                        try:
                            previous=verified_history(history_root,selected_plan,school_account)
                            needs_receipts=needs_receipts or any(a.regno not in previous['verified']|previous['manual'] for a in selected_plan.awards)
                        except (browser.UncertainSubmissionError,browser.ChangedSavedAwardError):
                            needs_receipts=True
                    if needs_receipts:
                        from .internal_assessment_saved_report import read_saved_report,recognise_saved_awards
                        ensure_driver()
                        saved_report=(saved_report_reader or read_saved_report)(driver,cls,school_account,session,report_session or 'Auto',stop_event)
                        receipt_path=audit.folder/'official_saved_assessment_report.json'
                        receipt_path.write_text(json.dumps(saved_report,indent=2),encoding='utf-8')
                        recovered=recognise_saved_awards(history_root,valid_plans,school_account,saved_report,receipt_path)
                        audit.append('official_saved_report_checked',evidence=str(receipt_path),sha256=saved_report['sha256'],records=saved_report['total'],recognised=recovered,
                                     meaning='Requested/Approved paired identities recorded on JKBOSE. A status-only report does not verify saved marks; Principal approval remains separate.')
                        print(f'[SAVED REPORT] {saved_report["total"]} official saved awards read across {saved_report["pages"]} pages; {sum(recovered.values())} selected awards newly recognised/recovered automatically. Requested is recorded, awaiting approval.')
                for index,job in enumerate(jobs):
                    browser._stop(stop_event)
                    if job['plan'] is not None: prior_by_subject[job['subject']]=verified_history(history_root,job['plan'],school_account)
                needs_selected_submission=any(job['plan'] and any(award.regno not in prior_by_subject[job['subject']]['verified']|prior_by_subject[job['subject']]['manual'] for award in job['plan'].awards) for job in jobs)
                if report_session is not None and valid_plans and (not execute or needs_selected_submission):
                    from .internal_assessment_students_report import read_student_report,write_snapshot,match_student_report
                    ensure_driver()
                    if match_check:
                        from .internal_assessment_students_report import validate_snapshot
                        student_report=match_check['snapshot']
                        validate_snapshot(student_report,cls,school_account,student_report['report_session'])
                        audit.append('students_report_reused_from_match',match_check=match_check['path'],sha256=match_check['sha256'])
                        print('[MATCH REUSED] Validated Students Report reused from 2. Match read only. Live batch checks remain active.')
                    else:
                        student_report=(report_reader or read_student_report)(driver,cls,school_account,session,report_session,stop_event)
                    audit.student_report_matching=write_snapshot(audit.folder,student_report)
                    audit.append('students_report_read',**audit.student_report_matching)
                    print(f'[STUDENTS REPORT] Class {cls}: {audit.student_report_matching["student_count"]} exam-roll applications; {audit.student_report_matching["skipped_no_exam_roll_count"]} without exam roll skipped; workbook {session} -> board report {student_report["report_session"]}. Enrolment only; no practical attendance/saved marks.')
                    # Persist the complete report/Excel reconciliation before
                    # navigating to marks or recording any submission intent.
                    from .internal_assessment_students_report import export_report_matching
                    roster_check=export_report_matching(audit,jobs,student_report,prior_by_subject,[],holds)
                    audit.student_report_matching.update(roster_check)
                    audit.append('students_report_preflight_exported',**roster_check['exports'],status=roster_check['status'])
                    if execute:
                        conflicts=[plan.subject for plan in valid_plans if not match_student_report(plan,student_report,prior_by_subject[plan.subject])['identity_roster_reconciled']]
                        if conflicts: raise ValueError('Students report enrolment/identity mismatch. No marks entered: '+', '.join(conflicts)+'. Review the private matching reports.')
            preflight={}
            if execute and not preview_only:
                # Complete enrolment matching comes from Students Report. In that
                # workflow inspect only selected, unfinished assessment subjects.
                # All-source/history absence holds were loaded above and stay active.
                needs_submission=any(job['plan'] and any(award.regno not in prior_by_subject[job['subject']]['verified']|prior_by_subject[job['subject']]['manual'] for award in job['plan'].awards) for job in jobs)
                if student_report is not None:
                    inspection_jobs=[job for job in jobs if job['plan'] and any(
                        award.regno not in prior_by_subject[job['subject']]['verified']|prior_by_subject[job['subject']]['manual']
                        for award in job['plan'].awards)]
                    audit.append('selected_subject_preflight_scope',subjects=[job['subject'] for job in inspection_jobs],
                                 meaning='Complete Students report identity matching; assessment interface/count checks only for selected unfinished subjects. Known all-source/history holds remain active.')
                else:
                    # Legacy direct callers without a Students Report retain the
                    # broader live scan; they must not claim report reconciliation.
                    inspection_jobs=load_jobs(source,source_sheet,cls,session,all_subjects=True,
                                              combine_biology=valid_plans[0].combine_biology) if needs_submission else []
                incomplete=[]
                for inspection in inspection_jobs:
                    browser._stop(stop_event); plan=inspection['plan']
                    if plan is None: raise ValueError('Cannot verify class-wide attendance: '+inspection['subject']+': '+inspection['error'])
                    apply_holds(plan,holds)
                    if verify_saved and 'saved_report' in locals():
                        from .internal_assessment_saved_report import annotate_rejected_corrections
                        annotate_rejected_corrections(plan,saved_report,school_account)
                    ensure_driver()
                    checked=browser.run(driver,plan,history_root,school_account,confirm_callback,stop_event,
                                        submit=False,confirm_session=confirm_session,lock_held=True,absence_callback=live_absence,interface_only=student_report is not None)
                    preflight[plan.subject]=checked
                    audit.append('subject_read_only_preflight',subject=plan.subject,matching_report=checked.get('matching_report'),status=checked['status'])
                    if not checked.get('matching_summary',{}).get('coverage_complete',False): incomplete.append(plan.subject)
                if incomplete:
                    audit.append('batchwise_submission_policy',unread_subjects=incomplete,
                                 meaning='Displayed-batch validation. Unread pages are not certified; known source/history/live holds remain mandatory.',automatic_after_match=bool(match_check))
                    print('[BATCH MODE] Later pending pages are unread. Every displayed batch will be validated; known absence/history holds remain active. '+('The approved run preflight authorises valid saves without further dialogs.' if automatic_submit else 'Each valid batch requires confirmation.'))
                for plan in valid_plans:
                    apply_holds(plan,holds)
                    prior_by_subject[plan.subject]=verified_history(history_root,plan,school_account)
            for index,job in enumerate(jobs):
                browser._stop(stop_event)
                if hashlib.sha256(source.read_bytes()).hexdigest()!=source_hash: raise ValueError('Workbook changed during the class run. Stop and reopen the preview.')
                plan=job['plan']
                if plan is None:
                    audit.update(index,{'subject':job['subject'],'status':'source_blocked','reason':job['error']}); continue
                if preview_only:
                    subject_audit=Audit(history_root,plan)
                    subject_audit.append('offline_preview',source=plan.manifest(),awards=plan.rows())
                    subject_audit.report(plan)
                    result={'subject':plan.subject,'status':'offline_preview','history':str(subject_audit.folder),
                            'manual_required_count':sum(bool(award.issues) for award in plan.awards)}
                else:
                    prior=prior_by_subject[plan.subject]
                    if (execute or verify_saved) and all(award.regno in prior['verified']|prior['manual'] for award in plan.awards):
                        result=browser.skip_verified_subject(plan,history_root,school_account,prior)
                    else:
                        ensure_driver()
                        result=browser.run(driver,plan,history_root,school_account,confirm_callback,stop_event,
                                           submit=execute,confirm_session=confirm_session,lock_held=True,absence_callback=live_absence,interface_only=student_report is not None,
                                           receipt_callback=final_batch_receipt if verify_saved else None,
                                           automatic_submit=automatic_submit and execute,
                                           match_evidence=(match_check['path'] if match_check else audit.student_report_matching.get('snapshot','')) if automatic_submit else '')
                        if execute and plan.subject in preflight:
                            for key in ['matching_report','matching_summary']:
                                if key in preflight[plan.subject]: result[key]=preflight[plan.subject][key]
                audit.update(index,result)
                if result['status']=='cancelled': audit.status='cancelled'; break
            else: audit.status='finished_with_manual_review' if any(result['status'] in {'source_blocked','manual_required','matching_review_required'} for result in audit.results) else 'finished'
        except browser.AssessmentStopped:
            audit.status='stopped'; audit.append('class_run_stopped'); raise
        except Exception as error:
            audit.status='failed'
            if 'index' in locals():
                result={'subject':jobs[index]['subject'],'status':'failed','reason':str(error)[:1000]}
                for key in ['matching_report','matching_summary']:
                    if hasattr(error,key): result[key]=getattr(error,key)
                audit.update(index,result)
            audit.append('class_run_failed',error_type=type(error).__name__,reason=str(error)[:1000]); raise
        finally:
            roster_error=None; original_failure=audit.status in {'failed','stopped'}
            if 'preflight' in locals():
                audit.preflight=list(preflight.values())
                for result in audit.results:
                    checked=preflight.get(result['subject'],{})
                    for key in ['matching_report','matching_summary']:
                        if key not in result and key in checked: result[key]=checked[key]
            if student_report is not None:
                from .internal_assessment_students_report import export_report_matching
                try:
                    matched=export_report_matching(audit,jobs,student_report,prior_by_subject,audit.results,holds)
                    audit.student_report_matching.update(matched)
                    audit.append('students_report_matching_exported',**matched['exports'],status=matched['status'])
                    if matched['status']=='roster_review_required' and audit.status=='finished': audit.status='finished_with_manual_review'
                    print(f'[ROSTER EXCEL] {matched["exports"]["excel"]}\n[ROSTER PDF] {matched["exports"]["pdf"]}')
                except Exception as error:
                    roster_error=error; audit.export_error='Students report export: '+str(error)[:900]
                    audit.append('students_report_export_failed',reason=str(error)[:1000])
                    if not original_failure: audit.status='failed'
            audit.append('class_run_ended',status=audit.status)
            report=audit.report()
            if driver is not None:
                try: driver.quit()
                except Exception: pass
            if not preview_only:
                from .internal_assessment_matching import export_matching_reports
                try:
                    exports=export_matching_reports(audit.folder,report,holds)
                    audit.exports=exports
                    audit.append('matching_reports_exported',**exports)
                    report=audit.report()
                    print(f'[MATCHING EXCEL] {exports["excel"]}\n[MATCHING PDF] {exports["pdf"]}')
                except Exception as export_error:
                    audit.export_error=str(export_error)[:1000]
                    audit.append('matching_report_export_failed',reason=str(export_error)[:1000])
                    if audit.status not in {'failed','stopped'}:
                        audit.status='failed'; audit.report(); raise
                    report=audit.report()
            if not preview_only:
                try:
                    from .internal_assessment_absence_register import refresh_absence_register
                    audit.absence_register=refresh_absence_register(source,history_root,cls,session,school_account)
                    audit.append('absence_register_refreshed',path=audit.absence_register)
                    print('[ABSENCE REGISTER] '+audit.absence_register)
                except Exception as register_error:
                    audit.append('absence_register_not_refreshed',reason=str(register_error)[:1000])
                    print('[ABSENCE REGISTER NOT REFRESHED] '+str(register_error)+'. Durable submission history retained.')
            if not preview_only and verify_saved:
                try:
                    from .internal_assessment_tracking import refresh_tracking_copy
                    audit.tracking_workbook=refresh_tracking_copy(source,history_root,cls,session,school_account)
                    audit.append('tracking_workbook_refreshed',path=audit.tracking_workbook)
                    print(f'[TRACKING EXCEL] {audit.tracking_workbook}')
                except Exception as tracking_error:
                    audit.tracking_error=str(tracking_error)[:1000]
                    audit.append('tracking_workbook_not_refreshed',reason=audit.tracking_error)
                    print('[TRACKING NOT REFRESHED] '+audit.tracking_error+'. Original workbook and submission history retained.')
                report=audit.report()
            from .internal_assessment_history import publish_history
            try:
                history_index = publish_history(history_root)
                print(f'[OPEN HISTORY] {history_index}')
            except Exception as view_error:
                audit.append('history_view_failed', reason=str(view_error)[:1000])
                audit.report()
                print(f'[HISTORY VIEW ERROR] {view_error}. Original evidence retained in {audit.folder}')
            if roster_error is not None and not original_failure: raise roster_error
        if (not execute and not preview_only and verify_saved and student_report is not None
                and len(valid_plans)==len(jobs) and audit.status in {'finished','finished_with_manual_review'}
                and audit.student_report_matching.get('status')=='roster_reconciled'):
            from .internal_assessment_match_check import save_match_check
            match_path=save_match_check(history_root,valid_plans,school_account,report_session,audit.folder,audit.student_report_matching['snapshot'])
            print(f'[MATCH READY] Automatic Submit is authorised for unchanged selections for 24 hours. Match evidence: {match_path}')
        return report
