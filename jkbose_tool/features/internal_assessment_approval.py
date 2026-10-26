"""Principal approval: exact pending-request matching, single actions and durable history."""
from pathlib import Path
from datetime import datetime,timezone
from html.parser import HTMLParser
from urllib.parse import urlparse
from collections import Counter
import hashlib,json,os,re,time
from selenium.common.exceptions import ElementClickInterceptedException
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select,WebDriverWait
from .internal_assessment import Audit,RunLock,prepare_history_root,norm,number,subject_name,match_batch,SUBJECTS,class_key,session_key
from .internal_assessment_postback import begin_read,wait_read
from .internal_assessment_students_report import open_report_from_sidebar,is_report_page,read_student_report,match_student_report,write_snapshot,enrolled_in_subject,has_assigned_exam_roll

URL='https://jkbose.jk.gov.in/utilities/auth_pages/generatePracticalMarks.aspx'
HEADERS=['Class','Roll Number','Registration Number','Name','Subject','Given Marks','Maximum Marks','Requested On','Action']
EXPORT_HEADERS=['ID','Class','Roll Number','Registration Number','Name','Subject','Present_Absent','Given_marks','Maximum Marks','Requested On']
SNAPSHOT=r"""
const grid=document.getElementById('bodypH_gridRR');
return {account:/Hello\s+([a-z0-9_-]+)/i.exec(document.body.innerText)?.[1]||'',
class:document.getElementById('bodypH_ddlClass')?.value,
headers:grid?[...grid.rows[0].cells].map(c=>c.innerText.trim()):[],
rows:grid?[...grid.rows].slice(1).filter(r=>r.cells.length===9).map(r=>({cells:[...r.cells].slice(0,8).map(c=>c.innerText.trim()),
approve:r.querySelector('input[id*="_btnApprove_"]')?.id||'',reject:r.querySelector('input[id*="_btnReject_"]')?.id||''})):[],
page:grid?[...grid.querySelectorAll('table span')].map(e=>e.innerText.trim()):[],
pager:grid?[...grid.querySelectorAll('a[href]')].map(e=>({text:e.innerText.trim(),href:e.getAttribute('href')})):[],
total:document.getElementById('bodypH_lblTotal')?.innerText.trim()||'',
empty_text:grid?.innerText.trim()||'',
approve_visible:!!document.getElementById('approvePopup')&&getComputedStyle(document.getElementById('approvePopup')).display!=='none',
reject_visible:!!document.getElementById('rejectPopup')&&getComputedStyle(document.getElementById('rejectPopup')).display!=='none',
approve_id:document.getElementById('bodypH_hfApproveId')?.value||''};
"""

def fingerprint(value):return hashlib.sha256(json.dumps(value,sort_keys=True).encode()).hexdigest()
def key(row):return row['regno'],row['exam_roll'],row['subject']
def accounts(school,principal):
    match=re.fullmatch(r'sch(\d{6})',school.lower())
    if not match or principal.lower()!='prin'+match[1]:raise ValueError('School and Principal accounts must identify the same school (schNNNNNN / prinNNNNNN).')

class TableParser(HTMLParser):
    def __init__(self):super().__init__(convert_charrefs=True);self.rows=[];self.row=None;self.cell=None
    def handle_starttag(self,tag,attrs):
        if tag=='tr':self.row=[]
        elif tag in {'td','th'} and self.row is not None:self.cell=[]
    def handle_data(self,data):
        if self.cell is not None:self.cell.append(data)
    def handle_endtag(self,tag):
        if tag in {'td','th'} and self.cell is not None:
            self.row.append(''.join(self.cell).strip());self.cell=None
        elif tag=='tr' and self.row is not None:
            if self.row:self.rows.append(self.row)
            self.row=None

def parse_export(path,cls,total):
    path=Path(path)
    if path.stat().st_size>20*1024*1024:raise ValueError('Principal export exceeds the supported size.')
    parser=TableParser();parser.feed(path.read_text(encoding='utf-8-sig'))
    if not parser.rows or parser.rows[0]!=EXPORT_HEADERS:raise ValueError('Principal Excel export headers changed; no approval allowed.')
    rows=parser.rows[1:]
    # The server can emit the header alone for an empty pending queue.
    if len(rows)!=total or any(len(r)!=10 for r in rows):raise ValueError('Principal export is incomplete or changed while downloading.')
    result=[]
    for r in rows:
        if not re.fullmatch(r'\d+',r[0]) or r[1]!=cls or not re.fullmatch(r'[A-Za-z0-9-]+',r[3]) or not re.fullmatch(r'\d{9}',r[2]):raise ValueError('Invalid Principal request identity/class.')
        result.append(dict(id=r[0],class_name=r[1],exam_roll=r[2],regno=r[3],name=r[4],subject=subject_name(r[5]),
                           attendance='P' if norm(r[6]) in {'p','present'} else 'A' if norm(r[6]) in {'a','ab','absent'} else '?',
                           marks=r[7],maximum=r[8],requested_on=r[9]))
    if (len({r['id'] for r in result})!=len(result) or len({key(r) for r in result})!=len(result)
            or len({(r['regno'],r['subject']) for r in result})!=len(result) or len({(r['exam_roll'],r['subject']) for r in result})!=len(result)):
        raise ValueError('Duplicate/conflicting Principal request IDs or paired student identities.')
    return result

def validate_grid(data,cls,principal):
    if data.get('account','').lower()!=principal.lower() or data.get('class')!=cls:raise ValueError('Principal account/class changed. No approval allowed.')
    if data.get('total','')=='' and data.get('empty_text')=='There are no pending requests to display.' and not data.get('rows') and data.get('headers')==['There are no pending requests to display.']:
        data['total']='0'
    if not re.fullmatch(r'\d+',data.get('total','')):raise ValueError('Principal pending total is unreadable.')
    total=int(data['total'])
    if total>20000 or (total and data.get('headers')!=HEADERS):raise ValueError('Principal grid columns/count changed.')
    rows=data.get('rows',[])
    if (total and (not rows or len(rows)>10)) or (not total and rows):raise ValueError('Principal grid count and rows disagree.')
    if 0<total<=10 and not data.get('page') and not data.get('pager') and len(rows)==total:data['page']=['1']
    if total and (len(data.get('page',[]))!=1 or not re.fullmatch(r'\d+',data['page'][0])):raise ValueError('Principal current page is unreadable.')
    if total and not 1<=int(data['page'][0])<=(total+9)//10:raise ValueError('Principal page exceeds the pending count.')
    if total and len(rows)!=min(10,total-(int(data['page'][0])-1)*10):raise ValueError('Principal grid does not contain the complete displayed page.')
    for r in rows:
        if len(r.get('cells',[]))!=8 or r['cells'][0]!=cls or not re.fullmatch(r'bodypH_gridRR_btnApprove_\d+',r.get('approve','')):raise ValueError('Principal grid row or decision control changed.')
    return total

def grid_row_matches(raw,request):
    c=raw['cells']
    return (c[:3]==[request['class_name'],request['exam_roll'],request['regno']] and norm(c[3])==norm(request['name'])
            and subject_name(c[4])==request['subject'] and (norm(c[5]) in {'','a','ab','abs','absent','0'} and norm(request['marks']) in {'','a','ab','abs','absent','0'} if request['attendance']=='A' else number(c[5])==number(request['marks'])) and number(c[6])==number(request['maximum']) and c[7]==request['requested_on'])

def match_requests(plans,rows,roster_checks=None):
    by_subject={p.subject:p for p in plans};result=[]
    for row in rows:
        if row['subject'] not in by_subject:continue
        plan=by_subject[row['subject']]
        live=dict(regno=row['regno'],exam_roll=row['exam_roll'],name=row['name'],subject=row['subject'],attendance=row['attendance'],max=row['maximum'],value=row['marks'],min='0',step='any')
        _,issues=match_batch(plan,{'rows':[live],'displayed_maximum':row['maximum']})
        if row['class_name']!=plan.student_class:issues.append('Request class differs from the source cohort.')
        if roster_checks is not None:
            record=roster_checks.get(plan.subject,{}).get((row['regno'],row['exam_roll']))
            if not record or record['match_status']!='matched_roster':issues.append('Request identity/subject does not match the complete Students report.')
            elif record.get('manual_issues'):issues.extend(record['manual_issues'])
        issues=list(dict.fromkeys(issues))
        result.append(dict(request=row,eligible=not issues,reasons=issues))
    return result

APPROVAL_MODES={'excel','class_all','class_except'}

def approval_confirmation_text(data):
    scope=data['scope'];ready=sum(e['eligible'] for e in data['rows']);held=len(data['rows'])-ready
    return (f'Approve all {ready} eligible awards in Class {scope["class"]}, session {scope["session"]}?\n'
            f'{held} excluded or held awards will remain untouched.\n\n'
            'This action cannot be undone by this software. It will confirm each eligible award on JKBOSE.\n'
            'Continue only after reviewing the complete preview and the class/session.')

def parse_exclusions(value):
    """Exact paired identifiers only; never turn a partial pair into a broad skip."""
    if isinstance(value,str):value=[line for line in value.splitlines() if line.strip()]
    pairs=set()
    for entry in value or []:
        parts=re.split(r'[,;:\t ]+',entry.strip()) if isinstance(entry,str) else list(entry)
        if len(parts)!=2:raise ValueError('Each exclusion needs registration number and nine-digit exam roll, separated by a comma.')
        reg,exam=(str(v).strip() for v in parts)
        if not re.fullmatch(r'[A-Za-z0-9-]+',reg) or not re.fullmatch(r'\d{9}',exam):raise ValueError('Invalid exclusion pair: use registration number and nine-digit exam roll.')
        if any((r==reg and e!=exam) or (e==exam and r!=reg) for r,e in pairs):raise ValueError('Conflicting exclusion identities; correct both identifiers before continuing.')
        pairs.add((reg,exam))
    return pairs

def exclusion_students(path,cls,session,school,principal):
    """Read local fetched requests or an awards workbook for a skip-only picker."""
    path=Path(path);result=[]
    if path.suffix.lower()=='.json':
        data=json.loads(path.read_text(encoding='utf-8-sig'))
        if data.get('scope')!=dict(school=school.lower(),principal=principal.lower(),**{'class':class_key(cls)},session=session_key(session)):
            raise ValueError('Fetched approval file belongs to a different account, class or session.')
        result=[e['request'] for e in data.get('requests',[])]
    else:
        if path.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=False)
            try:
                sheet=book['JKBOSE_Internal_Assessment'] if 'JKBOSE_Internal_Assessment' in book.sheetnames else book.active
                rows=list(sheet.values)
            finally:book.close()
        elif path.suffix.lower() in {'.xls','.html'}:
            parser=TableParser();parser.feed(path.read_text(encoding='utf-8-sig'));rows=parser.rows
        else:raise ValueError('Choose a fetched approval report JSON or an Excel awards/requests file.')
        if not rows:raise ValueError('Student file is empty.')
        labels=[norm(v) for v in rows[0]]
        def column(*names):return next((i for i,v in enumerate(labels) if v in names),None)
        reg=column('boardregno','registrationnumber');exam=column('examrollno','rollnumber');name=column('studentname','name');cohort=column('class');sess=column('session')
        if None in (reg,exam,name,cohort):raise ValueError('Student file needs class, registration, exam roll and name columns.')
        for row in rows[1:]:
            if row[cohort] is None:continue
            if class_key(row[cohort])!=class_key(cls):continue
            if sess is not None and session_key(row[sess])!=session_key(session):continue
            result.append(dict(regno=str(row[reg] or '').strip(),exam_roll=str(row[exam] or '').strip(),name=str(row[name] or '')))
    students={}
    for row in result:
        pair=next(iter(parse_exclusions([(row['regno'],row['exam_roll'])])))
        if pair in students and norm(students[pair]['name'])!=norm(row['name']):raise ValueError('Conflicting names for an exclusion identity.')
        students[pair]=dict(regno=pair[0],exam_roll=pair[1],name=row['name'])
    parse_exclusions(list(students))
    if not students:raise ValueError('No students for this class/session were found in the selected file.')
    return sorted(students.values(),key=lambda r:(r['exam_roll'],r['regno']))

def select_requests(plans,rows,roster_checks,*,mode,exclusions,roster,holds):
    if not rows:return []
    if mode=='excel':result=match_requests(plans,rows,roster_checks)
    else:
        result=[];by_pair={}
        for student in roster['rows']:
            if has_assigned_exam_roll(student):by_pair.setdefault((student['regno'],student['exam_roll']),[]).append(student)
        for row in rows:
            issues=[];candidates=by_pair.get((row['regno'],row['exam_roll']),[])
            if row['subject'] not in SUBJECTS.values():issues.append('Unsupported Internal Assessment subject; manual review.')
            if len(candidates)!=1:issues.append('Student identifiers are missing or ambiguous in the Students report.')
            else:
                student=candidates[0]
                if norm(student['name'])!=norm(row['name']):issues.append('Student name differs from the paired Students report identity.')
                if not enrolled_in_subject(student,row['subject']):issues.append('Student is not enrolled in this subject.')
                if student['approval'].lower()!='approved' or student['exam_fee'].lower()!='paid':issues.append('Students report eligibility/fee requires manual review.')
            try:
                maximum=number(row['maximum']);marks=number(row['marks'])
                if maximum<=0 or marks<maximum*number('0.4') or marks>maximum:issues.append('Present marks must be within the official maximum and at least 40%.')
            except (ValueError,TypeError,ArithmeticError):issues.append('Marks/maximum are not valid numbers.')
            result.append(dict(request=row,eligible=not issues,reasons=issues))
    identities={} ; exam_owners={}
    for r in rows:
        identities.setdefault(r['regno'],set()).add(r['exam_roll']);exam_owners.setdefault(r['exam_roll'],set()).add(r['regno'])
    if roster:
        for r in roster['rows']:
            if has_assigned_exam_roll(r):
                identities.setdefault(r['regno'],set()).add(r['exam_roll']);exam_owners.setdefault(r['exam_roll'],set()).add(r['regno'])
    absent={r['regno'] for r in rows if r['attendance']!='P'}
    held={h['registration'] for h in holds}
    for entry in result:
        r=entry['request'];issues=entry['reasons']
        if len(identities.get(r['regno'],set()))!=1 or len(exam_owners.get(r['exam_roll'],set()))!=1:issues.append('Conflicting registration/exam-roll ownership; manual identity review.')
        if r['regno'] in absent or r['regno'] in held:issues.append('Student has a known absence/identity hold; manual Principal approval only in every subject.')
        if r['attendance']!='P':issues.append('Absent or unknown attendance is never approved automatically.')
        if (r['regno'],r['exam_roll']) in exclusions:issues.append('Excluded student: all subjects skipped by exact registration/exam-roll pair.')
        entry['reasons']=list(dict.fromkeys(issues));entry['eligible']=not entry['reasons']
    return result

class ApprovalJournal:
    def __init__(self,root,scope):
        if set(scope)!={'school','principal','class','session'}:raise ValueError('Invalid Principal journal scope.')
        accounts(scope['school'],scope['principal'])
        if scope['class'] not in {'10','11','12'} or not re.fullmatch(r'\d{4}-\d{2}',scope['session']):raise ValueError('Invalid Principal journal class/session.')
        base=Path(root).resolve();self.folder=(base/'Principal_Approvals'/f'Class {scope["class"]} - {scope["session"]} - {scope["principal"]}').resolve()
        if not self.folder.is_relative_to(base):raise ValueError('Principal journal must remain inside its configured private history.')
        self.folder.mkdir(parents=True,exist_ok=True);self.path=self.folder/'events.jsonl';self.scope=scope;self.events=[];self.previous='0'*64
        if self.path.exists():
            for line in self.path.read_text(encoding='utf-8').splitlines():
                e=json.loads(line);digest=e.pop('hash')
                if e.get('previous')!=self.previous or Audit._hash(e)!=digest or e.get('scope')!=scope:raise ValueError('Principal approval journal integrity/scope failed.')
                e['hash']=digest;self.events.append(e);self.previous=digest
    def append(self,kind,**fields):
        e=dict(time=datetime.now(timezone.utc).isoformat(),kind=kind,scope=self.scope,previous=self.previous,**fields);e['hash']=Audit._hash(e)
        with self.path.open('a',encoding='utf-8') as out:out.write(json.dumps(e)+'\n');out.flush();os.fsync(out.fileno())
        self.events.append(e);self.previous=e['hash']
    def unresolved(self):
        result={}
        for e in self.events:
            if e['kind'] in {'approval_open_attempt','approval_reappeared'}:result[e['request']['id']]=e
            elif e['kind'] in {'approval_verified','manual_approval_resolution'}:result.pop(e['request_id'],None)
        return result
    def approved(self):
        result=set()
        for e in self.events:
            if e['kind']=='approval_verified' or (e['kind']=='manual_approval_resolution' and e['outcome']=='approved'):result.add(e['request_id'])
            elif e['kind']=='manual_approval_resolution':result.discard(e['request_id'])
        return result
    def reconcile(self,request_id,outcome,actor,note):
        if request_id not in self.unresolved() or outcome not in {'approved','not_approved'} or len(actor.strip())<3 or len(note.strip())<20:raise ValueError('Choose an unresolved request and provide administrator identity and a specific official verification note (20+ characters).')
        request=self.unresolved()[request_id]['request']
        self.append('manual_approval_resolution',request_id=request_id,request=request,outcome=outcome,actor=actor.strip(),note=note.strip())

def reconcile_interrupted_approvals(journal,pending,evidence):
    """Recover only a submitted request's exact single removal from hashed exports.

    Pending-queue absence alone is never an approval receipt. Both complete
    exports must belong to this journal and all unrelated requests must agree.
    """
    def checked_export(proof):
        path=Path(proof['file']).resolve(strict=True)
        if not path.is_relative_to(journal.folder.resolve()):raise ValueError('Recovery evidence is outside the Principal journal.')
        if hashlib.sha256(path.read_bytes()).hexdigest()!=proof['sha256']:raise ValueError('Recovery export integrity failed.')
        if proof['total']==0 and path.suffix=='.json':
            data=json.loads(path.read_text(encoding='utf-8'));validate_grid(data,journal.scope['class'],journal.scope['principal'])
            if data['total']!='0':raise ValueError('Recovery empty-queue evidence is not empty.')
            return []
        return parse_export(path,journal.scope['class'],proof['total'])
    recovered=[]
    unresolved=journal.unresolved()
    if not unresolved:return recovered
    try:
        after=checked_export(evidence)
        if {r['id']:r for r in after}!={r['id']:r for r in pending}:return recovered
    except (ValueError,OSError,KeyError,TypeError):return recovered
    for request_id,event in list(unresolved.items()):
        request=event['request'];start=next(i for i,e in enumerate(journal.events)if e['hash']==event['hash'])
        attempted=any(e['kind']=='approval_submit_attempt' and e.get('request_id')==request_id and e.get('request')==request for e in journal.events[start+1:])
        if not attempted:continue
        try:
            before=checked_export(event['evidence'])
            if next((r for r in before if r['id']==request_id),None)!=request:continue
            verify_removal(before,after,request)
        except (ValueError,OSError,KeyError,TypeError):continue
        journal.append('approval_verified',request_id=request_id,request=request,source_sha256=event.get('source_sha256',''),evidence=evidence,
                       verification='Recovered interrupted verification: prior exact request-bound submit attempt, hashed complete before/after exports, only that exact request removed; every other request unchanged. Not an independent Approved-status report.',
                       recovered=True,before_evidence=event['evidence'])
        recovered.append(request_id)
    return recovered

class PrincipalPortal:
    def __init__(self,driver,folder,cls,principal,stop_event=None):
        self.driver=driver;self.folder=Path(folder);self.cls=cls;self.principal=principal;self.stop_event=stop_event
        self.folder.mkdir(parents=True,exist_ok=True)
    def stop(self):
        from .internal_assessment_browser import _stop
        _stop(self.stop_event)
    def snapshot(self):
        self.stop()
        if not is_report_page(self.driver.current_url,URL):raise ValueError('Principal approval page redirected or login expired.')
        data=self.driver.execute_script(SNAPSHOT);validate_grid(data,self.cls,self.principal);return data
    def open(self):
        open_report_from_sidebar(self.driver,URL,'Principal Internal Assessment',self.stop_event)
        el=self.driver.find_element(By.ID,'bodypH_ddlClass')
        if Select(el).first_selected_option.get_attribute('value')!=self.cls:
            before=begin_read(self.driver,'bodypH_ddlClass','Principal class selection');Select(el).select_by_value(self.cls);wait_read(self.driver,before,self.stop)
        return self.snapshot()
    def page(self,n):
        data=self.snapshot()
        if data['page']==[str(n)]:return data
        href=f"javascript:__doPostBack('ctl00$bodypH$gridRR','Page${n}')"
        links=[e for e in self.driver.find_elements(By.CSS_SELECTOR,'#bodypH_gridRR a[href]') if e.get_attribute('href')==href and e.is_enabled()]
        if len(links)!=1:raise ValueError('Principal pager destination is missing/ambiguous; no decision clicked.')
        before=begin_read(self.driver,'','Principal request page');links[0].click();wait_read(self.driver,before,self.stop)
        data=self.snapshot()
        if data['page']!=[str(n)]:raise ValueError('Principal pager did not reach the requested page.')
        return data
    def click_download(self,expected):
        # Only this read-only control is retried. Approval/rejection clicks are
        # never retried. JKBOSE's success toast can cover the bottom-right button.
        def ready(driver):
            self.stop();current=self.snapshot()
            if current['total']!=expected['total'] or current['approve_visible'] or current['reject_visible']:
                raise ValueError('Principal queue/popup changed before the read-only export.')
            button=driver.find_element(By.ID,'bodypH_btnDownloadExcel')
            driver.execute_script("arguments[0].scrollIntoView({block:'center',inline:'nearest'});",button)
            try:button.click();return True
            except ElementClickInterceptedException:return False
        WebDriverWait(self.driver,30,poll_frequency=.25).until(ready)

    def read_all(self):
        data=self.open()
        if data['approve_visible'] or data['reject_visible']:raise ValueError('Close any existing Principal decision popup before starting.')
        if int(data['total'])==0:
            path=self.folder/('empty-queue-'+datetime.now(timezone.utc).strftime('%H%M%S%f')+'.json')
            path.write_text(json.dumps(data,indent=2),encoding='utf-8')
            return [],dict(file=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),total=0)
        before={p.name for p in self.folder.iterdir()}
        self.click_download(data)
        stable={}
        def downloaded(d):
            self.stop();files=[p for p in self.folder.glob('*.xls') if p.name not in before and p.stat().st_size>0]
            if len(files)>1:raise ValueError('Multiple Principal downloads; no unambiguous export.')
            if not files or any(self.folder.glob('*.crdownload')):return False
            p=files[0];size=p.stat().st_size;previous=stable.get(p);stable[p]=size
            return p if previous==size else False
        path=WebDriverWait(self.driver,60,poll_frequency=.3).until(downloaded)
        after=self.snapshot()
        if after['total']!=data['total']:raise ValueError('Principal queue changed while downloading; rerun read-only preview.')
        rows=parse_export(path,self.cls,int(data['total']))
        for raw in after['rows']:
            if len([r for r in rows if grid_row_matches(raw,r)])!=1:raise ValueError('Downloaded Principal requests differ from the visible grid.')
        return rows,dict(file=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),total=len(rows))
    def seek(self,request):
        data=self.page(1)
        for n in range(1,(int(data['total'])+9)//10+1):
            if n>1:data=self.page(n)
            found=[r for r in data['rows'] if grid_row_matches(r,request)]
            if len(found)>1:raise ValueError('Duplicate visible Principal request.')
            if found:return found[0],data
        raise ValueError('Validated Principal request is no longer in the pending queue.')
    def open_approval(self,request):
        row,data=self.seek(request)
        if data['approve_visible'] or data['reject_visible']:raise ValueError('Unexpected Principal decision popup.')
        before=begin_read(self.driver,row['approve'],'Open Principal approval confirmation')
        self.driver.find_element(By.ID,row['approve']).click();wait_read(self.driver,before,self.stop)
        self.approval_context=dict(request=dict(request),before=data,confirmed=False)
        current=self.snapshot()
        if (current['total']!=data['total'] or not current['approve_visible'] or current['reject_visible'] or current['approve_id']!=request['id']
                or not any(grid_row_matches(raw,request) for raw in current['rows'])):raise ValueError('Approval popup is not bound to the validated request; stop without confirmation.')
    def confirm(self,request):
        data=self.snapshot()
        if not data['approve_visible'] or data['reject_visible'] or data['approve_id']!=request['id']:raise ValueError('Principal confirmation target changed.')
        before=begin_read(self.driver,'bodypH_btnApproveSubmit','Principal approval response')
        self.driver.find_element(By.ID,'bodypH_btnApproveSubmit').click();wait_read(self.driver,before,self.stop)
        self.approval_context['confirmed']=True
    def verify_transition(self,pending,request):
        """Fast class-wide verification: exact popup/postback + complete page + count.

        Remaining unseen rows retain initial verified values until the final full
        export audit. This evidence is explicitly distinct from a full export.
        """
        context=getattr(self,'approval_context',None)
        if not context or context['request']!=request or not context['confirmed']:raise ValueError('No completed response bound to this approval request.')
        if len(pending)!=int(context['before']['total']) or next((r for r in pending if r['id']==request['id']),None)!=request:raise ValueError('Approval queue baseline changed.')
        data=self.snapshot();remaining=[r for r in pending if r['id']!=request['id']]
        if int(data['total'])!=len(remaining) or data['approve_visible'] or data['reject_visible']:raise ValueError('Approval response did not reduce the selected class queue by exactly one.')
        if remaining:
            target_page=min(int(context['before']['page'][0]),(len(remaining)+9)//10)
            if data['page']!=[str(target_page)]:data=self.page(target_page)
        if any(raw['cells'][1:3]==[request['exam_roll'],request['regno']] and subject_name(raw['cells'][4])==request['subject'] for raw in data['rows']):raise ValueError('Approved student/subject is still displayed; do not retry blindly.')
        seen=set()
        for raw in data['rows']:
            matches=[r for r in remaining if grid_row_matches(raw,r)]
            if len(matches)!=1 or matches[0]['id']in seen:raise ValueError('Visible remaining requests differ from the verified queue.')
            seen.add(matches[0]['id'])
        if remaining and data['page']==context['before']['page']:
            for raw in context['before']['rows']:
                if grid_row_matches(raw,request):continue
                if not any(current['cells']==raw['cells']for current in data['rows']):raise ValueError('Another displayed request changed during approval.')
        proof_file=self.folder/('queue-transition-'+datetime.now(timezone.utc).strftime('%H%M%S%f')+'.json')
        proof_file.write_text(json.dumps(dict(request=request,before=context['before'],after=data,completed_bound_response=True),indent=2),encoding='utf-8')
        self.approval_context=None
        return remaining,dict(file=str(proof_file),sha256=hashlib.sha256(proof_file.read_bytes()).hexdigest(),total=len(remaining),kind='live_queue_transition',coverage='Complete displayed page and live class total; unseen rows are checked by final full-queue audit.')


def verify_removal(before,after,request):
    old={r['id']:r for r in before};new={r['id']:r for r in after}
    if old.get(request['id'])!=request or set(new)!=set(old)-{request['id']} or any(new[k]!=old[k] for k in new):
        raise ValueError('Approval response did not prove removal of only the exact request. Do not retry automatically.')

def tracking_receipts(root,school,cls,session):
    """Project intact approval history; never infer approval from queue absence alone."""
    match=re.fullmatch(r'sch(\d{6})',school.lower())
    if not match:return {},{}
    scope=dict(school=school.lower(),principal='prin'+match[1],**{'class':cls},session=session)
    path=Path(root)/'Principal_Approvals'/f'Class {cls} - {session} - {scope["principal"]}'/'events.jsonl'
    if not path.exists():return {},{}
    journal=ApprovalJournal(root,scope);approved={};unresolved=journal.unresolved();active=journal.approved()
    for event in journal.events:
        if (event['kind']=='approval_verified' or event['kind']=='manual_approval_resolution' and event['outcome']=='approved'):
            request=event['request']
            if request['id'] in active and request['id'] not in unresolved:approved[key(request)]=dict(request,verified_time=event['time'],verified_method='manual' if event['kind']=='manual_approval_resolution' else 'request_removal')
    return approved,{key(event['request']):event['request'] for event in unresolved.values()}

def run_approval(jobs,source,history_root,school_account,principal_account,*,execute=False,confirm_session=False,
                 report_session='Auto',confirm_callback=None,stop_event=None,on_driver_ready=None,driver_factory=None,login_callback=None,report_reader=None,portal_factory=None,
                 approval_mode='excel',exclusions=None,student_class=None,session='',on_progress=None):
    from ..browser import get_driver,login_only
    from .internal_assessment_policy import CohortAbsencePolicy,source_absences,apply_holds
    from .internal_assessment_browser import _stop
    school_account=school_account.lower();principal_account=principal_account.lower();accounts(school_account,principal_account)
    if not confirm_session:raise ValueError('Confirm the academic/examination session and exclusive Principal operation.')
    if approval_mode not in APPROVAL_MODES:raise ValueError('Choose a valid Principal approval mode.')
    excluded=parse_exclusions(exclusions)
    if approval_mode=='class_except' and not excluded:raise ValueError('Add at least one excluded student pair.')
    if approval_mode!='class_except' and excluded:raise ValueError('Exclusions are supported only in class-except mode.')
    plans=[j['plan'] for j in jobs or []] if approval_mode=='excel' else []
    if approval_mode=='excel':
        if not plans or any(p is None for p in plans) or len({p.subject for p in plans})!=len(plans):raise ValueError('Choose valid distinct source subjects for Principal approval.')
        source=Path(source).resolve(strict=True);digest=hashlib.sha256(source.read_bytes()).hexdigest();cls=plans[0].student_class;session=plans[0].session
        if any(Path(p.source)!=source or p.source_sha256!=digest or p.student_class!=cls or p.session!=session for p in plans):raise ValueError('Principal plans do not match this source/class/session.')
    else:
        source=None;digest='';cls=class_key(student_class);session=session_key(session)
        if cls not in {'10','11','12'} or not re.fullmatch(r'\d{4}-\d{2}',session):raise ValueError('Choose the class and academic session.')
    root=prepare_history_root(history_root);scope=dict(school=school_account,principal=principal_account,**{'class':cls},session=session)
    driver=None
    with RunLock(root):
        policy=CohortAbsencePolicy(root,school_account,cls,session)
        if source:
            current=source_absences(source,plans[0].sheet)
            policy.add(current,origin='consolidated_source',evidence=digest);policy.reconcile_source(source,plans[0].sheet,current_holds=current)
        for plan in plans:apply_holds(plan,policy.holds)
        journal=ApprovalJournal(root,scope)
        run=journal.folder/'Runs'/datetime.now(timezone.utc).strftime('%Y-%m-%d_%H-%M-%S-%f');run.mkdir(parents=True)
        report=dict(approval_mode=approval_mode,exclusions=sorted(excluded),mode='approve' if execute else 'read_only',scope=scope,source=str(source) if source else '',source_sha256=digest,status='started',verified_approved=[],requests=[],unresolved=list(journal.unresolved()))
        def progress(queue,stage):
            approved=journal.approved();initial=report.get('initial_requests',queue);tracked=set(approved)|{r['id']for r in initial}
            students={(r['regno'],r['exam_roll'])for r in initial}
            for event in journal.events:
                if event.get('request_id')in approved and event.get('request'):students.add((event['request']['regno'],event['request']['exam_roll']))
            value=dict(stage=stage,student_class=cls,session=session,verified=len(approved),total_records=len(tracked),remaining=len(queue),students_total=len(students),students_remaining=len({(r['regno'],r['exam_roll'])for r in queue}),approved_this_run=len(report['verified_approved']),held=sum(not e['eligible']for e in report['requests']))
            report['progress']=value
            if on_progress:
                try:on_progress(value)
                except Exception:print('[PROGRESS DISPLAY] Display update unavailable; approval history is retained.')
            print(f'[APPROVAL PROGRESS] {value["verified"]}/{value["total_records"]} tracked awards verified; {value["remaining"]} pending awards; {value["students_remaining"]} students pending; {value["held"]} held/excluded awards.')
        def write():
            temp=run/'report.tmp';temp.write_text(json.dumps(report,indent=2),encoding='utf-8');os.replace(temp,run/'report.json')
            rows=['<!doctype html><meta charset="utf-8"><title>Principal Approval History</title><h1>Principal Approval History</h1>']
            from html import escape
            rows.append('<p>Outcome: '+escape(report['status'])+'. Approval is separate from submission.</p>')
            rows.append('<p>Selection: '+escape(report['approval_mode'])+'; excluded student pairs: '+str(len(report['exclusions']))+'.</p>')
            rows.append('<p>Verified in this run: '+str(len(report['verified_approved']))+'.</p><table border="1"><tr><th>Subject</th><th>Registration</th><th>Exam roll</th><th>Name</th><th>Given / Max</th><th>Outcome / review</th></tr>')
            verified=set(report['verified_approved']);unresolved=set(journal.unresolved())
            if unresolved:rows.append('<p>Earlier unresolved requests: '+str(len(unresolved))+'. These are held for manual reconciliation; unrelated eligible requests may continue.</p>')
            for entry in report['requests']:
                r=entry['request'];outcome='Uncertain approval; verify official status before retrying' if r['id'] in unresolved else 'Approval verified' if r['id'] in verified else 'Eligible; not approved in this run' if entry['eligible'] else '; '.join(entry['reasons'])
                rows.append('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in [r['subject'],r['regno'],r['exam_roll'],r['name'],r['marks']+' / '+r['maximum'],outcome])+'</tr>')
            (run/'OPEN APPROVAL HISTORY.html').write_text(''.join(rows)+'</table>',encoding='utf-8')
            # A readable latest copy never acts as an approval receipt.
            (journal.folder/'OPEN APPROVAL HISTORY.html').write_text(''.join(rows)+'</table>',encoding='utf-8')
            from urllib.parse import quote
            links=['<!doctype html><meta charset="utf-8"><title>Principal histories</title><h1>Principal Approval Histories</h1><ul>']
            for folder in sorted(journal.folder.parent.iterdir()):
                if folder.is_dir() and (folder/'OPEN APPROVAL HISTORY.html').is_file():links.append('<li><a href="'+quote(folder.name)+'/OPEN%20APPROVAL%20HISTORY.html">'+escape(folder.name)+'</a></li>')
            (journal.folder.parent/'OPEN HISTORY.html').write_text(''.join(links)+'</ul>',encoding='utf-8')
        try:
            _stop(stop_event)
            driver=(driver_factory or get_driver)(profile_dir=str(Path(__file__).resolve().parents[2]/'.chrome_profile_principal'),download_dir=str(run/'Downloads'))
            if on_driver_ready:on_driver_ready(driver)
            (login_callback or login_only)(driver,role='principal')
            portal=(portal_factory or PrincipalPortal)(driver,run/'Downloads',cls,principal_account,stop_event)
            pending,evidence=portal.read_all();report['initial_export']=evidence;report['initial_requests']=list(pending)
            if any((reg==r['regno'] and exam!=r['exam_roll']) or (exam==r['exam_roll'] and reg!=r['regno']) for reg,exam in excluded for r in pending):
                raise ValueError('An exclusion conflicts with the fetched student identifiers; correct the pair before approval.')
            active_subjects={r['subject'] for r in pending};roster_checks={};roster=None;matches=[]
            if pending and (approval_mode!='excel' or any(p.subject in active_subjects for p in plans)):
                roster=(report_reader or read_student_report)(driver,cls,principal_account,session,report_session,stop_event)
                if roster.get('account','').lower()!=principal_account or roster.get('workbook_session')!=session:raise ValueError('Principal Students report account/session differs.')
                report['student_report']=write_snapshot(run,roster)
                for plan in plans:
                    if plan.subject not in active_subjects:continue
                    matched=match_student_report(plan,roster,{'verified':set(),'manual':set()});matches.append(matched)
                    roster_checks[plan.subject]={(r['excel_registration'],r['excel_exam_roll']):r for r in matched['records'] if r['excel_registration']}
                (run/'roster_matching.json').write_text(json.dumps(matches,indent=2),encoding='utf-8')
            recovered=reconcile_interrupted_approvals(journal,pending,evidence)
            report['recovered_approvals']=recovered
            if recovered:print(f'[APPROVAL RECOVERED] {len(recovered)} interrupted receipt(s) verified by exact before/after exports; no approval clicked again.')
            def classify(queue):
                entries=select_requests(plans,queue,roster_checks,mode=approval_mode,exclusions=excluded,roster=roster,holds=CohortAbsencePolicy(root,school_account,cls,session).holds)
                unresolved=journal.unresolved();held_ids=set(unresolved);held_keys={key(e['request']) for e in unresolved.values()}
                for entry in entries:
                    if entry['request']['id'] in held_ids or key(entry['request']) in held_keys:
                        entry['reasons'].append('Earlier approval for this request is uncertain; reconcile its official status before retrying.');entry['eligible']=False
                return entries
            report['requests']=classify(pending);report['unmatched_exclusions']=[list(pair) for pair in sorted(excluded-{(r['regno'],r['exam_roll']) for r in pending})];report['total_pending']=len(pending);progress(pending,'preview');write()
            uncertain=journal.unresolved()
            if uncertain:print(f'[APPROVAL HOLDS] {len(uncertain)} earlier uncertain request(s) will not be retried. Other validated requests may continue.')
            reappeared={r['id'] for r in pending}&journal.approved()
            if reappeared:
                for request in pending:
                    if request['id'] in reappeared:journal.append('approval_reappeared',request=request)
                raise ValueError('Previously verified approvals reappeared in Principal requests; reconcile before continuing.')
            eligible=[e['request'] for e in report['requests'] if e['eligible']]
            order={plan.subject:i for i,plan in enumerate(plans)}
            if approval_mode=='excel':eligible.sort(key=lambda r:order[r['subject']])
            print(f'[PRINCIPAL PREVIEW] {len(eligible)} eligible; {len(report["requests"])-len(eligible)} held; {len(pending)} total pending. Principal login: {principal_account}.')
            journal.append('approval_preview',approval_mode=approval_mode,exclusions=sorted(excluded),source_sha256=digest,execute=execute,evidence=evidence,eligible=[r['id'] for r in eligible],report=str(run/'report.json'))
            if not execute:report['status']='checked_read_only';return report
            if not eligible:report['status']='nothing_eligible';return report
            if not confirm_callback or not confirm_callback(dict(scope=scope,approval_mode=approval_mode,exclusions=sorted(excluded),report_session=roster['report_session'],source=str(source) if source else '',rows=report['requests'])):report['status']='cancelled';return report
            fresh,fresh_evidence=portal.read_all()
            if {r['id']:r for r in fresh}!={r['id']:r for r in pending}:raise ValueError('Principal queue changed after preview. No approval started; review again.')
            pending=fresh
            for request in eligible:
                _stop(stop_event)
                if source and hashlib.sha256(source.read_bytes()).hexdigest()!=digest:raise ValueError('Source workbook changed; no further approval allowed.')
                # Reuse source matching before every action, including durable class-wide holds.
                current_policy=CohortAbsencePolicy(root,school_account,cls,session)
                for plan in plans:apply_holds(plan,current_policy.holds)
                if not next(e for e in classify(pending) if e['request']['id']==request['id'])['eligible']:raise ValueError('An approval acquired a hold after preview.')
                journal.append('approval_open_attempt',request=request,source_sha256=digest,evidence=fresh_evidence)
                portal.open_approval(request)
                _stop(stop_event)
                if source and hashlib.sha256(source.read_bytes()).hexdigest()!=digest:raise ValueError('Source changed while the Principal popup was open; confirmation not clicked.')
                journal.append('approval_submit_attempt',request_id=request['id'],request=request)
                portal.confirm(request)
                fast=getattr(portal,'verify_transition',None)
                if approval_mode!='excel' and callable(fast):after,proof=fast(pending,request)
                else:
                    after,proof=portal.read_all();verify_removal(pending,after,request)
                journal.append('approval_verified',request_id=request['id'],request=request,source_sha256=digest,evidence=proof,
                               verification=('Exact request-bound confirmation and completed response; live class total decreased by one, target absent and full displayed page validated; final full-queue audit required.' if proof.get('kind')=='live_queue_transition' else 'Single Approved confirmation bound to request ID; completed response and exact full-export removal with all other requests unchanged. Not independently read from a separate status report.'))
                report['verified_approved'].append(request['id']);pending=after;fresh_evidence=proof;progress(pending,'approving');write()
                print(f'[APPROVAL VERIFIED] {request["subject"]}: one exact request removed; {len(pending)} total remain.')
            if approval_mode!='excel':
                final_rows,final_proof=portal.read_all()
                if {r['id']:r for r in final_rows}!={r['id']:r for r in pending}:raise ValueError('Final class queue differs from tracked progress; inspect before any further decision.')
                report['final_queue_audit']=final_proof
            report['status']='finished_with_holds' if journal.unresolved() or any(not e['eligible'] for e in report['requests']) else 'finished'
            progress(pending,'complete' if not pending else 'held');write()
            return report
        except Exception as exc:
            report['status']='failed';report['error']=str(exc)[:1200]
            journal.append('approval_run_failed',reason=str(exc)[:1200],unresolved=list(journal.unresolved()));raise
        finally:
            report['unresolved']=list(journal.unresolved());write()
            try:
                from .internal_assessment_tracking import refresh_tracking_copy
                if source:report['tracking_workbook']=refresh_tracking_copy(source,root,cls,session,school_account)
                else:report['tracking_workbook']='Not selected in class-wide mode; approval receipts retained in Principal history.'
                write();print('[TRACKING EXCEL] '+report['tracking_workbook'])
            except Exception as tracking_error:
                report['tracking_error']=str(tracking_error)[:1000];write()
                print('[TRACKING NOT REFRESHED] '+str(tracking_error)+'. Approval journal retained.')
            print(f'[PRINCIPAL HISTORY] {journal.folder/"OPEN APPROVAL HISTORY.html"}')
            if driver is not None:
                if report['status']=='failed' and on_driver_ready:
                    driver._jkbose_keep_open=True
                    report['browser_left_open']=True;write();print('[BROWSER KEPT OPEN] Inspect the page or use Force Close Browser. No failed approval is retried.')
                else:
                    try:driver.quit()
                    except Exception:pass
