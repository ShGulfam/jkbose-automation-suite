"""Offline validation shared by Check and Submit. Never contacts a browser."""
from pathlib import Path
import re
import pandas as pd
from openpyxl import load_workbook
from ..subjects import canonicalize_row_columns

EXAM_SCHEMA = 'hss-jkbose-exam-v1'

def normalize_class(value):
    match = re.fullmatch(r'(11|12)(?:th)?', str(value or '').strip(), re.I)
    return match.group(1) if match else ''

def excluded_reason(row):
    for field in ('ExamDropped', 'isExamDropped', 'examDropped', 'Dropped'):
        if str(row.get(field, '')).strip().lower() in ('true', 'yes', '1', 'dropped'):
            return 'Exam-dropped student: excluded from automation'
    for field in ('Exam Status', 'ExamStatus', 'examStatus', 'Status'):
        if str(row.get(field, '')).strip().lower() in ('dropped', 'exam dropped', 'withdrawn', 'cancelled', 'rejected'):
            return 'Dropped/withdrawn/rejected application: manual review only'
    return ''

def read_exam_dataframe(file, sheet='Sheet2', class_override=None):
    path = Path(file)
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Choose an .xlsx workbook; legacy .xls must be saved as .xlsx first.')
    override = normalize_class(class_override) if class_override else ''
    if class_override and not override:
        raise ValueError('Class must be 11 or 12.')
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet not in book.sheetnames:
            raise ValueError(f"Worksheet '{sheet}' is missing. Choose the automation Sheet2 export.")
        ws = book[sheet]
        headers = [str(c.value or '').strip() for c in next(ws.iter_rows())]
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ValueError('Duplicate workbook headers; reconcile the file before automation.')
        rows, physical_indexes, registrations = [], [], set()
        for excel_row, cells in enumerate(ws.iter_rows(min_row=2), 2):
            if not any(c.value is not None for c in cells):
                continue
            if any(c.data_type == 'f' for c in cells):
                raise ValueError(f'Excel row {excel_row} contains formulas; export values only.')
            raw = {h: cells[i].value if i < len(cells) and cells[i].value is not None else '' for i, h in enumerate(headers) if h}
            row = canonicalize_row_columns(raw)
            registration = row.get('RegNo', '')
            reg = str(registration).strip()
            if not re.fullmatch(r'[A-Za-z0-9-]{6,}', reg) or re.search(r'0{5,}$', reg):
                raise ValueError(f'Excel row {excel_row}: missing or placeholder registration.')
            if not isinstance(registration, str) and len(reg) > 15:
                raise ValueError(f'Excel row {excel_row}: registration must be stored as exact Text.')
            if reg.lower() in registrations:
                raise ValueError(f'Duplicate registration at Excel row {excel_row}; select the approved application first.')
            registrations.add(reg.lower())
            supplied = normalize_class(row.get('Class', ''))
            if override and supplied and override != supplied:
                raise ValueError(f'Excel row {excel_row} belongs to another class; class override refused.')
            cls = override or supplied
            if not cls:
                raise ValueError(f'Excel row {excel_row}: choose Class 11 or 12.')
            row['RegNo'], row['Class'] = reg, cls
            schema = str(row.get('ExportSchema', '')).strip()
            if schema and schema != EXAM_SCHEMA:
                raise ValueError('Unknown exam export schema; download a compatible workbook.')
            if schema:
                if not re.fullmatch(r'20\d{2}-\d{2}', str(row.get('Session', '')).strip()):
                    raise ValueError(f'Excel row {excel_row}: academic session missing.')
                if not re.fullmatch(r'(?!0+$)\d{1,3}', str(row.get('ClassRollNo', '')).strip()):
                    raise ValueError(f'Excel row {excel_row}: no assigned Class Roll No.; manual review only.')
                if str(row.get('ApprovedForAutomation', '')).strip().lower() != 'yes':
                    raise ValueError(f'Excel row {excel_row}: not approved for automation.')
                if not str(row.get('Name', '')).strip():
                    raise ValueError(f'Excel row {excel_row}: current student name missing.')
            rows.append({key: str(value) if value is not None else '' for key, value in row.items()})
            physical_indexes.append(excel_row - 2)
        if not rows:
            raise ValueError('The selected worksheet contains no student rows.')
        return pd.DataFrame(rows, index=physical_indexes).fillna('')
    finally:
        book.close()

def load_check_inputs(file, student_class=None, sheet='Sheet2'):
    path = Path(file)
    if path.suffix.lower() == '.xlsx':
        frame = read_exam_dataframe(path, sheet, student_class)
        return [dict(row, _skip_reason=excluded_reason(row)) for row in frame.to_dict('records')]
    if path.suffix.lower() != '.txt':
        raise ValueError('Check accepts an automation .xlsx workbook or an RR .txt list.')
    cls = normalize_class(student_class)
    if not cls:
        raise ValueError('Choose Class 11 or 12 before checking a text list.')
    lines = path.read_text(encoding='utf-8-sig').splitlines()
    seen = set()
    rows = []
    for line in lines:
        reg = line.strip()
        if reg and (not re.fullmatch(r'[A-Za-z0-9-]{6,}', reg) or re.search(r'0{5,}$', reg)):
            raise ValueError('The RR text list contains an invalid or placeholder registration.')
        if reg and reg.lower() in seen:
            raise ValueError('The RR text list contains duplicate registrations.')
        if reg:
            seen.add(reg.lower())
        rows.append({'RegNo': reg, 'Class': cls, '_skip_reason': ''})
    if not any(r['RegNo'] for r in rows):
        raise ValueError('The RR text list is empty.')
    return rows
