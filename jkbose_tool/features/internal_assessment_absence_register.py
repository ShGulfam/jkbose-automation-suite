"""Readable class absence register projected from immutable submission evidence."""
from pathlib import Path
import hashlib,json,os
from .internal_assessment import Audit,norm
from .internal_assessment_history import storage_root,manifest_paths

def refresh_absence_register(source,history_root,student_class,session,account):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from .internal_assessment_batch import load_jobs
    source=Path(source).resolve(strict=True);digest=hashlib.sha256(source.read_bytes()).hexdigest()
    root=storage_root(history_root);records={}
    def key(subject,a):return (subject,a.regno,a.exam_roll)
    for job in load_jobs(source,student_class=student_class,session=session,all_subjects=True):
        if job['plan']:
            for a in job['plan'].awards:
                if a.attendance=='A':records[key(job['subject'],a)]=dict(name=a.name,status='Source Absent - not verified submitted',evidence='',time='',source_current=True)
    for path in manifest_paths(root):
        meta=json.loads(path.read_text(encoding='utf-8'))
        if meta.get('student_class')!=student_class or meta.get('session')!=session:continue
        audit,plan=Audit.existing(path.parent)
        accounts={e['school_account'].lower() for e in audit.events if e['kind']=='run_started' and e.get('school_account')}
        if accounts!={account.lower()}:continue
        awards={a.regno:a for a in plan.awards if a.attendance=='A'}
        for e in audit.events:
            if e['kind'] not in {'submit_attempt','verified_saved','manual_reconciliation','portal_award_rejected'}:continue
            for reg in (e.get('registrations') or ([e['regno']] if e.get('regno') else [])):
                if reg not in awards:continue
                a=awards[reg];k=key(plan.subject,a);old=records.get(k,{})
                if old.get('time','')>e['time']:continue
                if e['kind']=='submit_attempt':status='Uncertain - Absent submit attempted'
                elif e['kind']=='verified_saved':status='Submitted as Absent - verified'
                elif e['kind']=='portal_award_rejected':status='Submitted as Absent - rejected on portal'
                elif e['disposition']=='saved_manually_verified':status='Submitted as Absent - manually verified'
                else:status='Not saved - manually verified'
                records[k]=dict(name=a.name,status=status,evidence=str(audit.folder),time=e['time'],source_current=old.get('source_current',False))
    b=Workbook();w=b.active;w.title='Absent Awards';w.append([f'Class {student_class} - Internal Assessment Absence Register - {session}'])
    w.append(['Original source',str(source)]);w.append(['Source SHA256',digest])
    w.append(['Absent submission is separate from approval. Every student absent in any subject is excluded from automatic Principal approval in all subjects.'])
    w.append(['S.No.','Subject S.No.','Class','Session','Subject','Registration No.','Exam Roll No.','Student Name','Attendance','Submission Status','Automatic Principal Approval','Recorded At (UTC)','Evidence Folder','Current Source Absent'])
    counts={}
    for i,((subject,reg,exam),r) in enumerate(sorted(records.items()),1):
        counts[subject]=counts.get(subject,0)+1
        w.append([i,counts[subject],student_class,session,subject,reg,exam,r['name'],'Absent',r['status'],'NEVER - manual approval required',r['time'],r['evidence'],'Yes' if r['source_current'] else 'No - historical submission; review current source'])
    for cell in w[5]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='334155');cell.alignment=Alignment(wrap_text=True)
    widths=[9,15,9,13,25,24,18,30,13,46,40,28,65,45]
    for i,width in enumerate(widths,1):w.column_dimensions[w.cell(5,i).column_letter].width=width
    for row in w.iter_rows(min_row=6):
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        for c in [row[5],row[6]]:c.number_format='@'
    w.freeze_panes='F6';w.auto_filter.ref=f'A5:N{max(5,w.max_row)}'
    dest=source.parent/f'Class_{student_class}_Internal_Assessment_Absence_Register_{session.replace("-","_")}.xlsx'
    temp=dest.with_name(dest.stem+'.tmp.xlsx')
    try:b.save(temp)
    finally:b.close()
    if hashlib.sha256(source.read_bytes()).hexdigest()!=digest:
        temp.unlink(missing_ok=True);raise ValueError('Source changed; absence register was not replaced.')
    os.replace(temp,dest)
    return str(dest)
