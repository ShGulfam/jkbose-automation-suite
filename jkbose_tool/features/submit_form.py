"""
Feature: "submit form" -- for each row in an Excel sheet, look up the
student by RegNo, fill the RR renewal / exam form (personal info,
address, subjects, photo/signature) -- UNLESS it's a Class 12 row, which
is a renewal of an already-registered Class 11 student and needs no
filling at all -- then pause for you to solve the captcha and click
Submit manually.

Field IDs come from jkbose_tool.config.FIELD_MAP, built from a live
`inspect` dump of renewalformxixii.aspx. If the site changes again,
re-run `inspect` and update FIELD_MAP in config.py rather than editing
this file -- the loop below is generic.

SPEED NOTE: every field in FIELD_MAP is looked up with an IMMEDIATE
check (no wait) first, falling back to a short config.QUICK_WAIT (a few
seconds) only if it's not there yet. This matters a lot in practice --
FIELD_MAP lists every possible field, but a lot of them may not exist on
the page in its current state (e.g. while "RR Renewal Is Closed", or
fields that only appear for certain streams). The old approach waited up
to config.ELEMENT_WAIT (25s) per missing field, which for ~30 mapped
fields plus photo/signature retries could add minutes per row. By the
time this loop runs, the page has already settled from the Get-Details
postback, so a genuinely-present field is found instantly anyway -- the
long wait was only ever paying off for fields that don't exist at all.
"""
import os
import time
import traceback
import re
import unicodedata
import pandas as pd
from openpyxl import load_workbook
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import InvalidSessionIdException, NoSuchWindowException, WebDriverException

from .. import config
from .exam_input import read_exam_dataframe, excluded_reason
from .exam_guard import require_exam_form, require_lookup_details, ExamPortalUnavailable
from ..browser import page_change_token, safe_click, wait_for_postback
from ..subjects import (
    canonical_stream, stream_value, subjects_from_row, canonical_subject, norm,
    SUBJECT_CODES, VOCATIONAL_CODES, canonicalize_row_columns, split_subject_roles
)



class BrowserSessionLost(RuntimeError):
    """Fatal interruption: the Selenium/Chrome session can no longer be used.

    This is deliberately distinct from an ordinary row-level validation or
    element error.  Continuing the batch after Chrome has disconnected only
    produces the same ``invalid session id`` error for every remaining row.
    ``row_index`` is zero-based so callers can tell the user where to resume.
    """
    def __init__(self, message, row_index=None, regno=None):
        super().__init__(message)
        self.row_index = row_index
        self.regno = regno


def _is_dead_browser_exception(exc):
    """Return True when an exception means this WebDriver is unrecoverable.

    Selenium wraps Chrome disconnects in several exception classes/messages
    depending on exactly where the browser vanished.  Keep the list narrow:
    normal site/field failures must remain ordinary row errors.
    """
    if isinstance(exc, (InvalidSessionIdException, NoSuchWindowException)):
        return True
    text = str(exc).lower()
    fatal_markers = (
        'invalid session id',
        'session deleted as the browser has closed the connection',
        'disconnected: not connected to devtools',
        'not connected to devtools',
        'chrome not reachable',
        'target window already closed',
        'no such window',
        'web view not found',
    )
    return any(marker in text for marker in fatal_markers)


def _raise_if_dead(exc):
    if _is_dead_browser_exception(exc):
        raise BrowserSessionLost('Chrome/Selenium session is no longer connected.') from exc


def _assert_browser_alive(driver):
    """Very cheap heartbeat before starting a new student row."""
    try:
        driver.execute_script('return document.readyState')
    except Exception as exc:
        _raise_if_dead(exc)
        raise


def _ensure_driver_alive(driver):
    """Compatibility heartbeat used by the manual-submit watcher.

    Keep a single implementation of the Selenium liveness check.  Earlier
    builds accidentally called this name without defining it, which made the
    safe wait path stop the batch and allowed the outer cleanup to close
    Chrome.  Delegating to ``_assert_browser_alive`` avoids duplicated logic.
    """
    return _assert_browser_alive(driver)


def _completion_message_from_text(text):
    """Return a compact completion message when *text* clearly says the form is done.

    We deliberately match completion *states*, not the bare word ``Submit`` that is
    present on every pending form.  This handles wording variations used by JKBOSE
    across Class 11/12 pages and JavaScript alerts.
    """
    raw = ' '.join(str(text or '').split())
    if not raw:
        return ''
    low = raw.lower()

    exact_phrases = (
        'already submitted', 'already been submitted', 'form has already been submitted',
        'exam form already submitted', 'renewal form already submitted',
        'already renewed', 'renewal already done', 'already filled', 'form already filled',
        'already registered for this examination', 'already registered for examination',
        'submitted successfully', 'successfully submitted', 'form submitted successfully',
        'application submitted successfully', 'exam form submitted', 'renewal submitted',
        'renewal completed', 'renewal complete', 'registration completed',
        'registration complete', 'application already submitted', 'already applied',
        'form is already submitted', 'form is submitted', 'record already submitted',
    )
    for phrase in exact_phrases:
        pos = low.find(phrase)
        if pos >= 0:
            start = max(0, pos - 70)
            end = min(len(raw), pos + len(phrase) + 120)
            return raw[start:end]

    # Flexible variants such as "candidate has already submitted the examination form"
    # or "exam form was submitted earlier".  Keep the windows short to avoid
    # confusing ordinary instructional text with a completion state.
    patterns = (
        r'\balready\b.{0,55}\b(submitted|renewed|filled|registered|applied)\b',
        r'\b(submitted|renewed|filled|registered|applied)\b.{0,55}\balready\b',
        r'\b(form|application|renewal|registration|exam form)\b.{0,35}\b(has been|was|is)\s+submitted\b',
        r'\b(submitted|renewed|registered)\b.{0,25}\b(successfully|earlier|previously)\b',
    )
    for pat in patterns:
        m = re.search(pat, low, flags=re.I)
        if m:
            a = max(0, m.start() - 70)
            b = min(len(raw), m.end() + 120)
            return raw[a:b]
    return ''


def _detect_already_submitted(driver, regno=None, class_val=None, allow_manual_success=False):
    """Accept only a response tied to the current verified lookup, never body instructions."""
    context = getattr(driver, '_exam_lookup_context', None)
    if not regno or context != (str(regno), str(class_val)):
        return False, ''
    try:
        alert = driver.switch_to.alert
        text = (alert.text or '').strip()
        msg = _completion_message_from_text(text)
        conditional = re.search(r'\bif\b|not\s+submitted|unsuccess|failed|\bpending\b|cannot\s+submit|unable', text, re.I)
        already = re.search(r'\balready\b', text, re.I)
        if msg and not conditional and (str(regno) in text or already or allow_manual_success):
            try:
                alert.accept()
            except Exception:
                pass
            return True, msg
    except Exception as exc:
        _raise_if_dead(exc)

    require_exam_form(driver, class_val)
    for node in driver.find_elements(By.CSS_SELECTOR, "[role='alert'], .alert-success, [id$='lblMessage'], [id$='lblmsg']"):
        text = (node.text or '').strip()
        msg = _completion_message_from_text(text)
        if node.is_displayed() and str(regno) in text and msg and not re.search(r'\bif\b|not\s+submitted|unsuccess|failed|\bpending\b', text, re.I):
            return True, msg
    return False, ''


def _detect_submit_success(driver, regno=None, class_val=None):
    """Verify the page after the user's manual Submit action when possible."""
    done, msg = _detect_already_submitted(driver, regno, class_val, allow_manual_success=True)
    if done:
        return True, msg
    # A successful postback sometimes uses a generic success message without
    # repeating "already".  _completion_message_from_text already covers those.
    return False, ''


class _ExcelProgressWriter:
    """Update only Status/Notes cells while preserving the workbook.

    The previous implementation rebuilt the whole sheet with pandas after
    every processed student. Besides being slow for large batches, that can
    discard unrelated workbook sheets/formatting. This writer keeps one
    workbook open, changes only the two progress cells, and saves in batches.
    """
    def __init__(self, path, sheet_name, save_every=5, history=None):
        self.history = history
        self.path = path
        self.wb = load_workbook(path)
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet_name}")
        self.ws = self.wb[sheet_name]
        headers = {str(c.value).strip(): c.column for c in self.ws[1] if c.value is not None}
        for name in ("Status", "Notes"):
            if name not in headers:
                col = self.ws.max_column + 1
                self.ws.cell(row=1, column=col, value=name)
                headers[name] = col
        self.status_col = headers["Status"]
        self.notes_col = headers["Notes"]
        self.save_every = max(1, int(save_every))
        self.dirty = 0

    def update(self, dataframe_index, status, notes=""):
        if self.history:
            self.history.progress(dataframe_index, status, notes)
        excel_row = int(dataframe_index) + 2  # header is row 1
        self.ws.cell(row=excel_row, column=self.status_col, value=status)
        self.ws.cell(row=excel_row, column=self.notes_col, value=notes)
        self.dirty += 1
        if self.dirty >= self.save_every:
            self.save()

    def save(self):
        if self.dirty:
            self.wb.save(self.path)
            self.dirty = 0


def _build_path_from_helper(name, photos_dir):
    photos_dir = photos_dir or ""
    if not name:
        return ""
    if os.path.isabs(name) and os.path.exists(name):
        return name
    p = os.path.join(photos_dir, name)
    return p if os.path.exists(p) else ""


def _find_fast(driver, field_id, timeout=None):
    """Find a field quickly, while never hiding a dead browser session."""
    try:
        els = driver.find_elements(By.ID, field_id)
    except Exception as exc:
        _raise_if_dead(exc)
        raise
    if els:
        return els[0]
    try:
        return WebDriverWait(driver, timeout or config.QUICK_WAIT).until(
            EC.presence_of_element_located((By.ID, field_id))
        )
    except Exception as exc:
        _raise_if_dead(exc)
        return None


def _upload_and_verify(driver, input_id, path):
    if not path or not os.path.exists(path):
        return False, f"Missing file: {path}"

    el = _find_fast(driver, input_id)
    if el is None:
        return False, f"Field not present on page: {input_id}"

    last_err = None
    for attempt in range(1, config.UPLOAD_RETRIES + 1):
        try:
            driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
            el.send_keys(path)
            # File-input send_keys is synchronous in Selenium/Chrome; verify
            # immediately instead of imposing a fixed delay on every upload.
            files_len = driver.execute_script(
                "return (function(id){var e=document.getElementById(id); return e?e.files.length:0;})(arguments[0]);",
                input_id
            )
            if files_len and int(files_len) > 0:
                return True, f"Uploaded (attempt {attempt}), files.length={files_len}"
            last_err = f"files.length={files_len}"
        except Exception as e:
            _raise_if_dead(e)
            last_err = str(e)
        # element exists but upload didn't register -- brief retry, not a full re-search
        el = _find_fast(driver, input_id, timeout=2)
        if el is None:
            return False, f"Field disappeared from page: {input_id}"
        time.sleep(0.2)
    return False, f"Upload failed after {config.UPLOAD_RETRIES} attempts: {last_err}"


def _fill_field(driver, kind, field_id, value):
    """Fill one field per its FIELD_MAP 'kind'. Returns a short log message,
    or raises on failure so the caller can record it in Notes. Uses
    _find_fast instead of a long WebDriverWait -- see module docstring."""
    el = _find_fast(driver, field_id)
    if el is None:
        raise Exception(f"field not present: {field_id}")

    if kind in ("select", "select_postback"):
        sel = Select(el)
        try:
            current_opt = sel.first_selected_option
            current_value = (current_opt.get_attribute("value") or "").strip()
            current_text = (current_opt.text or "").strip()
        except Exception:
            current_value = current_text = ""

        desired = str(value).strip()
        # Critical speed guard: selecting an already-selected ASP.NET
        # postback dropdown often fires no change event. Waiting for a
        # postback after that would waste PAGE_CHANGE_TIMEOUT seconds.
        if desired in (current_value, current_text):
            return f"{field_id} already {desired} (unchanged)"

        before = page_change_token(driver) if kind == "select_postback" else None
        try:
            sel.select_by_value(desired)
        except Exception:
            sel.select_by_visible_text(desired)
        if kind == "select_postback":
            wait_for_postback(driver, before)
        return f"{field_id} -> {desired}"

    if kind == "text":
        desired = str(value)
        current = el.get_attribute("value") or ""
        if current == desired:
            return f"{field_id} already correct (unchanged)"
        el.clear()
        el.send_keys(desired)
        return f"{field_id} -> {desired}"

    if kind == "checkbox_postback":
        truthy = str(value).strip().lower() in ("yes", "y", "true", "1")
        if not truthy:
            return None
        before = page_change_token(driver)
        if not el.is_selected():
            safe_click(driver, el)
            wait_for_postback(driver, before)
        return f"{field_id} checked"

    return None




def _select_subjects_by_visible_name(driver, subjects):
    """Select ordinary subjects from the live optional-subject radio grid.

    Excel may contain subject *names* only.  We therefore match the visible
    <label> text after the Stream postback has rendered the correct grid; the
    numeric value/code is discovered from the radio element at runtime.
    """
    wanted = [canonical_subject(x) for x in subjects if str(x).strip()]
    if not wanted:
        return [], []

    labels = driver.find_elements(By.CSS_SELECTOR, f"#{config.OPTIONAL_SUBJECT_GRID_ID} label")
    live = {}
    for lab in labels:
        try:
            text = (lab.text or '').strip()
            target = lab.get_attribute('for') or ''
            if text and target:
                live[norm(text)] = (text, target)
        except Exception:
            pass

    selected, missing = [], []
    used_groups = {}
    for subject in wanted:
        match = live.get(norm(subject))
        if not match:
            missing.append(subject)
            continue
        live_text, radio_id = match
        radio = driver.find_element(By.ID, radio_id)
        group_name = radio.get_attribute('name') or radio_id
        if group_name in used_groups and used_groups[group_name] != live_text:
            raise ValueError(
                f"Subject conflict: '{used_groups[group_name]}' and '{live_text}' belong to the same JKBOSE group"
            )
        used_groups[group_name] = live_text
        if not radio.is_selected():
            safe_click(driver, radio)
        selected.append(f"{live_text} [{radio.get_attribute('value') or ''}]")
    return selected, missing


def _subject_code_columns(columns):
    """Return supported legacy/numeric optional-subject code columns in sheet order."""
    out = []
    for c in columns:
        n = norm(c)
        if (n.startswith("group") and n.endswith("choice")) or n.startswith("subjectcode") or n.startswith("subject code"):
            out.append(c)
    return out


def subject_codes_from_row(row):
    """Collect unique numeric optional-subject codes from legacy/code columns."""
    vals = []
    cols = row.index if hasattr(row, "index") else row.keys()
    for c in _subject_code_columns(cols):
        raw = str(row.get(c, "") or "").strip()
        if not raw or raw == "-1":
            continue
        # tolerate values read by Excel as 135.0
        if raw.endswith(".0") and raw[:-2].isdigit():
            raw = raw[:-2]
        if raw not in vals:
            vals.append(raw)
    return vals


def _select_subjects_by_code(driver, codes):
    """Select optional-subject radios by their live HTML value attribute.

    This preserves the legacy Group*_choice workflow while avoiding brittle
    hard-coded radio element IDs.  Stream must already be selected so the
    current grid contains the correct radio values.
    """
    wanted = [str(x).strip() for x in codes if str(x).strip()]
    if not wanted:
        return [], []
    radios = driver.find_elements(By.CSS_SELECTOR, f"#{config.OPTIONAL_SUBJECT_GRID_ID} input[type='radio']")
    live = {}
    for radio in radios:
        try:
            value = (radio.get_attribute("value") or "").strip()
            if value:
                live[value] = radio
        except Exception:
            pass
    selected, missing = [], []
    used_groups = {}
    for code in wanted:
        radio = live.get(code)
        if radio is None:
            missing.append(code)
            continue
        group_name = radio.get_attribute("name") or radio.get_attribute("id") or code
        if group_name in used_groups and used_groups[group_name] != code:
            raise ValueError(
                f"Subject-code conflict: values {used_groups[group_name]} and {code} belong to the same JKBOSE group"
            )
        used_groups[group_name] = code
        if not radio.is_selected():
            safe_click(driver, radio)
        selected.append(code)
    return selected, missing


def build_preflight(excel_file, sheet_name='Sheet2', photos_dir='', class_override=None, subject_input_mode='names'):
    """Parse an input workbook into a compact preview/validation model.

    No browser is opened.  This catches missing registration numbers/files and
    obviously incompatible local subject names before a long submit run.
    Unknown subjects are warnings, not hard failures, because the live JKBOSE
    page is authoritative and may contain newer options than the local catalog.
    """
    df = read_exam_dataframe(excel_file, sheet_name, class_override)
    # Catch a common Excel preparation mistake: pasting the selected photo
    # folder path into row 1, which turns it into an accidental column header.
    suspicious_headers = [str(c) for c in df.columns if (':\\' in str(c) or str(c).startswith('\\\\'))]
    preview = []
    for i, row in df.iterrows():
        row = canonicalize_row_columns(row)
        cls = str(class_override).strip() if class_override else str(row.get('Class','')).strip()
        reg = str(row.get('RegNo','')).strip()
        stream = canonical_stream(row.get('Stream',''))
        mode = str(subject_input_mode or 'names').strip().lower()
        all_subjects = subjects_from_row(row) if (cls != '12' and mode == 'names') else []
        code_subjects = subject_codes_from_row(row) if (cls != '12' and mode == 'codes') else []
        compulsory, vocational, optional = split_subject_roles(all_subjects, stream)
        subjects = all_subjects
        photo = str(row.get('PhotoPath','')).strip() or _build_path_from_helper(str(row.get('pho','')).strip(), photos_dir)
        sign = str(row.get('SignaturePath','')).strip() or _build_path_from_helper(str(row.get('sig','')).strip(), photos_dir)
        issues=[]
        if excluded_reason(row): issues.append(excluded_reason(row))
        if not reg: issues.append('Missing RegNo')
        if cls not in ('11','12'): issues.append('Class must be 11 or 12')
        if suspicious_headers:
            issues.append('Remove accidental path column header: ' + suspicious_headers[0][:55])
        if cls == '11':
            if not stream: issues.append('Missing Stream')
            if row.get('ExportSchema'):
                for field in ('FirstName', 'FatherName', 'MotherName', 'DOB', 'Gender', 'Religion', 'Category', 'Mobile', 'CurrVillage', 'CurrPincode', 'CurrDistrict'):
                    if not str(row.get(field, '')).strip(): issues.append(f'Missing {field}')
                if mode == 'names' and not all_subjects: issues.append('No current subjects supplied')
            if mode == 'names':
                known = SUBJECT_CODES.get(stream)
                if known:
                    for sub in optional:
                        if norm(sub) not in {norm(x) for x in known}:
                            issues.append(f"Check optional subject on live page: {sub}")
                if len(vocational) > 1:
                    issues.append("More than one vocational subject supplied")
                legacy_values = subject_codes_from_row(row)
                if legacy_values and not subjects:
                    issues.append('Code columns contain values, but Subject Input Mode is Names')
            else:
                if not code_subjects and not str(row.get('VocationalCode','')).strip():
                    issues.append('No subject codes found in Group*_choice / SubjectCode* columns')
                if subjects_from_row(row) and not code_subjects:
                    issues.append('Subject-name columns contain values, but Subject Input Mode is Codes')
            if not photo or not os.path.exists(photo): issues.append('Photo missing')
            if not sign or not os.path.exists(sign): issues.append('Signature missing')
        preview.append({
            'row': i+2, 'regno': reg, 'class': cls, 'stream': stream,
            'subjects': ', '.join(subjects), 'compulsory': ', '.join(compulsory),
            'optional': (', '.join(optional) if mode == 'names' else ', '.join(code_subjects)),
            'vocational': (', '.join(vocational) if mode == 'names' else str(row.get('VocationalCode','')).strip()),
            'subject_mode': mode,
            'photo': photo, 'signature': sign, 'issues': '; '.join(issues)
        })
    return preview

def _normalized_status(status):
    """Normalize Excel status text aggressively but predictably.

    Excel/copy-paste can introduce NBSPs, Unicode dashes, zero-width marks and
    repeated spaces.  Completion filtering is safety-critical because a row
    already marked ``Submitted - verified`` must never reach the browser/manual
    prompt again.
    """
    text = unicodedata.normalize('NFKC', str(status or ''))
    # Remove invisible formatting characters (zero-width/BOM/etc.).
    text = ''.join(ch for ch in text if unicodedata.category(ch) != 'Cf')
    text = text.strip().lower()
    text = text.replace('\u00a0', ' ').replace('–', '-').replace('—', '-').replace('−', '-')
    text = re.sub(r'\s+', ' ', text)
    return text


def _status_is_completed(status):
    """Return True only for statuses that are safe to skip before browser work.

    The primary/authoritative completed status is ``Submitted - verified``.
    Spacing, dash style, punctuation, NBSPs and case are ignored.  We also keep
    ``Already Submitted`` because that status is written only after an explicit
    JKBOSE already-completed message.  Anything containing ``unverified`` is
    always retryable.
    """
    text = _normalized_status(status)
    if not text:
        return False
    words = re.sub(r'[^a-z0-9]+', ' ', text).strip()
    if 'unverified' in words or words.startswith(('pending', 'error', 'failed', 'not submitted', 'not verified')):
        return False
    return words in {
        'submitted verified',
        'verified submitted',
        'already submitted',
    }



def _normalize_regno(value):
    """Normalize a registration number read from Excel/UI for exact row matching."""
    text = unicodedata.normalize('NFKC', str(value or ''))
    text = ''.join(ch for ch in text if unicodedata.category(ch) != 'Cf')
    text = text.replace('\u00a0', ' ').strip()
    # Excel sometimes exposes integer-like cells as ``12345.0`` through mixed sheets.
    if re.fullmatch(r'[A-Za-z0-9]+\.0', text):
        text = text[:-2]
    return text


def _find_resume_after_index(df, last_submitted_regno):
    """Return the dataframe index of an exact Last-submitted RegNo match.

    The match is deliberately exact after light normalization.  A missing or
    duplicate registration number is treated as an input error instead of
    guessing a resume point.
    """
    target = _normalize_regno(last_submitted_regno)
    if not target:
        return None
    matches = []
    for idx, row in df.iterrows():
        if _normalize_regno(row.get('RegNo', '')) == target:
            matches.append(idx)
    if not matches:
        raise ValueError(
            f"Last submitted RegNo '{target}' was not found in the selected sheet. "
            "Nothing was submitted; check the RegNo/sheet and try again."
        )
    if len(matches) > 1:
        rows = ', '.join(str(int(i) + 2) if str(i).isdigit() else str(i) for i in matches[:8])
        raise ValueError(
            f"Last submitted RegNo '{target}' occurs more than once in the selected sheet "
            f"(Excel rows {rows}). Remove duplicates or choose a unique resume point."
        )
    return matches[0]


def _scroll_to_submit_area(driver):
    """Bring captcha/Submit into view after a student has been fetched/filled."""
    try:
        target = None
        for element_id in (config.CAPTCHA_INPUT_ID, config.SUBMIT_BTN_ID):
            els = driver.find_elements(By.ID, element_id)
            if els:
                target = els[0]
                break
        if target is not None:
            driver.execute_script(
                "arguments[0].scrollIntoView({behavior:'auto', block:'center', inline:'nearest'});",
                target,
            )
        else:
            driver.execute_script(
                "window.scrollTo({top: Math.max(document.body.scrollHeight, document.documentElement.scrollHeight), "
                "behavior:'auto'});"
            )
        return True
    except Exception as exc:
        _raise_if_dead(exc)
        return False

def _safe_element_value(driver, element_id):
    try:
        els = driver.find_elements(By.ID, element_id)
        if not els:
            return None
        return (els[0].get_attribute('value') or '').strip()
    except Exception as exc:
        _raise_if_dead(exc)
        return None


def _safe_named_value(driver, name):
    try:
        els = driver.find_elements(By.NAME, name)
        if not els:
            return None
        return (els[0].get_attribute('value') or '')
    except Exception as exc:
        _raise_if_dead(exc)
        return None


def _install_submit_click_marker(driver):
    """Mark the live Submit button without changing the page's own handler."""
    try:
        driver.execute_script(
            """
            window.__jkboseManualSubmitClicked = false;
            var b = document.getElementById(arguments[0]);
            if (b && !b.__jkboseMarkerInstalled) {
                b.addEventListener('click', function(){ window.__jkboseManualSubmitClicked = true; }, true);
                b.__jkboseMarkerInstalled = true;
            }
            """,
            config.SUBMIT_BTN_ID,
        )
    except Exception as exc:
        _raise_if_dead(exc)


def _manual_submit_clicked(driver):
    try:
        return bool(driver.execute_script('return !!window.__jkboseManualSubmitClicked;'))
    except Exception as exc:
        # A successful form postback can replace the document and therefore the
        # marker itself.  That is not a browser failure; return False and let the
        # page-state checks below decide whether submission completed.
        _raise_if_dead(exc)
        return False


def _wait_for_manual_submission(driver, regno, stop_event=None, timeout=None):
    """Wait quietly until JKBOSE *verifiably* completes the current student.

    Safety rule: uncertainty NEVER advances the batch.  In GUI mode this wait
    has no automatic timeout; it stays on the same student until one of these
    happens:
      * JKBOSE shows an explicit submitted/already-submitted completion state;
      * the user presses Stop (returns ``None`` so the caller can end cleanly).

    A wrong captcha keeps the same record active. A reset, expired session or
    ambiguous navigation stops for manual review; it never proves submission.
    """
    start = time.monotonic()
    try:
        initial_url = driver.current_url or ''
    except Exception as exc:
        _raise_if_dead(exc)
        initial_url = ''
    initial_reg = _safe_element_value(driver, 'bodypH_txt_rege')
    initial_viewstate = _safe_named_value(driver, '__VIEWSTATE')
    _install_submit_click_marker(driver)
    saw_click = False
    last_hint = 0.0
    last_timeout_hint = 0.0

    print('[WAIT] Solve captcha and click JKBOSE Submit in Chrome. Waiting on this student until submission is verified.')

    while True:
        if stop_event is not None and stop_event.is_set():
            return None, 'Stopped by user while waiting for manual submission'
        _ensure_driver_alive(driver)

        # Strongest signal: explicit JKBOSE completion/already-submitted text.
        try:
            saw_click = saw_click or _manual_submit_clicked(driver)
            context = getattr(driver, '_exam_lookup_context', ('', ''))
            done, msg = _detect_already_submitted(driver, regno, context[1], allow_manual_success=saw_click)
            if done:
                return True, msg or 'JKBOSE reported submission complete'
        except ExamPortalUnavailable:
            raise
        except Exception as exc:
            _raise_if_dead(exc)

        if not saw_click:
            try:
                saw_click = _manual_submit_clicked(driver)
            except Exception as exc:
                _raise_if_dead(exc)

        current_reg = _safe_element_value(driver, 'bodypH_txt_rege')
        try:
            current_url = driver.current_url or ''
        except Exception as exc:
            _raise_if_dead(exc)
            current_url = initial_url

        # A cleared field or navigation can mean expiry or a failed postback.
        # Without a student-bound completion message this state is uncertain.
        if current_reg is None:
            submit_exists = True
            captcha_exists = True
            try:
                submit_exists = bool(driver.find_elements(By.ID, config.SUBMIT_BTN_ID))
                captcha_exists = bool(driver.find_elements(By.ID, config.CAPTCHA_INPUT_ID))
            except Exception as exc:
                _raise_if_dead(exc)
            if current_url != initial_url or (not submit_exists and not captcha_exists):
                raise ExamPortalUnavailable('Student form closed or navigated away without a verified submission response. Review manually before retrying.')
        elif regno and current_reg != str(regno).strip():
            raise ExamPortalUnavailable('Registration field reset without a verified submission response. Review manually before retrying.')

        current_viewstate = _safe_named_value(driver, '__VIEWSTATE')
        if initial_viewstate and current_viewstate and current_viewstate != initial_viewstate:
            try:
                captcha_exists = bool(driver.find_elements(By.ID, config.CAPTCHA_INPUT_ID))
                submit_exists = bool(driver.find_elements(By.ID, config.SUBMIT_BTN_ID))
            except Exception as exc:
                _raise_if_dead(exc)
                captcha_exists = submit_exists = True
            if not captcha_exists and not submit_exists:
                raise ExamPortalUnavailable('Form controls cleared without a verified submission response. Review manually before retrying.')

        # ASP.NET can replace the document after a failed captcha/postback, so
        # re-install the click marker while the same pending form remains.
        if not saw_click:
            _install_submit_click_marker(driver)

        elapsed = time.monotonic() - start
        if elapsed - last_hint >= 60:
            print('[WAIT] Still waiting on the same student — no verified submission detected yet.')
            last_hint = elapsed

        # Optional diagnostic threshold: do NOT return/advance when reached.
        if timeout is not None and elapsed >= timeout and elapsed - last_timeout_hint >= 300:
            print('[WAIT] Verification is still uncertain; staying on this student (nothing written to Excel).')
            last_timeout_hint = elapsed

        time.sleep(0.45)



def run(driver, excel_file: str, sheet_name: str, photos_dir: str,
        pause_callback=input, stop_event=None, class_override: str = None,
        subject_input_mode: str = 'names', auto_skip_submitted: bool = True,
        last_submitted_regno: str = None, history=None):
    """
    driver           : an already-logged-in Selenium driver, already inside the
                        Exam/RR section (see browser.login_and_enter_exam_section).
    excel_file       : path to the workbook with one row per student. Columns
                        should match config.FIELD_MAP's excel_column names
                        (e.g. RegNo, Class, Stream, FirstName, CurrPincode, ...),
                        plus PhotoPath/SignaturePath or pho/sig helper columns.
    sheet_name       : sheet to read/write within excel_file.
    photos_dir       : folder containing photo/signature files referenced by
                        the 'pho'/'sig' helper columns.
    pause_callback   : optional legacy callback used by CLI/manual workflows.
                        When None (the GUI default), the tool waits quietly for
                        the JKBOSE Submit action and advances automatically; no
                        modal Continue dialog is shown.
    stop_event       : optional threading.Event, checked between rows so a
                        caller (e.g. a GUI Stop button) can end the batch
                        early without killing the process.
    class_override   : "11" or "12" -- if given, used for every row instead
                        of that row's own 'Class' column. The Class dropdown
                        is set BEFORE the registration number is entered,
                        since the site requires that order.
    last_submitted_regno : optional manual resume boundary. When supplied, the
                        exact matching Excel row and every row above it are
                        skipped unconditionally, and EVERY row below it enters
                        the work queue regardless of stale Status values. This
                        is the safest recovery control after a partial run.
    auto_skip_submitted : when True (default), after Get Details the
                        live JKBOSE page/alert is checked for an explicit
                        already-submitted/already-renewed message. This applies
                        to BOTH Class 11 and Class 12; detected rows are marked
                        'Already Submitted' and continued automatically without
                        showing the manual Continue prompt.

    NOTE on Class 12: a Class-12 row is a renewal of an already-fully-
    registered Class-11 student, so nothing needs (re-)filling -- the flow
    is just RegNo -> Get Details -> captcha -> Submit. Rows resolving to
    Class 12 (via class_override or the row's own 'Class' column) skip
    FIELD_MAP filling and the photo/signature upload entirely. Class 11
    (or any other value) keeps the full fill-everything behavior.
    """
    wait = WebDriverWait(driver, config.ELEMENT_WAIT)

    df = read_exam_dataframe(excel_file, sheet_name, class_override)
    if "Status" not in df.columns:
        df["Status"] = ""
    if "Notes" not in df.columns:
        df["Notes"] = ""
    print(f"[INFO] Loaded {len(df)} rows from {excel_file} sheet '{sheet_name}'")
    progress = _ExcelProgressWriter(
        excel_file, sheet_name, getattr(config, "PROGRESS_SAVE_EVERY", 5), history=history
    )
    processed_this_run = 0
    interrupted_at = None

    # Build the work queue once. A manual Last-submitted RegNo is the
    # authoritative resume boundary when provided: rows through that exact row
    # are skipped and ALL rows below it are queued regardless of stale statuses.
    # Without a manual boundary, verified/completed Excel statuses are filtered.
    completed_rows = []
    pending_rows = []
    resume_idx = _find_resume_after_index(df, last_submitted_regno)
    if resume_idx is not None:
        resume_reg = _normalize_regno(last_submitted_regno)
        all_indices = list(df.index)
        try:
            pos = all_indices.index(resume_idx)
        except ValueError:
            raise ValueError(f"Could not resolve resume row for RegNo {resume_reg}")
        skipped_indices = all_indices[:pos + 1]
        pending_rows = all_indices[pos + 1:]
        print(
            f"[INFO] Manual resume point: RegNo {resume_reg} at Excel row {pos + 2}. "
            f"Skipping {len(skipped_indices)} row(s) through it; processing {len(pending_rows)} row(s) below it."
        )
        if any(_status_is_completed(df.loc[idx].get('Status', '')) for idx in pending_rows):
            print("[INFO] Manual resume point takes priority over Status below the boundary; those rows will be checked again.")
    else:
        for _i, _row in df.iterrows():
            _status = str(_row.get("Status", "")).strip()
            if _status_is_completed(_status):
                completed_rows.append((_i, _normalize_regno(_row.get("RegNo", "")), _status))
            else:
                pending_rows.append(_i)
        print(f"[INFO] Pre-scan: {len(completed_rows)} completed rows auto-skipped; {len(pending_rows)} rows need checking.")
        if completed_rows:
            sample = ', '.join(str(x[1]) for x in completed_rows[:5] if x[1])
            if sample:
                print(f"[INFO] Completed rows skipped before browser work (sample): {sample}" + (" ..." if len(completed_rows) > 5 else ""))

    for i in pending_rows:
        row = df.loc[i]
        row = canonicalize_row_columns(row)
        if excluded_reason(row):
            print(f"[SKIP] Excel row {i + 2}: {excluded_reason(row)}")
            if history: history.progress(i, 'Excluded', excluded_reason(row))
            continue
        if stop_event is not None and stop_event.is_set():
            print("[INFO] Stop requested -- ending batch before the next row.")
            break

        regno = _normalize_regno(row.get("RegNo", ""))
        if not regno:
            df.at[i, "Status"] = "Skipped: no RegNo"
            progress.update(i, "Skipped: no RegNo", "")
            continue

        print(f"\n[ROW {i + 1}/{len(df)}] RegNo={regno} — starting")
        notes = []
        try:
            _assert_browser_alive(driver)
            if history: history.bind_account(driver)
            driver.switch_to.default_content()
            require_exam_form(driver)

            # Avoid an unconditional refresh + fixed sleep for every student.
            # After a successful manual submit JKBOSE normally leaves the form
            # usable for the next lookup. Refresh only as recovery if the core
            # RegNo field is actually absent.
            if not driver.find_elements(By.ID, "bodypH_txt_rege"):
                print("[INFO] Form not ready after previous record -- refreshing once")
                driver.refresh()
                wait.until(EC.presence_of_element_located((By.ID, "bodypH_txt_rege")))

            # ---------- Class must be set BEFORE the registration lookup ----------
            class_val = str(class_override).strip() if class_override else str(row.get("Class", "")).strip()
            if class_val and class_val != "-1":
                try:
                    sel_elem = wait.until(
                        EC.presence_of_element_located((By.ID, config.CLASS_SELECT_ID))
                    )
                    sel_class = Select(sel_elem)
                    try:
                        current = sel_class.first_selected_option.get_attribute("value")
                    except Exception:
                        current = None
                    if current == class_val:
                        print(f"[INFO] Class already set to {class_val} -- skipping re-select")
                    else:
                        before = page_change_token(driver)
                        try:
                            sel_class.select_by_value(class_val)
                        except Exception:
                            sel_class.select_by_visible_text(class_val)
                        wait_for_postback(driver, before)
                        print(f"[INFO] Class set to {class_val} (postback waited)")
                except Exception as e:
                    _raise_if_dead(e)
                    raise ExamPortalUnavailable(f'Class selection failed: {e}') from e
            else:
                raise ExamPortalUnavailable('No valid Class value; lookup stopped.')

            require_exam_form(driver, class_val)

            # ---------- Look up the student ----------
            el_reg = wait.until(EC.presence_of_element_located((By.ID, "bodypH_txt_rege")))
            el_reg.clear()
            el_reg.send_keys(regno)
            print("[INFO] Registration number entered")

            before = page_change_token(driver)
            driver._exam_lookup_context = (regno, class_val)
            try:
                btn_get = driver.find_element(By.ID, config.GET_DETAILS_BTN_ID)
                safe_click(driver, btn_get)
                if not wait_for_postback(driver, before):
                    raise ExamPortalUnavailable('Get Details response was not confirmed.')
                print("[INFO] Clicked 'Get Details' and waited for postback")
                if _scroll_to_submit_area(driver):
                    print("[INFO] Scrolled to captcha/Submit area after Get Details")
            except Exception as e:
                _raise_if_dead(e)
                raise ExamPortalUnavailable(f'Get Details failed: {e}') from e

            # ---------- Auto-skip records JKBOSE says are already submitted ----------
            # Useful for both Class 11 and Class 12 batches containing a mix of
            # already-completed and still-pending students.  A strong site message
            # is required; otherwise the normal manual captcha/Submit flow remains.
            if auto_skip_submitted:
                already_done, done_message = _detect_already_submitted(driver, regno, class_val)
                if already_done:
                    status = "Already Submitted"
                    notes_text = done_message[:500]
                    df.at[i, "Status"] = status
                    df.at[i, "Notes"] = notes_text
                    progress.update(i, status, notes_text)
                    processed_this_run += 1
                    print(f"[ROW {i + 1}] JKBOSE reports this record is already submitted — auto-skipping.")
                    if done_message:
                        print(f"[INFO] Site message: {done_message}")
                    continue

            require_lookup_details(driver, regno, row)

            # ---------- Fill fields, unless this is a Class 12 renewal ----------
            # For Class 12, the student was already fully registered in
            # Class 11 -- nothing needs (re-)filling here. Get Details ->
            # captcha -> Submit is the whole job. Filling/uploading in that
            # case would be pointless at best and risk overwriting already-
            # correct data with blank/stale Excel columns at worst.
            if class_val == "12":
                print("[INFO] Class 12 renewal -- skipping field fill and photo/signature "
                      "upload (Get Details already has everything needed).")
            else:
                # ---------- Fill every mapped field present in this row ----------
                # Stream accepts friendly names (Science/Humanities/...) as well as
                # site codes (S/H/...).  Convert names to the current site value.
                row_for_fill = row.copy()
                if str(row_for_fill.get("Stream", "")).strip():
                    row_for_fill["Stream"] = stream_value(row_for_fill.get("Stream", ""))

                for col, field_id, kind in config.FIELD_MAP:
                    if kind in ("lookup", "class_preselect"):
                        continue  # RegNo and Class already handled above
                    val = str(row_for_fill.get(col, "")).strip()
                    if not val and col in ('MiddleName', 'LastName') and str(row_for_fill.get('Name', '')).strip():
                        try:
                            field = driver.find_element(By.ID, field_id)
                            field.clear()
                        except Exception as e:
                            _raise_if_dead(e)
                            notes.append(f'{col} could not be cleared: {e}')
                        continue
                    if not val or val == "-1":
                        continue
                    try:
                        msg = _fill_field(driver, kind, field_id, val)
                        if msg:
                            print(f"[INFO] {msg}")
                    except Exception as e:
                        _raise_if_dead(e)
                        print(f"[WARN] {col} ({field_id}) skipped/failed: {e}")
                        notes.append(f"{col} failed: {e}")

                # ---------- Subjects: names (default) OR legacy numeric codes ----------
                mode = str(subject_input_mode or "names").strip().lower()
                if mode == "codes":
                    codes = subject_codes_from_row(row)
                    if codes:
                        try:
                            selected, missing = _select_subjects_by_code(driver, codes)
                            if selected:
                                print("[INFO] Optional subject codes selected: " + ", ".join(selected))
                            if missing:
                                msg = "Optional subject codes not found in current Stream grid: " + ", ".join(missing)
                                print("[WARN] " + msg)
                                notes.append(msg)
                        except Exception as e:
                            _raise_if_dead(e)
                            print(f"[WARN] Subject-code selection failed: {e}")
                            notes.append(f"Subject-code selection failed: {e}")
                    else:
                        print("[INFO] Subject Input Mode=Codes; no optional Group*_choice/SubjectCode* values supplied")
                    # VocationalCode is filled by FIELD_MAP above, preserving the legacy workflow.
                else:
                    all_subjects = subjects_from_row(row)
                    compulsory_subjects, vocational_subjects, optional_subjects = split_subject_roles(all_subjects, row.get("Stream", ""))
                    if compulsory_subjects:
                        print("[INFO] Compulsory subjects (no click needed): " + ", ".join(compulsory_subjects))
                    if optional_subjects:
                        try:
                            selected, missing = _select_subjects_by_visible_name(driver, optional_subjects)
                            if selected:
                                print("[INFO] Optional subjects selected: " + ", ".join(selected))
                            if missing:
                                msg = "Optional subjects not found in current Stream grid: " + ", ".join(missing)
                                print("[WARN] " + msg)
                                notes.append(msg)
                        except Exception as e:
                            _raise_if_dead(e)
                            print(f"[WARN] Subject selection failed: {e}")
                            notes.append(f"Subject selection failed: {e}")

                    # Vocational name may be in VocationalSubject OR directly in a
                    # normal sub column (for example sub5='IT and ITES'). Numeric
                    # VocationalCode remains supported, but name mode prefers names.
                    explicit_voc = str(row.get("VocationalSubject", "")).strip()
                    voc_name = explicit_voc or (vocational_subjects[0] if vocational_subjects else "")
                    if len(vocational_subjects) > 1:
                        notes.append("Multiple vocational subjects supplied; using first: " + vocational_subjects[0])
                    if voc_name and not str(row.get("VocationalCode", "")).strip():
                        try:
                            msg = _fill_field(driver, "select", config.VOCATIONAL_SELECT_ID, voc_name)
                            if msg: print(f"[INFO] {msg}")
                        except Exception as e:
                            _raise_if_dead(e)
                            print(f"[WARN] VocationalSubject failed: {e}")
                            notes.append(f"VocationalSubject failed: {e}")

                # ---------- Photo & signature ----------
                photo_path = str(row.get("PhotoPath", "")).strip() or _build_path_from_helper(str(row.get("pho", "")).strip(), photos_dir)
                sign_path = str(row.get("SignaturePath", "")).strip() or _build_path_from_helper(str(row.get("sig", "")).strip(), photos_dir)

                ok_photo, msg_photo = _upload_and_verify(driver, config.PHOTO_UPLOAD_ID, photo_path)
                ok_sign, msg_sign = _upload_and_verify(driver, config.SIGNATURE_UPLOAD_ID, sign_path)
                print(f"[INFO] Photo: {msg_photo}; Sign: {msg_sign}")
                if not ok_photo:
                    notes.append("Photo upload failed")
                if not ok_sign:
                    notes.append("Signature upload failed")

            if notes:
                raise ExamPortalUnavailable('Incomplete fields/subjects/uploads require manual review: ' + ' | '.join(notes))

            require_exam_form(driver, class_val)
            # ---------- Manual captcha / submit ----------
            _scroll_to_submit_area(driver)
            print(f"[ROW {i + 1}] Form ready for RegNo {regno}. Captcha/Submit area brought into view.")

            if pause_callback is None:
                # GUI mode: no modal popup. Watch the live JKBOSE form and advance
                # automatically once the user's manual Submit completes.
                #
                # IMPORTANT: a non-browser/internal watcher error must NEVER make
                # us leave this student and fall through to driver cleanup. Retry
                # the watcher on the same live form; only a real dead-browser
                # exception or an explicit Stop may end the wait.
                watcher_error_count = 0
                while True:
                    try:
                        verified, verify_message = _wait_for_manual_submission(
                            driver, regno, stop_event=stop_event
                        )
                        break
                    except BrowserSessionLost:
                        raise
                    except Exception as exc:
                        _raise_if_dead(exc)
                        if isinstance(exc, ExamPortalUnavailable):
                            raise
                        watcher_error_count += 1
                        if watcher_error_count == 1 or watcher_error_count % 6 == 0:
                            print(
                                f"[WARN] Internal submission watcher error for RegNo {regno}: {exc}. "
                                "Staying on this student and retrying; Excel row is unchanged."
                            )
                        if stop_event is not None and stop_event.is_set():
                            verified, verify_message = None, 'Stopped by user while watcher was retrying'
                            break
                        time.sleep(2.0)
            else:
                # CLI/backward-compatible mode: keep the caller's explicit pause.
                pause_callback("After manual submit (or skip), press ENTER to continue to next record...")
                try:
                    verified, verify_message = _detect_submit_success(driver, regno, class_val)
                except Exception as exc:
                    _raise_if_dead(exc)
                    verified, verify_message = False, ''
            if verified is None:
                # User pressed Stop while this student was pending.  Preserve the
                # row exactly as it was: no synthetic status/note, no advancement.
                print(f"[STOP] Stopped while RegNo {regno} was still pending. Row left unchanged.")
                break

            if verified:
                status = "Submitted - verified"
                if verify_message:
                    notes.append(verify_message[:500])
                notes_text = " | ".join(notes)
                df.at[i, "Status"] = status
                df.at[i, "Notes"] = notes_text
                progress.update(i, status, notes_text)
                processed_this_run += 1
                print(f"[ROW {i + 1}] Submission verified — future runs will auto-skip this record.")
            else:
                # This path is retained only for legacy/CLI pause callbacks.
                # Never stamp an unverified row and never pretend it completed.
                print(f"[ROW {i + 1}] Submission not verified; row left unchanged.")

        except ExamPortalUnavailable as e:
            progress.update(i, 'Blocked - manual review', str(e)[:1000])
            print(f'[STOP] {e}. Remaining rows were not attempted.')
            break
        except BrowserSessionLost as e:
            interrupted_at = (i, regno, str(e))
            print(f"[FATAL] Browser connection lost at row {i + 1} (RegNo {regno}).")
            print("[FATAL] The remaining rows were NOT attempted. Excel progress is being saved.")
            print("[INFO] Run Submit Forms again after Chrome/JKBOSE is available; rows already marked 'Submitted - verified' will be skipped automatically.")
            break
        except Exception as e:
            if _is_dead_browser_exception(e):
                interrupted_at = (i, regno, str(e))
                print(f"[FATAL] Browser connection lost at row {i + 1} (RegNo {regno}).")
                print(f"[FATAL] {str(e).splitlines()[0] if str(e) else type(e).__name__}")
                print("[FATAL] Stopping now instead of generating the same error for every remaining row.")
                print("[INFO] Excel progress is saved. Re-run later to resume automatically from unfinished rows.")
                break

            print(f"[ERROR] Exception processing RegNo {regno}: {e}")
            traceback.print_exc()
            status = f"Error: {str(e)[:80]}"
            df.at[i, "Status"] = status
            df.at[i, "Notes"] = "Exception occurred"
            progress.update(i, status, "Exception occurred")
            processed_this_run += 1
            continue

    progress.save()
    if interrupted_at is not None:
        i, regno, reason = interrupted_at
        print(f"[STOPPED] Batch interrupted at row {i + 1}/{len(df)} (RegNo {regno}).")
        print(f"[DONE] Progress saved. Successfully/ordinarily processed this run: {processed_this_run}.")
        raise BrowserSessionLost(
            f"Browser disconnected at row {i + 1}/{len(df)} (RegNo {regno}). "
            "Progress saved; run again to resume.",
            row_index=i, regno=regno,
        )
    else:
        print(f"[DONE] Batch finished. Excel updated (Status & Notes). Processed this run: {processed_this_run}.")
