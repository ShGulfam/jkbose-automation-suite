"""Immutable selected-export snapshots and changes for the tracking workbook."""
from pathlib import Path
from datetime import datetime,timezone
from decimal import Decimal,InvalidOperation
import hashlib,json,os,re,time
from .internal_assessment import Audit,class_key,session_key,subject_name
from .internal_assessment_history import storage_root

HEADERS=['Class','Session','Evaluation Type','Subject Code','Subject Name','Board Reg. No.','Exam Roll No.','Class Roll No.','Student Name','Marks','Max Marks','Attendance','Source Issues','Export Schema','Source Award ID','Saved Award Student Name']

def snapshot(path):
    from openpyxl import load_workbook
    path=Path(path).resolve(strict=True);digest=hashlib.sha256(path.read_bytes()).hexdigest()
    book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        if 'JKBOSE_Internal_Assessment' not in book.sheetnames:return None
        sheet=book['JKBOSE_Internal_Assessment']
        if sheet.max_row>20000 or sheet.max_column>300:raise ValueError('Export dimensions exceed source-history limits.')
        rows=list(sheet.iter_rows());headers=[c.value for c in rows[0]][:16]
        legacy=headers[:15]==HEADERS[:15] and (len(headers)==15 or headers[15] in (None,'Submission Status'))
        if headers!=HEADERS and not legacy:raise ValueError('Consolidated handoff columns changed; source-version comparison requires the known schema.')
        values=[];seen=set()
        for row in rows[1:]:
            if all(c.value is None for c in row[:16]):continue
            if any(c.data_type=='f' for c in row[:16]):raise ValueError('Formula in source-version handoff; export values.')
            raw=[c.value for c in row[:16]];raw += [None]*(16-len(raw))
            if legacy:raw[15]=None
            if raw[13]!='hss-jkbose-ia-v1':raise ValueError('Unknown source-version schema.')
            key=(str(raw[3]),str(raw[5]))
            if key in seen:raise ValueError('Duplicate student/subject in source-version snapshot.')
            seen.add(key);values.append(raw)
        if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:raise ValueError('Workbook changed while recording export history.')
        return dict(path=str(path),filename=path.name,source_sha256=digest,headers=headers,rows=values)
    finally:book.close()

def comparable(value,index):
    if index in {9,10} and value not in (None,''):
        try:
            n=Decimal(str(value))
            if n.is_finite():return str(n.normalize())
        except InvalidOperation:pass
    if index==11:
        text=str(value or '').strip().lower()
        if text in {'present','p'}:return 'P'
        if text in {'absent','a','ab','abs'}:return 'A'
    return '' if value is None else str(value)

def compare(before,current):
    old={(str(r[3]),str(r[5])):r for r in before['rows']} if before else {};new={(str(r[3]),str(r[5])):r for r in current['rows']};changes=[]
    if before is None:return changes
    for key in sorted(old.keys()|new.keys()):
        previous=old.get(key);row=new.get(key);identity=row or previous
        base=dict(subject=identity[4],registration=str(identity[5]),exam_roll=str(identity[6]),name=identity[8])
        if previous is None or row is None:
            changes.append(dict(base,change='Added' if previous is None else 'Removed',field='Student subject award',previous='' if previous is None else str(previous[9])+' / '+str(previous[10])+' '+str(previous[11]),current='' if row is None else str(row[9])+' / '+str(row[10])+' '+str(row[11])));continue
        for index in [6,7,8,9,10,11,12,14,15]:
            if comparable(previous[index],index)!=comparable(row[index],index):changes.append(dict(base,change='Updated',field=HEADERS[index],previous=previous[index],current=row[index]))
    return changes

def record_source_version(source,root,student_class,session,account):
    current=snapshot(source)
    if current is None:return None
    cls=class_key(student_class);session=session_key(session)
    if not re.fullmatch(r'[A-Za-z0-9_-]{3,80}',account):raise ValueError('Source history requires a verified school account.')
    if any(class_key(r[0])!=cls or session_key(r[1])!=session for r in current['rows']):raise ValueError('Source-version cohort differs from selected class/session.')
    scope=dict(student_class=cls,session=session,account=account.lower());key=Audit._hash(scope)[:12]
    folder=storage_root(root)/'_source_versions'/f'Class {cls} - {session} - {key}';folder.mkdir(parents=True,exist_ok=True)
    index=folder/'versions.jsonl';previous='0'*64;last=None
    if index.exists():
        for line in index.read_text(encoding='utf-8').splitlines():
            e=json.loads(line);signature=e.pop('hash')
            if e.get('previous')!=previous or Audit._hash(e)!=signature or e.get('scope')!=scope:raise ValueError('Source-version history integrity failed.')
            name=e['file']
            if Path(name).name!=name:raise ValueError('Invalid source-version snapshot path.')
            raw=(folder/name).read_bytes()
            if hashlib.sha256(raw).hexdigest()!=e['snapshot_sha256']:raise ValueError('Source-version snapshot was changed.')
            last=json.loads(raw);previous=signature
    if last and last['current']['source_sha256']==current['source_sha256']:return last
    baseline=last['current'] if last else None
    if baseline is None:
        tracker=Path(source).with_name(Path(source).stem+'_Submission_Tracking.xlsx')
        if tracker.is_file():
            baseline=snapshot(tracker)
            if baseline and any(class_key(r[0])!=cls or session_key(r[1])!=session for r in baseline['rows']):raise ValueError('Previous tracking cohort differs from selected class/session.')
    result=dict(scope=scope,time=datetime.now(timezone.utc).isoformat(),previous_source=baseline['filename'] if baseline else '',previous_sha256=baseline['source_sha256'] if baseline else '',current=current,changes=compare(baseline,current))
    raw=json.dumps(result,indent=2).encode();name=f'{time.time_ns()}-{current["source_sha256"][:12]}.json';path=folder/name
    with path.open('xb') as stream:stream.write(raw);stream.flush();os.fsync(stream.fileno())
    event=dict(kind='source_version',scope=scope,file=name,snapshot_sha256=hashlib.sha256(raw).hexdigest(),source_sha256=current['source_sha256'],previous=previous);event['hash']=Audit._hash(event)
    with index.open('a',encoding='utf-8') as stream:stream.write(json.dumps(event)+'\n');stream.flush();os.fsync(stream.fileno())
    return result

def add_source_changes_sheet(book,version,statuses):
    if not version:return
    from openpyxl.styles import Font,PatternFill,Alignment
    sheet=book.create_sheet('Source_Changes');sheet.append(['Changes since previous hssshangus.in consolidated export'])
    sheet.append(['Previous export',version['previous_source'] or 'First selected export; baseline recorded'])
    sheet.append(['Current export',version['current']['filename']]);sheet.append(['New values apply only to pending or rejected awards. Recorded awards are reported without a retry.'])
    sheet.append(['S.No.','Subject S.No.','Subject','Registration No.','Exam Roll No.','Student Name','Change','Field','Previous Value','Current Value','Submission State','Action'])
    groups={}
    for i,item in enumerate(version['changes'],1):
        subject=subject_name(item['subject']);subject='Biology' if subject in {'Botany','Zoology'} else subject;groups[subject]=groups.get(subject,0)+1
        state=statuses.get((subject,item['registration']),['Not in current source'])[0]
        action='Already recorded - report only; never retry' if state=='Submitted' else 'Validate current values in pending portal before resubmission' if state=='Rejected' else 'Resolve uncertain history before any save' if state=='Uncertain' else 'Not in current export; no portal action inferred' if item['change']=='Removed' else 'Use current export after live batch validation'
        sheet.append([i,groups[subject],item['subject'],item['registration'],item['exam_roll'],item['name'],item['change'],item['field'],item['previous'],item['current'],state,action])
    if not version['changes']:sheet.append([None,None,None,None,None,None,'No award values changed' if version['previous_source'] else 'Baseline recorded'])
    for c in sheet[5]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='334155');c.alignment=Alignment(wrap_text=True)
    for i,width in enumerate([9,14,25,24,18,30,14,27,35,35,20,65],1):sheet.column_dimensions[sheet.cell(5,i).column_letter].width=width
    for row in sheet.iter_rows(min_row=6):
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        for c in [row[3],row[4]]:c.number_format='@'
    sheet.freeze_panes='G6';sheet.auto_filter.ref=f'A5:L{sheet.max_row}'
